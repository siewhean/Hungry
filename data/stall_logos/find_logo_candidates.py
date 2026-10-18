#!/usr/bin/env python3
"""Find rights-aware stall logo candidates (read-only; never downloads or rehosts images).

Inputs (all read-only):
  --venues   JSON list of venues: id, name, website_url (fetched via anon REST beforehand)
  --overpass optional Overpass JSON cache with OSM brand / brand:wikidata tags
Outputs:
  stall_logo_candidates.xlsx / .csv in --out

Sources, in priority order:
  a) Wikidata P154 (logo image) -> Wikimedia Commons file URL + licence (extmetadata)
  b) Official website homepage: <link rel=icon|apple-touch-icon>, og:logo / og:image,
     schema.org Organization.logo JSON-LD.  One homepage GET per host, robots.txt respected,
     1 request/second per host, 10 s timeout.  URLs are recorded only.
  c) Everything else: no candidate -> merchant upload via the merchant portal.
"""
import argparse, csv, html, json, re, sys, threading, time, urllib.parse, urllib.request, urllib.robotparser
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser

UA = "HungryBot (contact placeholder)"
WIKI_UA = UA + " python-urllib; Wikimedia API read-only lookups"
CHAIN_MIN = 3
KNOWN_BRANDS = {  # always looked up even if <3 outlets in data
    "mcdonald's", "kfc", "saizeriya", "toast box", "ya kun kaya toast", "old chang kee", "subway",
    "starbucks", "jollibee", "din tai fung", "koufu", "kopitiam", "fun toast", "burger king",
    "pizza hut", "breadtalk", "mos burger", "popeyes", "sukiya", "sushiro", "4fingers",
}
# Brand-name merges (normalized variant -> canonical normalized key)
MERGE = {"starbucks coffee": "starbucks", "roasted delight": "roasted delights", "koi the": "koi",
         "toastbox": "toast box", "domino's pizza": "domino's", "7 eleven": "7-eleven",
         "boost juice bars": "boost juice"}
GENERIC = set("""rice chicken mixed vegetable vegetables veg drink drinks stall stalls seafood fish soup noodle
noodles mee western food foods indian malay chinese nasi padang yong tau foo mala hotpot beverage beverages
economic economy fried roasted roast duck char kway teow prawn hokkien laksa bak chor wanton wonton satay
prata roti cai fan zi cze kopi coffee shop nil na n a unknown kiosk canteen cafe dessert desserts juice
fruit fruits porridge dim sum ban mian fishball lor carrot cake chap thai japanese korean vietnamese muslim
halal vegetarian bbq grill cooked hawker centre center market fresh hot cold the and & of noodle house
kitchen eating house restaurant food court stall no tea snacks snack bread toast wraps wrap sandwich
sandwiches burger burgers pasta pizza steak salad congee claypot curry chilli handmade traditional
famous best original old new authentic delights delight boneless herbal""".split())
SKIP_HOSTS = ("instagram.com", "facebook.com", "fb.com", "linktr.ee", "tiktok.com", "google.com", "goo.gl",
              "grab.com", "food.grab.com", "foodpanda.sg", "deliveroo.com.sg", "wa.me", "whatsapp.com",
              "beacons.ai", "linkin.bio", "shopee.sg", "chope.co", "quandoo.sg", "youtube.com", "twitter.com",
              "x.com", "linkedin.com", "tripadvisor.com", "tripadvisor.com.sg", "burpple.com", "yelp.com",
              "maps.app.goo.gl", "bit.ly", "taplink.cc", "carrd.co", "wixsite.com", "sites.google.com")
FOOD_DESC = re.compile(r"restaurant|fast[- ]food|\bfood|caf[eé]|coffee|\btea\b|bak(e|ery)|beverage|drink|"
                       r"juice|dessert|pizza|burger|noodle|convenience store|kopitiam|hawker|confection|snack|"
                       r"ice cream|bubble tea|eatery|dining|chicken|sushi|ramen|donut|doughnut|bread|cookie|"
                       r"catering|culinary|food court|coffeehouse", re.I)


def norm(name: str) -> str:
    n = name.lower().replace("’", "'").replace("‘", "'")
    n = re.sub(r"\s*[\(\[\{].*?[\)\]\}]", "", n)
    n = re.split(r"\s+(?:@|-|–|—|\||at)\s+|\s*@\s*", n)[0]
    n = re.sub(r"[^a-z0-9&'\- ]+", " ", n)
    n = re.sub(r"\b(pte|ltd|private|limited|singapore|sg|s'pore|outlet|branch|food holdings|holdings|"
               r"investment|investments|group|enterprise|enterprises|trading|management)\b", "", n)
    n = re.sub(r"\s+", " ", n).strip(" -'")
    return MERGE.get(n, n)


def is_generic(key: str) -> bool:
    toks = re.findall(r"[a-z0-9&']+", key)
    return not toks or all(t in GENERIC or t.isdigit() for t in toks)


def host_of(url: str) -> str:
    if "://" not in url:
        url = "http://" + url
    h = urllib.parse.urlparse(url.strip()).netloc.lower().split("@")[-1].split(":")[0]
    return h[4:] if h.startswith("www.") else h


def http_json(url, params, ua=WIKI_UA, retries=5):
    q = url + "?" + urllib.parse.urlencode(params)
    for i in range(retries):
        try:
            req = urllib.request.Request(q, headers={"User-Agent": ua})
            with urllib.request.urlopen(req, timeout=20) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if e.code == 429 and i < retries - 1:  # honour Wikimedia rate limiting
                time.sleep(int(e.headers.get("Retry-After") or 0) or 30 * (i + 1))
                continue
            print(f"WARN {url} {params.get('search') or params.get('ids') or params.get('titles')}: {e}",
                  file=sys.stderr)
            return None
        except Exception as e:  # noqa: BLE001 - network errors are retried then reported
            if i == retries - 1:
                print(f"WARN {url} {params.get('search') or params.get('ids') or params.get('titles')}: {e}",
                      file=sys.stderr)
                return None
            time.sleep(2 * (i + 1))


# ---------------- a) Wikidata / Commons ----------------
WD = "https://www.wikidata.org/w/api.php"
COMMONS = "https://commons.wikimedia.org/w/api.php"


def label_norm(s):
    return re.sub(r"[^a-z0-9&]+", "", s.lower().replace("’", "'"))


def wd_search(brand_key):
    d = http_json(WD, {"action": "wbsearchentities", "search": brand_key, "language": "en",
                       "type": "item", "limit": 7, "format": "json", "maxlag": 5})
    time.sleep(1.0)
    return None if d is None else d.get("search", [])


SPARQL = "https://query.wikidata.org/sparql"


def display_variants(key, vs):
    """Case variants of a brand name for exact-label SPARQL lookup."""
    out = {key, key.title(), key.upper(), key.capitalize()}
    for nm, _ in Counter(v["name"] for v in vs).most_common(3):
        base = re.split(r"\s+(?:@|-|–|—|\||at)\s+|\s*@\s*", re.sub(r"\s*[\(\[].*?[\)\]]", "", nm))[0].strip()
        out.update({base, nm.strip()})
    return {o for o in out if o and '"' not in o and "\\" not in o}


def wd_sparql_match(brands, cache_path):
    """Return {brand_key: [ {id,label,description,aliases} ]} for items whose en label/alias equals a variant."""
    try:
        cache = json.load(open(cache_path))
    except Exception:  # noqa: BLE001
        cache = {}
    var_to_keys = defaultdict(set)
    for k, vs in brands.items():
        for v in display_variants(k, vs):
            var_to_keys[v].add(k)
    todo = [v for v in sorted(var_to_keys) if v not in cache]
    for i in range(0, len(todo), 400):
        chunk = todo[i:i + 400]
        values = " ".join(json.dumps(v, ensure_ascii=False) + "@en" for v in chunk)
        q = ("SELECT ?l ?item ?desc WHERE { VALUES ?l { " + values + " } "
             "{ ?item rdfs:label ?l } UNION { ?item skos:altLabel ?l } "
             "?item wdt:P154 ?logo . "
             "OPTIONAL { ?item schema:description ?desc FILTER(lang(?desc)='en') } }")
        d = None
        for attempt in range(4):
            try:
                req = urllib.request.Request(SPARQL, data=urllib.parse.urlencode({"query": q, "format": "json"}).encode(),
                                             headers={"User-Agent": WIKI_UA, "Accept": "application/sparql-results+json"})
                with urllib.request.urlopen(req, timeout=60) as r:
                    d = json.load(r)
                break
            except Exception as e:  # noqa: BLE001
                print("WARN sparql", e, file=sys.stderr)
                time.sleep(20 * (attempt + 1))
        if d is None:
            continue
        for v in chunk:
            cache[v] = []
        for b in d["results"]["bindings"]:
            cache[b["l"]["value"]].append({"id": b["item"]["value"].rsplit("/", 1)[-1], "label": b["l"]["value"],
                                           "description": b.get("desc", {}).get("value", ""), "aliases": []})
        json.dump(cache, open(cache_path, "w"))
        time.sleep(2)
    res = defaultdict(list)
    for v, ks in var_to_keys.items():
        for k in ks:
            for r in cache.get(v, []):
                if r["id"] not in {x["id"] for x in res[k]}:
                    res[k].append(r)
    return res


def wd_entities(qids):
    out = {}
    for i in range(0, len(qids), 50):
        d = http_json(WD, {"action": "wbgetentities", "ids": "|".join(qids[i:i + 50]),
                           "props": "claims|labels|descriptions", "languages": "en", "format": "json"})
        out.update((d or {}).get("entities", {}))
        time.sleep(1.0)
    return out


def claim_vals(ent, prop):
    vals = []
    for c in ent.get("claims", {}).get(prop, []):
        dv = c.get("mainsnak", {}).get("datavalue", {}).get("value")
        if isinstance(dv, dict) and "id" in dv:
            vals.append(dv["id"])
        elif dv is not None:
            vals.append(dv)
    return vals


def strip_html(s):
    return html.unescape(re.sub(r"<[^>]+>", "", s or "")).strip()


def commons_info(files):
    out = {}
    for i in range(0, len(files), 40):
        titles = "|".join("File:" + f for f in files[i:i + 40])
        d = http_json(COMMONS, {"action": "query", "titles": titles, "prop": "imageinfo",
                                "iiprop": "url|size|extmetadata", "format": "json"})
        time.sleep(1.0)
        pages = (d or {}).get("query", {}).get("pages", {})
        norm_map = {n["to"]: n["from"] for n in (d or {}).get("query", {}).get("normalized", [])}
        for p in pages.values():
            ii = (p.get("imageinfo") or [{}])[0]
            em = ii.get("extmetadata", {})
            title = norm_map.get(p.get("title"), p.get("title", ""))
            out[title.removeprefix("File:")] = {
                "url": ii.get("url", "").split("?")[0], "page": ii.get("descriptionurl", ""),
                "width": ii.get("width"), "height": ii.get("height"),
                "licence": strip_html(em.get("LicenseShortName", {}).get("value", "")),
                "artist": strip_html(em.get("Artist", {}).get("value", "")),
                "attribution_required": strip_html(em.get("AttributionRequired", {}).get("value", "")),
                "restrictions": strip_html(em.get("Restrictions", {}).get("value", "")),
            }
    return out


def commons_rights(info):
    lic = info.get("licence", "").lower()
    restr = info.get("restrictions", "").lower()
    free = bool(re.search(r"\bcc[ -]?(by|0|zero)|public domain|\bpd\b|pd-", lic))
    if not free or "trademark" in restr or "textlogo" in lic or "ineligible" in lic:
        return "trademark_review"
    return "cleared_open_licence"


# ---------------- b) Official website ----------------
class HeadParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links, self.metas, self.jsonld, self._in_ld, self._buf = [], [], [], False, []

    def handle_starttag(self, tag, attrs):
        a = {k.lower(): (v or "") for k, v in attrs}
        if tag == "link" and a.get("href"):
            self.links.append((a.get("rel", "").lower(), a["href"], a.get("sizes", "")))
        elif tag == "meta" and a.get("content"):
            self.metas.append(((a.get("property") or a.get("name") or a.get("itemprop") or "").lower(), a["content"]))
        elif tag == "script" and "ld+json" in a.get("type", "").lower():
            self._in_ld, self._buf = True, []

    def handle_data(self, data):
        if self._in_ld:
            self._buf.append(data)

    def handle_endtag(self, tag):
        if tag == "script" and self._in_ld:
            self._in_ld = False
            self.jsonld.append("".join(self._buf))


def jsonld_logos(blob):
    found = []

    def walk(o):
        if isinstance(o, dict):
            t = o.get("@type")
            types = t if isinstance(t, list) else [t]
            if "logo" in o and any(x in ("Organization", "Restaurant", "LocalBusiness", "Brand",
                                          "FoodEstablishment", "Corporation", "CafeOrCoffeeShop", "Store")
                                   for x in types):
                lg = o["logo"]
                if isinstance(lg, dict):
                    lg = lg.get("url") or lg.get("contentUrl")
                if isinstance(lg, str):
                    found.append(lg)
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
    try:
        walk(json.loads(blob))
    except Exception:  # noqa: BLE001 - malformed JSON-LD is common; ignore
        pass
    return found


def fetch_site(host, start_url):
    """robots.txt check, then a single homepage GET. Returns (status, candidates)."""
    base = start_url if "://" in start_url else "https://" + start_url
    parsed = urllib.parse.urlparse(base)
    scheme = parsed.scheme or "https"
    root = f"{scheme}://{parsed.netloc}"
    home = root + "/"
    rp = urllib.robotparser.RobotFileParser()
    try:
        req = urllib.request.Request(root + "/robots.txt", headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=10) as r:
            rp.parse(r.read(200_000).decode("utf-8", "replace").splitlines())
    except urllib.error.HTTPError as e:
        rp.parse([] if e.code >= 400 and e.code != 401 and e.code != 403 else [])
        if e.code in (401, 403):
            return "robots_forbidden", []
    except Exception:  # noqa: BLE001 - no robots.txt reachable -> treat as allowed
        rp.parse([])
    if not rp.can_fetch(UA, home):
        return "disallowed_by_robots", []
    time.sleep(1.0)  # 1 req/s per host
    try:
        req = urllib.request.Request(home, headers={"User-Agent": UA, "Accept": "text/html"})
        with urllib.request.urlopen(req, timeout=10) as r:
            final = r.geturl()
            if "html" not in (r.headers.get("Content-Type") or "html"):
                return "not_html", []
            body = r.read(2_000_000).decode(r.headers.get_content_charset() or "utf-8", "replace")
    except Exception as e:  # noqa: BLE001
        return f"fetch_error:{type(e).__name__}", []
    p = HeadParser()
    try:
        p.feed(body)
    except Exception:  # noqa: BLE001
        pass
    cands = []
    j = lambda u: urllib.parse.urljoin(final, u.strip())  # noqa: E731
    for blob in p.jsonld:
        for u in jsonld_logos(blob):
            cands.append(("jsonld", j(u), ""))
    for k, v in p.metas:
        if k in ("og:logo", "logo"):
            cands.append(("og_logo", j(v), ""))
        elif k in ("og:image", "og:image:url", "og:image:secure_url"):
            cands.append(("og_image", j(v), ""))
    for rel, href, sizes in p.links:
        if "apple-touch-icon" in rel:
            cands.append(("official_site_favicon", j(href), sizes or "apple-touch-icon"))
        elif "icon" in rel.split():
            cands.append(("official_site_favicon", j(href), sizes))
    seen, out = set(), []
    for c in cands:
        if c[1] not in seen and not c[1].startswith("data:"):
            seen.add(c[1])
            out.append(c)
    return "ok", out


def best_per_source(cands):
    """Keep one candidate per source type; for favicons prefer the largest declared size."""
    by = defaultdict(list)
    for c in cands:
        by[c[0]].append(c)

    def size_of(c):
        m = re.search(r"(\d+)x(\d+)", c[2] or "")
        return int(m.group(1)) if m else (180 if "apple" in (c[2] or "") else 0)
    out = []
    for src, lst in by.items():
        out.append(max(lst, key=size_of) if src == "official_site_favicon" else lst[0])
    return out


# ---------------- main ----------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--venues", required=True)
    ap.add_argument("--overpass")
    ap.add_argument("--out", required=True)
    ap.add_argument("--cache", required=True, help="scratch dir for API response caches")
    ap.add_argument("--workers", type=int, default=6)
    args = ap.parse_args()

    venues = json.load(open(args.venues))
    groups = defaultdict(list)
    for v in venues:
        groups[norm(v["name"])].append(v)

    # OSM brand -> QID (strong evidence for matching)
    osm_qid = {}
    if args.overpass:
        try:
            el = json.load(open(args.overpass)).get("elements", [])
            cnt = Counter()
            for e in el:
                t = e.get("tags", {})
                if t.get("brand:wikidata") and (t.get("brand") or t.get("name")):
                    for nm in {t.get("brand"), t.get("brand:en")} - {None}:  # not name: avoids mis-tagged outlets
                        cnt[(norm(nm), t["brand:wikidata"].split(";")[0])] += 1
            for (k, q), n in cnt.most_common():
                osm_qid.setdefault(k, q)
        except Exception as e:  # noqa: BLE001
            print("WARN overpass cache unreadable:", e, file=sys.stderr)

    brands = {k: vs for k, vs in groups.items()
              if k and (len(vs) >= CHAIN_MIN or k in KNOWN_BRANDS or k in osm_qid) and not is_generic(k)}
    brands = {k: vs for k, vs in brands.items() if len(vs) >= CHAIN_MIN or k in KNOWN_BRANDS}
    generic_groups = sum(1 for k, vs in groups.items() if len(vs) >= CHAIN_MIN and is_generic(k))
    print(f"brands={len(brands)} generic_3plus_groups_skipped={generic_groups} osm_qids={len(osm_qid)}",
          file=sys.stderr)

    # --- Wikidata matching: batched SPARQL exact label/alias lookup (cached) ---
    sc = wd_sparql_match({k: vs for k, vs in brands.items() if k not in osm_qid}, f"{args.cache}/wd_sparql_cache.json")
    # Fallback: wbsearchentities (fuzzier, rate-limited) only for bigger chains SPARQL missed
    ws_path = f"{args.cache}/wd_search_cache.json"
    try:
        wsc = json.load(open(ws_path))
    except Exception:  # noqa: BLE001
        wsc = {}
    for k, vs in brands.items():
        if k in osm_qid or sc.get(k) or len(vs) < 10:
            continue
        if k not in wsc:
            res = wd_search(k)
            if res is None:
                continue
            wsc[k] = res
            json.dump(wsc, open(ws_path, "w"))
        sc[k] = wsc[k]

    cand_q = set(osm_qid.get(k) for k in brands if k in osm_qid)
    for k in brands:
        for r in sc.get(k, []):
            labels = [r.get("label", "")] + ([r["match"]["text"]] if r.get("match") else []) + r.get("aliases", [])
            if any(label_norm(x) == label_norm(k) for x in labels):
                cand_q.add(r["id"])
    ents = wd_entities(sorted(q for q in cand_q if q))

    brand_q = {}
    for k in brands:
        if k in osm_qid:
            brand_q[k] = (osm_qid[k], "osm_brand_wikidata")
            continue
        for r in sc.get(k, []):
            labels = [r.get("label", "")] + ([r["match"]["text"]] if r.get("match") else []) + r.get("aliases", [])
            if not any(label_norm(x) == label_norm(k) for x in labels):
                continue
            ent = ents.get(r["id"], {})
            desc = (ent.get("descriptions", {}).get("en", {}) or {}).get("value", "") or r.get("description", "")
            # require a food/business description and a logo claim to avoid same-name false positives
            if FOOD_DESC.search(desc) and claim_vals(ent, "P154"):
                brand_q[k] = (r["id"], "wikidata_search_exact_label")
                break

    logo_file = {}
    for k, (q, _) in brand_q.items():
        vals = claim_vals(ents.get(q, {}), "P154")
        if vals:
            logo_file[k] = vals[-1]  # last listed is usually the current logo
    cinfo = commons_info(sorted(set(logo_file.values())))

    rows = []
    covered = set()
    for k, vs in sorted(brands.items(), key=lambda kv: -len(kv[1])):
        disp = Counter(v["name"] for v in vs).most_common(1)[0][0]
        ids = ";".join(v["id"] for v in vs)
        if k in logo_file and logo_file[k] in cinfo and cinfo[logo_file[k]]["url"]:
            ci = cinfo[logo_file[k]]
            q, how = brand_q[k]
            attrib = ci["artist"] or ""
            if ci["attribution_required"].lower() == "true":
                attrib = (attrib + " (attribution required)").strip()
            rows.append({"venue_ids_or_brand": ids, "brand_key": k, "brand_name": disp, "outlet_count": len(vs),
                         "logo_url": ci["url"], "source": "wikimedia_commons", "licence": ci["licence"],
                         "attribution": attrib, "rights_status": commons_rights(ci),
                         "width": ci["width"], "height": ci["height"], "wikidata_id": q, "match_method": how,
                         "confidence": "high" if how == "osm_brand_wikidata" else "med",
                         "source_page": ci["page"], "notes": ("Restrictions: " + ci["restrictions"]) if ci["restrictions"] else "",
                         "owner_approved": ""})
            covered.add(k)

    # --- Websites ---
    host_venues = defaultdict(list)
    host_url = {}
    for v in venues:
        u = (v.get("website_url") or "").strip()
        if not u:
            continue
        h = host_of(u)
        if not h:
            continue
        host_venues[h].append(v)
        host_url.setdefault(h, u)
    site_cache_path = f"{args.cache}/site_cache.json"
    try:
        site_cache = json.load(open(site_cache_path))
    except Exception:  # noqa: BLE001
        site_cache = {}
    todo = [h for h in host_venues if h not in site_cache and not any(h == s or h.endswith("." + s) for s in SKIP_HOSTS)]
    lock = threading.Lock()
    done = [0]

    def work(h):
        st, c = fetch_site(h, host_url[h])
        with lock:
            site_cache[h] = {"status": st, "cands": c}
            done[0] += 1
            if done[0] % 50 == 0:
                json.dump(site_cache, open(site_cache_path, "w"))
                print(f"sites {done[0]}/{len(todo)}", file=sys.stderr)
    with ThreadPoolExecutor(args.workers) as ex:
        list(ex.map(work, todo))
    json.dump(site_cache, open(site_cache_path, "w"))

    site_status = Counter()
    hosts_with_cands = 0
    for h, vs in sorted(host_venues.items(), key=lambda kv: -len(kv[1])):
        if any(h == s or h.endswith("." + s) for s in SKIP_HOSTS):
            site_status["skipped_social_or_aggregator"] += 1
            continue
        sc_ = site_cache.get(h, {"status": "not_fetched", "cands": []})
        site_status[sc_["status"].split(":")[0]] += 1
        cands = best_per_source([tuple(c) for c in sc_["cands"]])
        if cands:
            hosts_with_cands += 1
        disp = Counter(v["name"] for v in vs).most_common(1)[0][0]
        for src, url, sizes in cands:
            m = re.search(r"(\d+)x(\d+)", sizes or "")
            rows.append({"venue_ids_or_brand": ";".join(v["id"] for v in vs), "brand_key": norm(disp),
                         "brand_name": disp, "outlet_count": len(vs), "logo_url": url,
                         "source": "og_image" if src in ("og_image", "og_logo") else src,
                         "licence": "unknown (merchant-owned; not licensed)", "attribution": h,
                         "rights_status": "needs_merchant_permission",
                         "width": int(m.group(1)) if m else "", "height": int(m.group(2)) if m else "",
                         "wikidata_id": "", "match_method": "official_website_homepage",
                         "confidence": ("low" if src == "og_image" else
                                        "high" if len({norm(v["name"]) for v in vs}) <= 2 else "med"),
                         "source_page": f"https://{h}/", "notes": ("og:logo" if src == "og_logo" else ""),
                         "owner_approved": ""})
            covered.add(norm(disp))

    # Chains with no candidate -> merchant upload
    for k, vs in sorted(brands.items(), key=lambda kv: -len(kv[1])):
        if k in covered:
            continue
        disp = Counter(v["name"] for v in vs).most_common(1)[0][0]
        rows.append({"venue_ids_or_brand": ";".join(v["id"] for v in vs), "brand_key": k, "brand_name": disp,
                     "outlet_count": len(vs), "logo_url": "", "source": "none", "licence": "", "attribution": "",
                     "rights_status": "merchant_upload_required", "width": "", "height": "",
                     "wikidata_id": brand_q.get(k, ("", ""))[0], "match_method": "", "confidence": "",
                     "source_page": "", "notes": "No licensable logo found; request via merchant portal upload.",
                     "owner_approved": ""})

    cols = ["venue_ids_or_brand", "brand_name", "brand_key", "outlet_count", "logo_url", "source", "licence",
            "attribution", "rights_status", "confidence", "width", "height", "wikidata_id", "match_method", "source_page",
            "notes", "owner_approved"]
    for r in rows:  # Excel cell limit
        if len(r["venue_ids_or_brand"]) > 32000:
            r["venue_ids_or_brand"] = r["venue_ids_or_brand"][:32000] + ";...TRUNCATED(see csv)"
    with open(f"{args.out}/stall_logo_candidates.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    commons_rows = [r for r in rows if r["source"] == "wikimedia_commons"]
    venues_with_site = sum(len(v) for v in host_venues.values())
    summary = [
        ("Total venues", len(venues)),
        ("Brand groups (3+ outlets or known brand, generic names excluded)", len(brands)),
        ("Generic 3+ name groups skipped (e.g. 'chicken rice', 'drink stall')", generic_groups),
        ("Brands matched to a Wikidata item", len(brand_q)),
        ("  via OSM brand:wikidata", sum(1 for _, h in brand_q.values() if h == "osm_brand_wikidata")),
        ("  via Wikidata exact label/alias (SPARQL or wbsearchentities) + food description + P154", sum(1 for _, h in brand_q.values() if h != "osm_brand_wikidata")),
        ("Brands with a Commons logo (P154)", len(commons_rows)),
        ("  venues covered by Commons logos", sum(r["outlet_count"] for r in commons_rows)),
    ]
    for lic, n in Counter(r["licence"] for r in commons_rows).most_common():
        summary.append((f"  Commons licence: {lic}", n))
    for st, n in Counter(r["rights_status"] for r in commons_rows).most_common():
        summary.append((f"  Commons rights_status: {st}", n))
    summary += [
        ("Venues with website_url", venues_with_site),
        ("Distinct website hosts", len(host_venues)),
    ]
    for st, n in Counter(r["rights_status"] for r in rows).most_common():
        summary.append((f"All rows rights_status: {st}", n))
    for st, n in Counter(r["confidence"] or "n/a" for r in rows).most_common():
        summary.append((f"All rows confidence: {st}", n))
    for st, n in site_status.most_common():
        summary.append((f"  host status: {st}", n))
    web_rows = [r for r in rows if r["rights_status"] == "needs_merchant_permission"]
    summary += [
        ("Hosts with >=1 website logo candidate", hosts_with_cands),
        ("Website candidate rows", len(web_rows)),
    ]
    for s, n in Counter(r["source"] for r in web_rows).most_common():
        summary.append((f"  website source: {s}", n))
    summary.append(("Brand groups with no candidate (merchant upload)",
                    sum(1 for r in rows if r["source"] == "none")))
    covered_ids = set()
    for r in rows:
        if r["source"] != "none":
            covered_ids.update(r["venue_ids_or_brand"].split(";"))
    summary.append(("Venues with no candidate at all (merchant upload)", len(venues) - len(covered_ids & {v['id'] for v in venues})))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = Workbook()
        ws = wb.active
        ws.title = "Candidates"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        fills = {"cleared_open_licence": "E2F0D9", "trademark_review": "FFF2CC",
                 "needs_merchant_permission": "FCE4D6", "merchant_upload_required": "EDEDED"}
        for r in rows:
            ws.append([r[c] for c in cols])
            f = fills.get(r["rights_status"])
            if f:
                ws.cell(ws.max_row, cols.index("rights_status") + 1).fill = PatternFill("solid", fgColor=f)
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        ws.add_data_validation(dv)
        col_oa = "ABCDEFGHIJKLMNOPQ"[cols.index("owner_approved")]
        dv.add(f"{col_oa}2:{col_oa}{max(ws.max_row, 2)}")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, wdt in zip("ABCDEFGHIJKLMNOPQ", (30, 30, 22, 9, 60, 20, 26, 30, 24, 10, 7, 7, 12, 24, 40, 40, 10)):
            ws.column_dimensions[col].width = wdt
        s = wb.create_sheet("Summary")
        s.append(["metric", "value"])
        for k, v in summary:
            s.append([k, v])
        s.column_dimensions["A"].width = 70
        s.append([])
        s.append(["Rights note", "Logos are trademarks. 'cleared_open_licence' means only the Commons file "
                  "copyright licence is free; trademark use still needs brand consent for anything implying "
                  "endorsement. Website URLs are merchant-owned: get permission before display."])
        wb.save(f"{args.out}/stall_logo_candidates.xlsx")
    except ImportError:
        print("openpyxl missing; CSV only", file=sys.stderr)
    json.dump(summary, open(f"{args.cache}/logo_summary.json", "w"), indent=1)
    for k, v in summary:
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
