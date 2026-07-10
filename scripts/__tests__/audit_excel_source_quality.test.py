#!/usr/bin/env python3
"""Unit tests for the Hungry Excel source-quality display gate."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from audit_excel_source_quality import (  # noqa: E402
    audit_supabase_cache,
    classify_row,
    unapproved_display_rows,
)


def row(**overrides):
    value = {
        "Menu Source": "official_online_seed:https://merchant.example/menu",
        "Menu Items": "Chicken Rice, Laksa, Nasi Lemak, Mee Soto, Satay, Roti Prata, Kaya Toast, Char Kway Teow",
        "Image URL": "https://merchant.example/stall.jpg",
        "Dishes Image URL": "https://merchant.example/dish.jpg",
        "Image Rights Status": "merchant_owned",
        "Menu Rights Status": "merchant_owned",
        "Approved for Display": "yes",
        "Source URL": "https://merchant.example",
        "Google Website": "",
        "Google Place ID": "",
        "Google Rating": "",
        "Google Review Count": "",
    }
    value.update(overrides)
    return value


class AuditExcelSourceQualityTest(unittest.TestCase):
    def test_licensed_or_merchant_owned_row_is_display_approved(self):
        flags, _status = classify_row(row())

        self.assertIn("menu_display_approved", flags)
        self.assertIn("image_display_approved", flags)
        self.assertNotIn("production_menu_unverified", flags)
        self.assertNotIn("production_image_unverified", flags)
        self.assertEqual([], unapproved_display_rows([{"QA Flags": ";".join(flags)}]))

    def test_delivery_app_rows_are_quarantined_even_when_marked_approved(self):
        flags, _status = classify_row(
            row(
                **{
                    "Menu Source": "official_online_seed:https://foodpanda.sg/restaurant/test",
                    "Source URL": "https://foodpanda.sg/restaurant/test",
                    "Image Rights Status": "unverified",
                    "Menu Rights Status": "unverified",
                },
            ),
        )

        self.assertIn("delivery_app_source_quarantine", flags)
        self.assertIn("production_menu_unverified", flags)
        self.assertIn("production_image_unverified", flags)
        self.assertEqual(1, len(unapproved_display_rows([{"QA Flags": ";".join(flags)}])))

    def test_official_source_still_needs_explicit_approval_and_rights_status(self):
        flags, _status = classify_row(
            row(
                **{
                    "Approved for Display": "",
                    "Menu Rights Status": "",
                    "Image Rights Status": "",
                },
            ),
        )

        self.assertIn("menu_operator_source", flags)
        self.assertNotIn("menu_display_approved", flags)
        self.assertNotIn("image_display_approved", flags)
        self.assertEqual(1, len(unapproved_display_rows([{"QA Flags": ";".join(flags)}])))

    def test_placeholder_images_are_not_display_approved(self):
        flags, _status = classify_row(
            row(
                **{
                    "Image URL": "https://images.unsplash.com/photo-1517248135467-4c7edcad34c4",
                    "Dishes Image URL": "https://images.unsplash.com/photo-1546069901-ba9599a7e63c",
                    "Image Rights Status": "merchant_owned",
                },
            ),
        )

        self.assertIn("storefront_image_placeholder", flags)
        self.assertIn("dish_image_placeholder", flags)
        self.assertIn("production_image_unverified", flags)
        self.assertEqual(1, len(unapproved_display_rows([{"QA Flags": ";".join(flags)}])))

    def test_cache_audit_uses_supabase_approval_and_rights_columns(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "audit.xlsx"
            workbook = Workbook()
            workbook.active.title = "Sheet1"
            workbook["Sheet1"].append(["Business ID", "Stall / Restaurant Name"])
            workbook["Sheet1"].append(["business-1", "Safe Venue"])
            venues = workbook.create_sheet("Supabase Venues")
            venues.append(
                [
                    "id",
                    "menu_approved_for_display",
                    "menu_rights_status",
                    "media_approved_for_display",
                    "media_rights_status",
                    "image_url",
                    "logo_url",
                ]
            )
            venues.append(
                [
                    "venue-1",
                    True,
                    "operator_owned",
                    True,
                    "permitted_api",
                    "https://merchant.example/venue.jpg",
                    "",
                ]
            )
            dishes = workbook.create_sheet("Supabase Dishes")
            dishes.append(
                [
                    "id",
                    "menu_approved_for_display",
                    "menu_rights_status",
                    "media_approved_for_display",
                    "media_rights_status",
                    "image_url",
                ]
            )
            dishes.append(
                [
                    "dish-1",
                    True,
                    "licensed",
                    True,
                    "unverified",
                    "https://merchant.example/dish.jpg",
                ]
            )
            workbook.save(path)
            workbook.close()

            result = audit_supabase_cache(path)

        self.assertEqual(result["rows_checked"], 2)
        self.assertFalse(result["passed"])
        self.assertEqual(result["flag_counts"]["media_approval_rights_conflict"], 1)
        self.assertEqual(result["flag_counts"]["media_not_authorized"], 1)
        self.assertNotIn("menu_not_authorized", result["flag_counts"])


if __name__ == "__main__":
    unittest.main()
