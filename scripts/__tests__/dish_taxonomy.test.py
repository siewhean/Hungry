#!/usr/bin/env python3

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from dish_taxonomy import DISH_TAXONOMY, lookup_dish_taxonomy  # noqa: E402
from export_excel_catalog_for_app import (  # noqa: E402
    build_catalog,
    infer_food_type_tag,
    infer_taste_tags,
    infer_taste_vector,
    load_sheet_rows,
)


class DishTaxonomyTest(unittest.TestCase):
    def test_contains_at_least_fifty_structured_dish_types(self):
        self.assertGreaterEqual(len(DISH_TAXONOMY), 50)
        for entry in DISH_TAXONOMY:
            self.assertTrue(entry.cuisine_category)
            self.assertGreaterEqual(len(entry.taste_tags), 2)
            self.assertEqual(
                {"sweet", "spicy", "sour", "salty", "rich", "aromatic"},
                set(entry.taste_vector),
            )

    def test_longest_alias_wins_for_specific_common_dish(self):
        entry = lookup_dish_taxonomy("Signature Roasted Hainanese Chicken Rice Set")
        self.assertIsNotNone(entry)
        self.assertEqual("Hainanese Chicken Rice", entry.canonical_name)
        self.assertEqual("Singaporean", entry.cuisine_category)

    def test_exporter_uses_taxonomy_before_generic_cuisine_fallback(self):
        self.assertEqual("Korean", infer_food_type_tag(["Asian"], "Beef Bibimbap"))
        self.assertIn("Aromatic", infer_taste_tags("Asian", "Beef Bibimbap"))
        self.assertEqual(8, infer_taste_vector("Asian", "Beef Bibimbap")["aromatic"])

    def test_unknown_dish_preserves_existing_exporter_fallback(self):
        self.assertEqual("Mediterranean", infer_food_type_tag(["Mediterranean"], "House Special"))
        self.assertEqual(["Local"], infer_taste_tags("", "House Special"))

    def test_explicit_source_food_type_remains_authoritative(self):
        self.assertEqual(
            "Merchant Category",
            infer_food_type_tag(["Asian"], "Bibimbap", explicit="Merchant Category"),
        )

    def test_catalog_build_uses_taxonomy_for_workbook_dish(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(
                [
                    "Business ID",
                    "Stall / Restaurant Name",
                    "Cuisine",
                    "Menu Items",
                    "Price Min (SGD)",
                    "Price Max (SGD)",
                ]
            )
            sheet.append(["SFA-TEST", "Test Stall", "Asian", "Beef Bibimbap", 12, 12])
            venues = workbook.create_sheet("Supabase Venues")
            venues.append(
                [
                    "id",
                    "business_id",
                    "name",
                    "menu_approved_for_display",
                    "menu_rights_status",
                    "media_approved_for_display",
                    "media_rights_status",
                ]
            )
            venues.append(
                ["venue-test", "SFA-TEST", "Test Stall", True, "operator_owned", False, "unverified"]
            )
            dishes = workbook.create_sheet("Supabase Dishes")
            dishes.append(
                [
                    "id",
                    "venue_id",
                    "name",
                    "menu_approved_for_display",
                    "menu_rights_status",
                    "media_approved_for_display",
                    "media_rights_status",
                ]
            )
            dishes.append(
                ["dish-test", "venue-test", "Beef Bibimbap", True, "licensed", False, "unverified"]
            )
            workbook.save(workbook_path)

            catalog = build_catalog(load_sheet_rows(workbook_path), workbook_path, 0)
            columns = {column: index for index, column in enumerate(catalog["dishColumns"])}
            dish = catalog["dishes"][0]
            self.assertEqual("Korean", dish[columns["foodTypeTag"]])
            self.assertIn("Aromatic", dish[columns["tasteTags"]])
            self.assertEqual(8, dish[columns["tasteVector"]]["aromatic"])


if __name__ == "__main__":
    unittest.main()
