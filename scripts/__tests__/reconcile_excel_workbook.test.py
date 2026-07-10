#!/usr/bin/env python3
"""Tests for deterministic workbook-only reconciliation."""

from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from export_excel_catalog_for_app import build_catalog, load_sheet_rows  # noqa: E402
from reconcile_excel_workbook import ReconciliationError, reconcile_workbook  # noqa: E402


SOURCE_HEADERS = [
    "Business ID",
    "Stall ID",
    "Licence No",
    "Stall / Restaurant Name",
    "Postal Code",
    "Unit No.",
    "Menu Items",
    "Last Verified Date",
    "Image URL",
    "Dishes Image URL",
    "Logo URL",
    "Latitude",
    "Longitude",
]
VENUE_HEADERS = [
    "id",
    "business_id",
    "stall_id",
    "licence_no",
    "name",
    "postal_code",
    "unit_number",
    "menu_approved_for_display",
    "menu_rights_status",
    "media_approved_for_display",
    "media_rights_status",
    "image_url",
    "logo_url",
]
DISH_HEADERS = [
    "id",
    "venue_id",
    "name",
    "description",
    "menu_approved_for_display",
    "menu_rights_status",
    "media_approved_for_display",
    "media_rights_status",
    "image_url",
]


def append_dict(sheet, headers, values):
    sheet.append([values.get(header) for header in headers])


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    source = workbook.active
    source.title = "Sheet1"
    source.append(SOURCE_HEADERS)
    safe_source = {
        "Business ID": "biz-safe",
        "Stall / Restaurant Name": "Safe Stall",
        "Postal Code": "238801",
        "Unit No.": "#01-01",
        "Last Verified Date": "2026-06-01",
        "Image URL": "https://cdn.example/placeholder.jpg",
        "Latitude": 1.3,
        "Longitude": 103.8,
    }
    append_dict(source, SOURCE_HEADERS, safe_source)
    append_dict(source, SOURCE_HEADERS, {**safe_source, "Menu Items": "Chicken Rice ($5.00)"})
    conflict_source = {
        "Business ID": "biz-conflict",
        "Stall / Restaurant Name": "Conflicting A",
        "Postal Code": "111111",
        "Unit No.": "#01-01",
        "Latitude": 1.31,
        "Longitude": 103.81,
    }
    append_dict(source, SOURCE_HEADERS, conflict_source)
    append_dict(
        source,
        SOURCE_HEADERS,
        {**conflict_source, "Stall / Restaurant Name": "Conflicting B"},
    )
    append_dict(
        source,
        SOURCE_HEADERS,
        {"Stall / Restaurant Name": "No Stable Identity", "Latitude": 1.32, "Longitude": 103.82},
    )
    append_dict(
        source,
        SOURCE_HEADERS,
        {
            "Business ID": "shared-business",
            "Stall / Restaurant Name": "Ambiguous Cache",
            "Postal Code": "222222",
            "Unit No.": "#02-02",
            "Latitude": 1.33,
            "Longitude": 103.83,
        },
    )

    venues = workbook.create_sheet("Supabase Venues")
    venues.append(VENUE_HEADERS)
    safe_venue = {
        "id": "venue-safe",
        "business_id": "biz-safe",
        "name": "Safe Stall",
        "postal_code": "238801",
        "unit_number": "#01-01",
        "menu_approved_for_display": True,
        "menu_rights_status": "operator_owned",
        "media_approved_for_display": True,
        "media_rights_status": "unverified",
        "image_url": "https://cdn.example/placeholder.jpg",
    }
    append_dict(venues, VENUE_HEADERS, safe_venue)
    append_dict(venues, VENUE_HEADERS, {**safe_venue, "logo_url": "https://merchant.example/logo.jpg"})
    for venue_id in ("venue-a", "venue-b"):
        append_dict(
            venues,
            VENUE_HEADERS,
            {
                "id": venue_id,
                "business_id": "shared-business",
                "name": "Ambiguous Cache",
                "menu_approved_for_display": True,
                "menu_rights_status": "licensed",
            },
        )
    conflict_venue = {
        "id": "venue-conflict",
        "business_id": "biz-venue-conflict",
        "name": "Venue Conflict A",
    }
    append_dict(venues, VENUE_HEADERS, conflict_venue)
    append_dict(venues, VENUE_HEADERS, {**conflict_venue, "name": "Venue Conflict B"})
    append_dict(venues, VENUE_HEADERS, {"name": "Missing Venue ID"})

    dishes = workbook.create_sheet("Supabase Dishes")
    dishes.append(DISH_HEADERS)
    safe_dish = {
        "id": "dish-safe",
        "venue_id": "venue-safe",
        "name": "Chicken Rice",
        "menu_approved_for_display": True,
        "menu_rights_status": "unverified",
        "media_approved_for_display": True,
        "media_rights_status": "unverified",
        "image_url": "https://cdn.example/placeholder.jpg",
    }
    append_dict(dishes, DISH_HEADERS, safe_dish)
    append_dict(dishes, DISH_HEADERS, {**safe_dish, "description": "Merchant description"})
    append_dict(
        dishes,
        DISH_HEADERS,
        {"id": "dish-orphan", "venue_id": "missing", "name": "Orphan"},
    )
    conflict_dish = {"id": "dish-conflict", "venue_id": "venue-safe", "name": "Dish A"}
    append_dict(dishes, DISH_HEADERS, conflict_dish)
    append_dict(dishes, DISH_HEADERS, {**conflict_dish, "name": "Dish B"})
    append_dict(
        dishes,
        DISH_HEADERS,
        {"id": "dish-removed-venue", "venue_id": "venue-conflict", "name": "Removed Venue Dish"},
    )
    workbook.save(path)
    workbook.close()


def rows_by_header(path: Path, sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        values = sheet.iter_rows(values_only=True)
        headers = list(next(values))
        return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]
    finally:
        workbook.close()


def workbook_values(path: Path) -> dict[str, list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            sheet_name: list(workbook[sheet_name].iter_rows(values_only=True))
            for sheet_name in workbook.sheetnames
        }
    finally:
        workbook.close()


class ReconcileExcelWorkbookTest(unittest.TestCase):
    def test_reconciles_without_overwriting_or_inventing_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            output = root / "candidate.xlsx"
            make_workbook(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

            metrics = reconcile_workbook(source, output)

            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
            self.assertTrue(output.exists())
            source_rows = rows_by_header(output, "Sheet1")
            self.assertEqual([row["Business ID"] for row in source_rows], ["biz-safe"])
            self.assertEqual(source_rows[0]["Menu Items"], "Chicken Rice ($5.00)")
            self.assertEqual(source_rows[0]["Last Verified Date"], "2026-06-01")
            self.assertIsNone(source_rows[0]["Image URL"])

            venues = rows_by_header(output, "Supabase Venues")
            self.assertEqual({row["id"] for row in venues}, {"venue-safe", "venue-a", "venue-b"})
            safe_venue = next(row for row in venues if row["id"] == "venue-safe")
            self.assertFalse(safe_venue["media_approved_for_display"])
            self.assertIsNone(safe_venue["image_url"])
            self.assertEqual(safe_venue["media_rights_status"], "unverified")

            dishes = rows_by_header(output, "Supabase Dishes")
            self.assertEqual([row["id"] for row in dishes], ["dish-safe"])
            self.assertFalse(dishes[0]["menu_approved_for_display"])
            self.assertFalse(dishes[0]["media_approved_for_display"])
            self.assertIsNone(dishes[0]["image_url"])
            self.assertEqual(dishes[0]["description"], "Merchant description")

            source_quarantine = rows_by_header(output, "Quarantine Sheet1")
            venue_quarantine = rows_by_header(output, "Quarantine Venues")
            dish_quarantine = rows_by_header(output, "Quarantine Dishes")
            self.assertEqual(len(source_quarantine), 5)
            self.assertEqual(len(venue_quarantine), 4)
            self.assertEqual(len(dish_quarantine), 5)
            self.assertTrue(any("ambiguous_cache_match" in row["_reason_codes"] for row in source_quarantine))
            self.assertTrue(any("orphan_dish_reference" in row["_reason_codes"] for row in dish_quarantine))

            report = {row["metric"]: row["value"] for row in rows_by_header(output, "Reconciliation Report")}
            self.assertEqual(report["export_preflight"], "passed")
            self.assertNotIn("generated_at", report)
            self.assertEqual(report["merged_source_groups"], 1)
            self.assertEqual(report["orphan_dish_rows"], 2)
            self.assertEqual(metrics["output_sheet1_rows"], 1)

            catalog = build_catalog(load_sheet_rows(output), output, 0)
            self.assertEqual(catalog["metadata"]["venueCount"], 1)
            self.assertEqual(catalog["metadata"]["dishCount"], 1)

    def test_same_input_and_output_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            make_workbook(source)
            with self.assertRaisesRegex(ReconciliationError, "input overwrite is forbidden"):
                reconcile_workbook(source, source)

    def test_reconciliation_sheet_values_are_deterministic(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            first = root / "first.xlsx"
            second = root / "second.xlsx"
            make_workbook(source)

            reconcile_workbook(source, first)
            reconcile_workbook(source, second)

            self.assertEqual(workbook_values(first), workbook_values(second))

    def test_existing_quarantine_rows_survive_a_followup_reconciliation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            first = root / "first.xlsx"
            followup = root / "followup.xlsx"
            make_workbook(source)
            reconcile_workbook(source, first)

            expected = {
                sheet: rows_by_header(first, sheet)
                for sheet in ("Quarantine Sheet1", "Quarantine Venues", "Quarantine Dishes")
            }
            reconcile_workbook(first, followup)

            for sheet_name, rows in expected.items():
                self.assertEqual(rows_by_header(followup, sheet_name), rows)


if __name__ == "__main__":
    unittest.main()
