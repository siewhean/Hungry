#!/usr/bin/env python3
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "Hungry F&B Database.xlsx"
ENRICHMENT_RESULTS_PATH = ROOT / "data" / "menu_enrichment_results.json"
SHEET_TAB = "Sheet1"

def clean(val):
    if val is None:
        return ""
    return str(val).strip()

def main():
    if not ENRICHMENT_RESULTS_PATH.exists():
        print(f"Enrichment results not found at {ENRICHMENT_RESULTS_PATH}. Please run enrich_full_menus.js first.")
        return

    print("Loading enrichment results...")
    with open(ENRICHMENT_RESULTS_PATH, "r", encoding="utf-8") as f:
        enrichment_results = json.load(f)

    # Map URI to enriched data
    uri_to_data = {item["uri"]: item for item in enrichment_results}
    print(f"Loaded {len(uri_to_data)} enriched stalls.")

    print("Backing up workbook...")
    backup_path = EXCEL_PATH.with_name(f"Hungry F&B Database.backup-apply-menu.xlsx")
    try:
        import shutil
        shutil.copy(EXCEL_PATH, backup_path)
        print(f"Saved backup copy to {backup_path}")
    except Exception as e:
        print(f"Warning: could not create backup copy: {e}")

    print("Loading Excel workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_TAB]

    headers = [clean(cell.value) for cell in ws[1]]
    menu_idx = headers.index("Menu")
    img_idx = headers.index("Dishes Image URL")
    uri_idx = headers.index("Google Maps URI")
    source_idx = headers.index("Menu Source") if "Menu Source" in headers else None

    if source_idx is None:
        print("Warning: 'Menu Source' column not found in sheet headers.")

    updated_count = 0
    skipped_count = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        uri = clean(row[uri_idx].value)
        if not uri or uri not in uri_to_data:
            skipped_count += 1
            continue

        data = uri_to_data[uri]
        dishes = data.get("dishes", [])
        if not dishes:
            skipped_count += 1
            continue

        # Create comma-separated lists
        names = []
        urls = []
        for dish in dishes:
            names.append(dish["name"].replace(",", ";")) # Replace comma inside dish names with semicolon to avoid splitting issues!
            img_url = dish.get("imageUrl", "")
            if not img_url:
                img_url = "https://raw.githubusercontent.com/siewhean/Hungry/main/assets/images/Logo/hungry_logo_icon_transparent.png"
            urls.append(img_url)

        menu_str = ", ".join(names)
        img_str = ",".join(urls) # Keep it comma-aligned with exactly same length

        # Update cells
        ws.cell(row=idx, column=menu_idx + 1, value=menu_str)
        ws.cell(row=idx, column=img_idx + 1, value=img_str)

        if source_idx is not None:
            source_val = "GrabFood Scrape" if data.get("source") == "grabfood" else "Claude AI Fallback"
            ws.cell(row=idx, column=source_idx + 1, value=source_val)

        updated_count += 1

    print(f"Applied enrichment: updated {updated_count} rows, skipped {skipped_count} rows.")
    wb.save(EXCEL_PATH)
    print("Excel database successfully saved.")

if __name__ == "__main__":
    main()
