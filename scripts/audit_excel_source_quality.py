#!/usr/bin/env python3
"""Audit source quality for the Hungry F&B workbook.

This is intentionally read-only. It does not try to "fix" rows from scraped
delivery-app data. It classifies every workbook row into source-quality flags so
the enrichment queue can prioritize real merchant/operator sources, compliant
images, and full menu verification.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook

from data_display_policy import (
    ALLOWED_RIGHTS_STATUSES as APPROVED_RIGHTS_STATUSES,
    approved_rights_status,
    display_authorized,
    is_placeholder_url,
    truthy,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "Hungry F&B Database.xlsx"
DEFAULT_SHEET = "Sheet1"
DEFAULT_REPORT = ROOT / "tasks" / "Roadmap" / "Database" / "excel-source-quality-audit-2026-06-09.json"
DEFAULT_QUEUE = ROOT / "tasks" / "Roadmap" / "Database" / "excel-source-quality-review-queue-2026-06-09.csv"

GENERIC_MENU_SOURCES = ("Cuisine Category Seed", "category_seed:")
OFFICIAL_MENU_PREFIXES = ("official_online_seed:", "online_price_menu:")
REFERENCE_MENU_PREFIXES = ("online_reference_seed:", "online_menu_candidate:", "source_description:")
UNSAFE_BULK_DOMAINS = ("grab.com", "foodpanda.", "deliveroo.")
GENERIC_IMAGE_HOSTS = ("images.unsplash.com",)
GOOGLE_HOSTS = ("googleapis.com", "googleusercontent.com", "gstatic.com")
URL_RE = re.compile(r"https?://[^\s,;)]+", re.IGNORECASE)


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def split_csv(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"\s*,\s*", value) if item.strip()]


def host(value: str) -> str:
    if not value.startswith(("http://", "https://")):
        return ""
    return urlparse(value).netloc.lower()


def hosts_in_text(value: str) -> list[str]:
    return [host(match.group(0)) for match in URL_RE.finditer(value)]


def source_family(menu_source: str) -> str:
    if not menu_source:
        return "missing"
    if menu_source == "Cuisine Category Seed":
        return "Cuisine Category Seed"
    if ":" in menu_source:
        return menu_source.split(":", 1)[0]
    return menu_source


def classify_row(row: dict[str, Any]) -> tuple[list[str], str]:
    flags: list[str] = []
    menu_source = clean(row.get("Menu Source"))
    menu_items = clean(row.get("Menu Items"))
    image_url = clean(row.get("Image URL"))
    dishes_image_url = clean(row.get("Dishes Image URL"))
    menu_approved = truthy(row.get("Menu Approved for Display")) or truthy(
        row.get("Approved for Display"),
    )
    image_approved = truthy(row.get("Image Approved for Display")) or truthy(
        row.get("Approved for Display"),
    )
    menu_rights_status = clean(row.get("Menu Rights Status"))
    image_rights_status = clean(row.get("Image Rights Status"))
    source_url = clean(row.get("Source URL"))
    google_website = clean(row.get("Google Website"))
    google_place_id = clean(row.get("Google Place ID"))
    google_rating = clean(row.get("Google Rating"))
    google_review_count = clean(row.get("Google Review Count"))

    menu_count = len(split_csv(menu_items))
    if not menu_source:
        flags.append("menu_source_missing")
    elif menu_source.startswith(GENERIC_MENU_SOURCES) or menu_source in GENERIC_MENU_SOURCES:
        flags.append("menu_generic_seed")
    elif menu_source.startswith(OFFICIAL_MENU_PREFIXES):
        flags.append("menu_operator_source")
    elif menu_source.startswith(REFERENCE_MENU_PREFIXES):
        flags.append("menu_reference_source")
    else:
        flags.append("menu_unclassified_source")

    if menu_count < 4:
        flags.append("menu_too_short")
    elif menu_count < 8:
        flags.append("menu_likely_partial")

    image_host = host(image_url)
    dish_image_host = host(dishes_image_url)
    if not image_url:
        flags.append("storefront_image_missing")
    elif any(domain in image_host for domain in GENERIC_IMAGE_HOSTS):
        flags.append("storefront_image_placeholder")
    elif any(domain in image_host for domain in GOOGLE_HOSTS):
        flags.append("storefront_image_google_display_only")
    else:
        flags.append("storefront_image_external_source")

    if not dishes_image_url:
        flags.append("dish_image_missing")
    elif any(domain in dish_image_host for domain in GENERIC_IMAGE_HOSTS):
        flags.append("dish_image_placeholder")
    elif any(domain in dish_image_host for domain in GOOGLE_HOSTS):
        flags.append("dish_image_google_display_only")
    else:
        flags.append("dish_image_external_source")

    source_hosts = " ".join([host(source_url), host(google_website), *hosts_in_text(menu_source)])
    if any(domain in source_hosts for domain in UNSAFE_BULK_DOMAINS):
        flags.append("delivery_app_source_quarantine")

    menu_rights_approved = approved_rights_status(menu_rights_status)
    image_rights_approved = approved_rights_status(image_rights_status)
    menu_display_safe = (
        menu_approved
        and menu_rights_approved
        and "delivery_app_source_quarantine" not in flags
    )
    image_display_safe = (
        image_approved
        and image_rights_approved
        and "storefront_image_placeholder" not in flags
        and "dish_image_placeholder" not in flags
        and "delivery_app_source_quarantine" not in flags
    )

    if menu_approved and not menu_display_safe:
        flags.append("production_menu_unverified")
    if image_approved and not image_display_safe:
        flags.append("production_image_unverified")
    if menu_display_safe:
        flags.append("menu_display_approved")
    if image_display_safe:
        flags.append("image_display_approved")

    if google_place_id:
        flags.append("google_matched")
    if google_rating and google_review_count:
        flags.append("google_rating_present")

    if "menu_operator_source" in flags and not any(
        item in flags
        for item in (
            "menu_too_short",
            "storefront_image_placeholder",
            "dish_image_placeholder",
            "delivery_app_source_quarantine",
        )
    ):
        status = "operator_verified_candidate"
    elif "menu_generic_seed" in flags or "storefront_image_placeholder" in flags or "dish_image_placeholder" in flags:
        status = "needs_source_verification"
    elif "menu_reference_source" in flags:
        status = "needs_manual_review"
    else:
        status = "review"

    return flags, status


def audit_supabase_cache(workbook_path: Path, limit: int | None = None) -> dict[str, Any]:
    """Audit the authoritative approval/rights columns in workbook cache sheets."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    cache_sheet_names = ("Supabase Venues", "Supabase Dishes")
    missing_sheets = [name for name in cache_sheet_names if name not in workbook.sheetnames]
    flag_counts: Counter[str] = Counter()
    blocked_samples: list[dict[str, Any]] = []
    blocked_count = 0
    rows_checked = 0
    try:
        for sheet_name in cache_sheet_names:
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if header_values is None:
                missing_sheets.append(sheet_name)
                continue
            headers = [clean(value) for value in header_values]
            sheet_rows_checked = 0
            for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                if limit is not None and sheet_rows_checked >= limit:
                    break
                sheet_rows_checked += 1
                rows_checked += 1
                row = {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                    if header
                }
                flags: list[str] = []
                menu_approved = truthy(row.get("menu_approved_for_display"))
                media_approved = truthy(row.get("media_approved_for_display"))
                menu_rights = row.get("menu_rights_status")
                media_rights = row.get("media_rights_status")
                media_urls = [row.get("image_url")]
                if sheet_name == "Supabase Venues":
                    media_urls.append(row.get("logo_url"))

                if not display_authorized(menu_approved, menu_rights):
                    flags.append("menu_not_authorized")
                if menu_approved and not approved_rights_status(menu_rights):
                    flags.append("menu_approval_rights_conflict")

                has_media = any(clean(value) for value in media_urls)
                if has_media and not display_authorized(media_approved, media_rights):
                    flags.append("media_not_authorized")
                if media_approved and not approved_rights_status(media_rights):
                    flags.append("media_approval_rights_conflict")
                if any(is_placeholder_url(value) for value in media_urls):
                    flags.append("placeholder_media")

                for flag in flags:
                    flag_counts[flag] += 1
                if flags:
                    blocked_count += 1
                    if len(blocked_samples) < 25:
                        blocked_samples.append(
                            {
                                "sheet": sheet_name,
                                "row_number": row_number,
                                "id": clean(row.get("id")),
                                "flags": flags,
                            }
                        )
    finally:
        workbook.close()

    if missing_sheets:
        flag_counts["missing_cache_sheet"] += len(missing_sheets)
        blocked_count += len(missing_sheets)
        for name in missing_sheets:
            if len(blocked_samples) < 25:
                blocked_samples.append(
                    {"sheet": name, "row_number": 1, "id": "", "flags": ["missing_cache_sheet"]}
                )

    return {
        "rows_checked": rows_checked,
        "missing_sheets": missing_sheets,
        "flag_counts": dict(flag_counts.most_common()),
        "blocked_rows": blocked_count,
        "sample_blocked_rows": blocked_samples,
        "passed": blocked_count == 0,
        "_blocked_count": blocked_count,
    }


def audit_workbook(workbook_path: Path, sheet_name: str, limit: int | None = None) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"Workbook missing sheet {sheet_name!r}. Available: {', '.join(workbook.sheetnames)}")

    worksheet = workbook[sheet_name]
    headers = [clean(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows: list[dict[str, Any]] = []
    status_counts: Counter[str] = Counter()
    flag_counts: Counter[str] = Counter()
    menu_source_counts: Counter[str] = Counter()
    image_host_counts: Counter[str] = Counter()
    dish_image_host_counts: Counter[str] = Counter()
    item_count_counts: Counter[int] = Counter()

    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: values[index] if index < len(values) else None for index, header in enumerate(headers)}
        if not any(clean(value) for value in row.values()):
            continue
        flags, status = classify_row(row)
        row["Row Number"] = row_number
        row["QA Status"] = status
        row["QA Flags"] = ";".join(flags)
        rows.append(row)

        status_counts[status] += 1
        for flag in flags:
            flag_counts[flag] += 1
        menu_source_counts[source_family(clean(row.get("Menu Source")))] += 1
        image_host_counts[host(clean(row.get("Image URL"))) or "missing"] += 1
        dish_image_host_counts[host(clean(row.get("Dishes Image URL"))) or "missing"] += 1
        item_count_counts[len(split_csv(clean(row.get("Menu Items"))))] += 1

        if limit is not None and len(rows) >= limit:
            break

    workbook.close()
    total = len(rows)
    cache_audit = audit_supabase_cache(workbook_path, limit)
    blocked_cache_rows = cache_audit.pop("_blocked_count")
    report = {
        "workbook": str(workbook_path),
        "sheet": sheet_name,
        "rows_checked": total,
        "status_counts": dict(status_counts.most_common()),
        "flag_counts": dict(flag_counts.most_common()),
        "menu_source_counts": dict(menu_source_counts.most_common(30)),
        "image_host_counts": dict(image_host_counts.most_common(20)),
        "dish_image_host_counts": dict(dish_image_host_counts.most_common(20)),
        "menu_item_count_distribution": dict(sorted(item_count_counts.items())),
        "critical_findings": {
            "placeholder_storefront_images": flag_counts["storefront_image_placeholder"],
            "placeholder_dish_images": flag_counts["dish_image_placeholder"],
            "generic_menu_seed_rows": flag_counts["menu_generic_seed"],
            "likely_partial_menu_rows": flag_counts["menu_likely_partial"] + flag_counts["menu_too_short"],
            "delivery_app_source_quarantine_rows": flag_counts["delivery_app_source_quarantine"],
            "production_menu_unverified_rows": flag_counts["production_menu_unverified"],
            "production_image_unverified_rows": flag_counts["production_image_unverified"],
            "menu_display_approved_rows": flag_counts["menu_display_approved"],
            "image_display_approved_rows": flag_counts["image_display_approved"],
        },
        "supabase_cache_display_audit": cache_audit,
    }
    return {"report": report, "rows": rows, "blocked_cache_rows": blocked_cache_rows}


def unapproved_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    blocked_flags = {
        "menu_generic_seed",
        "menu_source_missing",
        "menu_unclassified_source",
        "menu_reference_source",
        "storefront_image_placeholder",
        "dish_image_placeholder",
        "delivery_app_source_quarantine",
        "production_menu_unverified",
        "production_image_unverified",
    }
    blocked: list[dict[str, Any]] = []
    for row in rows:
        flags = {flag for flag in clean(row.get("QA Flags")).split(";") if flag}
        if (
            flags & blocked_flags
            or "menu_display_approved" not in flags
            or "image_display_approved" not in flags
        ):
            blocked.append(row)
    return blocked


def write_review_queue(rows: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "Row Number",
        "Business ID",
        "Stall / Restaurant Name",
        "Stall Type",
        "Cuisine",
        "Signature Dish",
        "Menu Source",
        "Menu Items",
        "Source URL",
        "Google Website",
        "Image URL",
        "Dishes Image URL",
        "QA Status",
        "QA Flags",
    ]
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            if row["QA Status"] != "operator_verified_candidate":
                writer.writerow(row)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--queue", type=Path, default=DEFAULT_QUEUE)
    parser.add_argument(
        "--fail-unapproved-display",
        action="store_true",
        help="Exit non-zero if the Supabase cache sheets contain display approval/rights blockers.",
    )
    args = parser.parse_args()

    result = audit_workbook(args.workbook, args.sheet, args.limit)
    report = result["report"]
    rows = result["rows"]

    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    write_review_queue(rows, args.queue)
    blocked_rows = result["blocked_cache_rows"]
    report["display_gate"] = {
        "mode": "licensed_or_owner_authorized_only",
        "blocked_rows": blocked_rows,
        "passed": blocked_rows == 0,
    }
    args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True))
    print(f"review_queue={args.queue}")
    if args.fail_unapproved_display and blocked_rows:
        print(
            f"display_gate_failed={blocked_rows} rows require licensed/owned menu and image provenance",
        )
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
