#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it in the active Python environment, "
        "or run with the Codex workspace Python runtime used for workbook tooling."
    ) from exc


BAD_EXACT_URLS = {
    "https://eats.oddle.me",
    "https://menustic.com",
    "https://www.foodline.sg/",
    "https://www.mitsueki.sg/",
}
LOW_CONFIDENCE_PRICE_MENU_URLS = {
    "https://capellahotels.com/assets/docs/singapore/King_of_Culinary_Menu_1.pdf",
}
BAD_DOMAIN_PATHS = (
    ("maps.app.goo.gl", ""),
    ("ordinarypatrons.com", "/singapore-food-dictionary/"),
    ("singaporeaneats.wordpress.com", "/category/"),
    ("wineconnection.com.sg", "/collections/glasses"),
    ("yellowpages.com.sg", "/company/"),
    ("burpple.com", "/@"),
    ("burpple.com", "/categories/"),
    ("bestlah.sg", "/food/best-"),
    ("danielfooddiary.com", "/"),
)


def normalize_spaces(value: Any) -> str:
    return " ".join(("" if value is None else str(value)).split())


def should_clear(source_value: Any) -> bool:
    source = normalize_spaces(source_value)
    if not source.startswith("online_menu_candidate: "):
        return False
    url = source.split(": ", 1)[1].strip()
    if url in BAD_EXACT_URLS:
        return True
    parsed = urlparse(url)
    domain = parsed.netloc.lower().removeprefix("www.")
    path = parsed.path.lower()
    return any(domain == bad_domain and path.startswith(bad_path) for bad_domain, bad_path in BAD_DOMAIN_PATHS)


def should_clear_menu_fields(source_value: Any) -> bool:
    source = normalize_spaces(source_value)
    if source.startswith("online_price_menu: "):
        url = source.split(": ", 1)[1].strip()
        return url in LOW_CONFIDENCE_PRICE_MENU_URLS
    if source.startswith("name_propagated: "):
        return any(url in source for url in LOW_CONFIDENCE_PRICE_MENU_URLS)
    return False


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = {normalize_spaces(cell.value): index for index, cell in enumerate(ws[1], start=1)}
    source_col = headers.get("Menu Source")
    if source_col is None:
        raise SystemExit("Missing required column: Menu Source")
    menu_field_cols = [
        headers.get("Signature Dish"),
        headers.get("Menu"),
        headers.get("Menu Items"),
        source_col,
    ]

    cleared: list[dict[str, Any]] = []
    name_col = headers.get("Stall / Restaurant Name")
    for row_num in range(2, ws.max_row + 1):
        cell = ws.cell(row_num, source_col)
        if not should_clear(cell.value):
            continue
        cleared.append(
            {
                "row": row_num,
                "name": ws.cell(row_num, name_col).value if name_col else None,
                "source": cell.value,
            }
        )
        if args.apply:
            cell.value = None

    menu_fields_cleared: list[dict[str, Any]] = []
    for row_num in range(2, ws.max_row + 1):
        cell = ws.cell(row_num, source_col)
        if not should_clear_menu_fields(cell.value):
            continue
        menu_fields_cleared.append(
            {
                "row": row_num,
                "name": ws.cell(row_num, name_col).value if name_col else None,
                "source": cell.value,
            }
        )
        if args.apply:
            for col_num in menu_field_cols:
                if col_num is not None:
                    ws.cell(row_num, col_num).value = None

    backup_path = None
    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "apply": args.apply,
                "backup": str(backup_path) if backup_path else None,
                "cleared_count": len(cleared),
                "cleared": cleared,
                "menu_fields_cleared_count": len(menu_fields_cleared),
                "menu_fields_cleared": menu_fields_cleared,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
