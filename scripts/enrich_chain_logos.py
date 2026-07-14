import json
import uuid
import pandas as pd
from openpyxl import load_workbook

# Clearbit Logo API: https://logo.clearbit.com/[domain]
CHAIN_DOMAINS = {
    "mcdonald's": "mcdonalds.com.sg",
    "kfc": "kfc.com.sg",
    "burger king": "burgerking.com.sg",
    "subway": "subway.com",
    "toast box": "toastbox.com.sg",
    "ya kun": "yakun.com",
    "koi": "koithe.com",
    "gong cha": "gong-cha-sg.com",
    "liho": "lihotea.com",
    "mr coconut": "mrcoconut.sg",
    "starbucks": "starbucks.com.sg",
    "coffee bean": "coffeebean.com.sg",
    "peperoni": "peperonipizzeria.com.sg",
    "domino's": "dominos.com.sg",
    "pizza hut": "pizzahut.com.sg",
    "astons": "astons.com.sg",
    "old chang kee": "oldchangkee.com"
}

def get_logo_url(stall_name: str) -> str:
    name_lower = str(stall_name).lower()
    for chain, domain in CHAIN_DOMAINS.items():
        if chain in name_lower:
            return f"https://logo.clearbit.com/{domain}"
    return ""

def main():
    file_path = "data/Hungry F&B Database.xlsx"
    print(f"Loading {file_path}...")

    # Use pandas to find rows easily, but write with openpyxl to preserve formatting
    df = pd.read_excel(file_path, sheet_name="Sheet1")

    # Check if 'Logo URL' column exists, if not add it
    if 'Logo URL' not in df.columns:
        df['Logo URL'] = ""

    updates = 0
    for idx, row in df.iterrows():
        stall_name = row.get("Stall / Restaurant Name", "")
        if pd.isna(stall_name):
            continue

        current_logo = row.get("Logo URL", "")
        if pd.isna(current_logo) or not current_logo:
            new_logo = get_logo_url(stall_name)
            if new_logo:
                df.at[idx, 'Logo URL'] = new_logo
                updates += 1

    print(f"Found {updates} chains to enrich with logos.")

    if updates > 0:
        print("Saving to Excel...")
        # Since reading/writing with pandas destroys advanced Excel formatting,
        # we update openpyxl directly
        wb = load_workbook(file_path)
        ws = wb["Sheet1"]

        # Find column indices
        headers = {cell.value: idx for idx, cell in enumerate(ws[1], 1)}
        if 'Logo URL' not in headers:
            # Append new column
            new_col_idx = ws.max_column + 1
            ws.cell(row=1, column=new_col_idx, value="Logo URL")
            headers['Logo URL'] = new_col_idx

        logo_col_idx = headers['Logo URL']
        stall_col_idx = headers['Stall / Restaurant Name']

        for row in range(2, ws.max_row + 1):
            stall_name = ws.cell(row=row, column=stall_col_idx).value
            if stall_name:
                logo_url = get_logo_url(stall_name)
                if logo_url:
                    ws.cell(row=row, column=logo_col_idx, value=logo_url)

        wb.save(file_path)
        print("Done saving!")

if __name__ == "__main__":
    main()
