#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

DEPENDENCY_HINT = (
    "Install the missing package in the active Python environment, or run with "
    "the Codex workspace Python runtime used for workbook tooling."
)

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(f"Missing dependency: openpyxl. {DEPENDENCY_HINT}") from exc

try:
    from pypdf import PdfReader
except ModuleNotFoundError:
    PdfReader = None


PRICE_RE = re.compile(r"(?:S\$|\$)\s*\d{1,3}(?:\.\d{1,2})?|\b\d{1,3}\.\d{2}\b", re.IGNORECASE)
PRICE_ONLY_RE = re.compile(r"^(?:S\$|\$)?\s*\d{1,3}(?:\.\d{1,2})?$", re.IGNORECASE)
FOOD_WORDS = {
    "ayam",
    "bak",
    "beef",
    "biryani",
    "bread",
    "burger",
    "cake",
    "cheese",
    "chicken",
    "chocolate",
    "coffee",
    "crab",
    "curry",
    "duck",
    "egg",
    "fish",
    "fries",
    "garlic",
    "hor fun",
    "kaya",
    "laksa",
    "lamb",
    "mee",
    "miso",
    "mutton",
    "noodle",
    "oyster",
    "pasta",
    "pizza",
    "pork",
    "prata",
    "prawn",
    "rice",
    "salad",
    "salmon",
    "sandwich",
    "satay",
    "seafood",
    "soup",
    "steak",
    "sushi",
    "tart",
    "toast",
    "tofu",
    "tuna",
    "waffle",
}
REJECT_RE = re.compile(
    r"\b(?:"
    r"address|all prices|available|beverage|booking|cart|contact|copyright|delivery|"
    r"download|event|facebook|follow|gallery|gift card|instagram|location|login|"
    r"menu|minimum|opening|order online|outlet|privacy|reservation|service charge|"
    r"subject to|terms|wine list"
    r")\b",
    re.IGNORECASE,
)
DRINK_ONLY_RE = re.compile(r"\b(?:beer|cocktail|gin|spirits|vodka|whisky|wine)\b", re.IGNORECASE)
SOURCE_PREFIX = "online_menu_candidate: "
TARGET_PREFIX = "online_price_menu: "


class TextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = clean(data)
        if text:
            self.parts.append(text)


@dataclass(frozen=True)
class FetchResult:
    url: str
    content_type: str
    body: bytes


@dataclass(frozen=True)
class ExtractionResult:
    status: str
    items: tuple[str, ...]
    source_url: str
    error: str = ""


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return clean(value) in {"", "-", "—", "N/A", "NA", "None", "null"}


def fetch_url(url: str, timeout_s: int) -> FetchResult:
    request = Request(
        url,
        headers={
            "User-Agent": "HungryDataEnrichmentBot/0.1 (+source-confirmed price menu extraction)",
            "Accept": "text/html,application/pdf;q=0.9,*/*;q=0.4",
        },
    )
    with urlopen(request, timeout=timeout_s) as response:
        body = response.read(2_500_000)
        content_type = response.headers.get("content-type", "")
        final_url = response.geturl()
    return FetchResult(url=final_url, content_type=content_type, body=body)


def html_units(body: bytes) -> list[str]:
    parser = TextParser()
    parser.feed(body.decode("utf-8", errors="replace"))
    return parser.parts


def pdf_units(body: bytes) -> list[str]:
    if PdfReader is None:
        raise RuntimeError(f"Missing dependency: pypdf. {DEPENDENCY_HINT}")
    reader = PdfReader(io.BytesIO(body))
    units: list[str] = []
    for page in reader.pages[:10]:
        text = page.extract_text() or ""
        units.extend(clean(line) for line in text.splitlines() if clean(line))
    return units


def text_units(fetched: FetchResult) -> tuple[list[str], str]:
    content_type = fetched.content_type.lower()
    start = fetched.body[:80].lower()
    if "pdf" in content_type or start.startswith(b"%pdf"):
        return pdf_units(fetched.body), "pdf"
    if "html" in content_type or b"html" in start:
        return html_units(fetched.body), "html"
    return [], "unsupported"


def normalize_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def contains_food_word(value: str) -> bool:
    lowered = value.lower()
    return any(word in lowered for word in FOOD_WORDS)


def clean_item(value: str) -> str:
    text = clean(value)
    text = re.sub(r"\([^)]*(?:calories|kcal|contains|spicy|vegetarian)[^)]*\)", " ", text, flags=re.IGNORECASE)
    text = PRICE_RE.sub(" ", text)
    text = re.sub(r"\b(?:from|nett|sgd)\b", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*•\s*[A-Z ]{1,8}$", " ", text)
    text = re.sub(r"\s+[NVD]{1,4}$", " ", text)
    text = re.sub(r"\s+recommended.*$", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"^[|:;,.•*+\-–—/\\]+", "", text)
    text = re.sub(r"[|:;,.•*+\-–—/\\]+$", "", text)
    text = clean(text)
    text = re.sub(r"^(?:add|choice of|extra|set of)\s+", "", text, flags=re.IGNORECASE)
    return clean(text).strip(" |:;,.•*+-–—/\\")


def valid_item(value: str) -> bool:
    text = clean_item(value)
    if not (5 <= len(text) <= 70):
        return False
    lowered = text.lower()
    if text.count(",") >= 2:
        return False
    if lowered.startswith(("for ", "served with ", "with ", "choice of ")):
        return False
    if " – " in text or " - " in text:
        return False
    if REJECT_RE.search(lowered):
        return False
    if DRINK_ONLY_RE.search(lowered) and not contains_food_word(text):
        return False
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", text)
    if len(words) < 2 or len(words) > 10:
        return False
    if text.isupper() and len(words) <= 2:
        return False
    if not contains_food_word(text):
        return False
    return True


def candidate_lines(units: list[str], pair_price_only: bool) -> list[str]:
    lines: list[str] = []
    previous = ""
    for unit in units:
        text = clean(unit)
        if not text:
            continue
        if PRICE_RE.search(text):
            if pair_price_only and PRICE_ONLY_RE.fullmatch(text) and previous:
                lines.append(previous)
            elif not PRICE_ONLY_RE.fullmatch(text):
                lines.append(text)
        previous = text if not PRICE_ONLY_RE.fullmatch(text) else previous
    return lines


def extract_from_url(url: str, timeout_s: int) -> ExtractionResult:
    try:
        fetched = fetch_url(url, timeout_s)
        units, source_type = text_units(fetched)
        if not units:
            return ExtractionResult(status="unsupported_or_empty", items=(), source_url=fetched.url)
        items = extract_items_from_source_units(units, source_type)
        status = "price_menu_items" if len(items) >= 3 else "insufficient_items"
        return ExtractionResult(status=status, items=items, source_url=fetched.url)
    except Exception as exc:
        return ExtractionResult(status="fetch_failed", items=(), source_url=url, error=str(exc)[:180])


def extract_items_from_source_units(units: list[str], source_type: str) -> tuple[str, ...]:
    items: list[str] = []
    seen: set[str] = set()
    for line in candidate_lines(units, pair_price_only=source_type == "pdf"):
        item = clean_item(line)
        key = normalize_key(item)
        if not key or key in seen or not valid_item(item):
            continue
        seen.add(key)
        items.append(item)
        if len(items) >= 10:
            break
    return tuple(items)


def headers_for(sheet) -> dict[str, int]:
    return {clean(cell.value): index for index, cell in enumerate(sheet[1], start=1)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100000)
    parser.add_argument("--max-sources", type=int, default=50)
    parser.add_argument("--min-items", type=int, default=3)
    parser.add_argument("--timeout", type=int, default=10)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    parser.add_argument("--updates-only", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = headers_for(ws)
    required = ["Stall / Restaurant Name", "Signature Dish", "Menu", "Menu Items", "Menu Source"]
    missing = [field for field in required if field not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    checked = 0
    updated = 0
    report: list[dict[str, Any]] = []
    last_row = min(ws.max_row, args.start_row + args.max_rows - 1)
    for row_num in range(args.start_row, last_row + 1):
        source_value = clean(ws.cell(row_num, headers["Menu Source"]).value)
        if not source_value.startswith(SOURCE_PREFIX):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu Items"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Signature Dish"]).value):
            continue
        if checked >= args.max_sources:
            break
        checked += 1
        source_url = source_value[len(SOURCE_PREFIX) :].strip()
        result = extract_from_url(source_url, args.timeout)
        row_report = {
            "row": row_num,
            "name": ws.cell(row_num, headers["Stall / Restaurant Name"]).value,
            "source": source_url,
            "status": result.status,
            "items": result.items,
            "error": result.error,
        }
        if len(result.items) >= args.min_items:
            item_text = ", ".join(result.items)
            if args.apply:
                ws.cell(row_num, headers["Signature Dish"]).value = result.items[0]
                ws.cell(row_num, headers["Menu"]).value = item_text
                ws.cell(row_num, headers["Menu Items"]).value = item_text
                ws.cell(row_num, headers["Menu Source"]).value = f"{TARGET_PREFIX}{result.source_url}"
            updated += 1
            row_report["would_update"] = True
        else:
            row_report["would_update"] = False
        if not args.updates_only or row_report["would_update"]:
            report.append(row_report)

    backup_path = None
    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "checked": checked,
                "updated": updated,
                "apply": args.apply,
                "backup": str(backup_path) if backup_path else None,
                "report": report,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
