#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import urllib.robotparser
import shutil
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import HTTPRedirectHandler, Request, build_opener

from openpyxl import load_workbook
from pypdf import PdfReader


MENU_LINK_HINTS = ("menu", "food", "dining")
EXCLUDED_LINK_HINTS = (
    "privacy",
    "terms",
    "careers",
    "login",
    "facebook.com",
    "instagram.com",
    "food-dictionary",
    "maps.app.goo.gl",
    "yellowpages.com.sg",
)
EXCLUDED_SOURCE_DOMAINS = ("facebook.com", "instagram.com", "tiktok.com", "linkedin.com", "maps.app.goo.gl")
THIRD_PARTY_ORDERING_DOMAINS = ("oddle.me", "grab.com", "foodpanda.sg", "foodpanda.com", "deliveroo.com", "chope.co", "getz.co", "tablecheck.com")
EXCLUDED_SOURCE_PATH_HINTS = (
    "/category/",
    "/collections/glasses",
    "/company/",
    "/singapore-food-dictionary",
)
STOCK_IMAGE_HOSTS = ("unsplash.com", "pexels.com", "shutterstock.com", "gettyimages.com", "istockphoto.com", "pixabay.com", "freepik.com", "adobestock.com")
PLACEHOLDER_IMAGE_MARKERS = ("placeholder", "default-image", "no-image", "transparent", "spacer", "1x1", "blank.gif", "dummy", "revslider", "logo", "icon", "sprite")
ROBOTS_CACHE: dict[str, urllib.robotparser.RobotFileParser] = {}


class SameOriginRedirectHandler(HTTPRedirectHandler):
    """Reject redirects before making a request to a different origin."""
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        old, new = urlparse(req.full_url), urlparse(newurl)
        if (old.scheme, old.netloc.lower()) != (new.scheme, new.netloc.lower()):
            raise ValueError("cross-origin redirect rejected")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


URL_OPENER = build_opener(SameOriginRedirectHandler())
USER_AGENT = "HungryDataEnrichmentBot/0.1"
EXCLUDED_ITEM_LINES = {
    "menu",
    "food menu",
    "drinks menu",
    "beverage",
    "beverages",
    "dessert",
    "desserts",
    "savoury",
    "savory",
    "takeaways",
    "contact",
    "location",
    "home",
    "gallery",
    "events",
    "reservations",
    "subscribe",
    "cart",
    "checkout",
    "continue shopping",
    "contact us",
    "download menu",
    "events and reservations",
    "item added to your cart",
    "loading your experience… this won’t take long.",
    "log in",
    "operating hours",
    "powered by shopify",
    "reservations & events",
    "singapore singapore",
    "skip to content",
    "we are on grabfood!",
    "we’re getting things ready",
}
FOOD_WORDS = {
    "beef",
    "beer",
    "bread",
    "burger",
    "cake",
    "cheese",
    "chicken",
    "chocolate",
    "coffee",
    "crab",
    "curry",
    "dessert",
    "duck",
    "egg",
    "fish",
    "fries",
    "garlic",
    "ice",
    "kaya",
    "laksa",
    "lamb",
    "mango",
    "mee",
    "milk",
    "miso",
    "noodle",
    "pasta",
    "pizza",
    "pork",
    "prawn",
    "rice",
    "salad",
    "salmon",
    "sambal",
    "satay",
    "seafood",
    "soup",
    "tea",
    "toast",
    "tofu",
    "tuna",
    "waffle",
    "wine",
}


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            self._href = dict(attrs).get("href")
            self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._href is not None:
            text = normalize_spaces(" ".join(self._text))
            self.links.append((self._href, text))
            self._href = None
            self._text = []


class TextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = normalize_spaces(data)
        if text:
            self.parts.append(text)


class ImageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.images: list[dict[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        if tag.lower() == "img":
            src = values.get("src") or values.get("data-src") or values.get("data-original") or ""
            if src:
                self.images.append({"url": src, "alt": normalize_spaces(values.get("alt")), "title": normalize_spaces(values.get("title")), "kind": "img"})
        elif tag.lower() == "meta" and values.get("property", "").lower() in {"og:image", "og:image:url", "twitter:image"}:
            src = values.get("content") or ""
            if src:
                self.images.append({"url": src, "alt": "", "title": "", "kind": values.get("property", "og:image")})


@dataclass(frozen=True)
class FetchResult:
    url: str
    content_type: str
    body: bytes


@dataclass(frozen=True)
class MenuExtraction:
    source_url: str
    items: list[str]
    status: str
    error: str = ""
    images: list[dict[str, str]] | None = None


def normalize_spaces(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return normalize_spaces(value) in {"", "-", "—", "N/A", "NA", "None", "null"}


def is_excluded_source_url(url: str) -> bool:
    parsed = urlparse(url)
    domain = parsed.netloc.lower()
    path = parsed.path.lower()
    if any(excluded_domain in domain for excluded_domain in EXCLUDED_SOURCE_DOMAINS):
        return True
    return any(hint in path for hint in EXCLUDED_SOURCE_PATH_HINTS)


def is_third_party_ordering_url(url: str) -> bool:
    host = urlparse(url).netloc.lower().removeprefix("www.")
    return any(host == domain or host.endswith(f".{domain}") for domain in THIRD_PARTY_ORDERING_DOMAINS)


def fetch_url(url: str, timeout_s: int) -> FetchResult:
    request = Request(
        url,
        headers={
            "User-Agent": f"{USER_AGENT} (+official menu source enrichment)",
            "Accept": "text/html,application/pdf;q=0.9,*/*;q=0.5",
        },
    )
    with URL_OPENER.open(request, timeout=timeout_s) as response:
        body = response.read(1_500_000)
        content_type = response.headers.get("content-type", "")
        final_url = response.geturl()
    initial = urlparse(url)
    final = urlparse(final_url)
    if (initial.scheme, initial.netloc.lower()) != (final.scheme, final.netloc.lower()):
        raise ValueError("redirect left the requested source origin")
    return FetchResult(url=final_url, content_type=content_type, body=body)


def robots_allows(url: str, timeout_s: int) -> bool:
    parsed = urlparse(url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    if origin not in ROBOTS_CACHE:
        parser = urllib.robotparser.RobotFileParser()
        parser.set_url(f"{origin}/robots.txt")
        try:
            result = fetch_url(parser.url, timeout_s)
            parser.parse(result.body.decode("utf-8", errors="replace").splitlines())
        except HTTPError as exc:
            if exc.code == 404:
                parser.parse([])
            else:
                parser.parse(["User-agent: *", "Disallow: /"])
        except (URLError, TimeoutError, OSError, ValueError):
            parser.parse(["User-agent: *", "Disallow: /"])
        ROBOTS_CACHE[origin] = parser
    return ROBOTS_CACHE[origin].can_fetch(USER_AGENT, url)


def html_text(body: bytes) -> str:
    parser = TextParser()
    parser.feed(body.decode("utf-8", errors="replace"))
    return "\n".join(parser.parts)


def html_images(page_url: str, body: bytes, official_home_url: str) -> list[dict[str, str]]:
    parser = ImageParser()
    parser.feed(body.decode("utf-8", errors="replace"))
    official_host = urlparse(official_home_url).netloc.lower().removeprefix("www.")
    results: list[dict[str, str]] = []
    seen: set[str] = set()
    for image in parser.images:
        url = urljoin(page_url, image["url"])
        parsed = urlparse(url)
        host = parsed.netloc.lower().removeprefix("www.")
        candidate_text = f"{url} {image['alt']} {image['title']}".lower()
        if parsed.scheme not in {"http", "https"} or (host != official_host and not host.endswith(f".{official_host}")):
            continue
        if any(marker in candidate_text for marker in PLACEHOLDER_IMAGE_MARKERS) or any(stock in host for stock in STOCK_IMAGE_HOSTS):
            continue
        if url in seen:
            continue
        seen.add(url)
        results.append({"image_url": url, "alt": image["alt"], "title": image["title"], "html_kind": image["kind"]})
    return results[:80]


def html_links(base_url: str, body: bytes) -> list[str]:
    parser = LinkParser()
    parser.feed(body.decode("utf-8", errors="replace"))
    candidates: list[str] = []
    seen: set[str] = set()
    for href, text in parser.links:
        combined = f"{href} {text}".lower()
        if not any(hint in combined for hint in MENU_LINK_HINTS):
            continue
        if any(hint in combined for hint in EXCLUDED_LINK_HINTS):
            continue
        full_url = urljoin(base_url, href)
        parsed = urlparse(full_url)
        if parsed.scheme not in {"http", "https"}:
            continue
        if is_excluded_source_url(full_url):
            continue
        normalized = full_url.split("#", 1)[0]
        if normalized in seen:
            continue
        seen.add(normalized)
        candidates.append(normalized)
    return candidates


def pdf_text(body: bytes) -> str:
    reader = PdfReader(io.BytesIO(body))
    parts: list[str] = []
    for page in reader.pages[:8]:
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n".join(parts)


def clean_item_line(line: str) -> str:
    text = normalize_spaces(line)
    text = re.sub(r"\bS?\$?\s*\d+(?:\.\d{1,2})?\b", " ", text)
    text = re.sub(r"\s{2,}", " ", text)
    text = re.sub(r"^[|•*\\-–—]+", "", text).strip(" |•*\\-–—")
    return text


def looks_like_menu_item(line: str) -> bool:
    text = clean_item_line(line)
    if not (3 <= len(text) <= 72):
        return False
    lowered = text.lower()
    if lowered in EXCLUDED_ITEM_LINES:
        return False
    if re.search(
        r"(copyright|domain owner|facebook|instagram|loading your experience|powered by|privacy|terms|address|opening hours|subscribe)",
        lowered,
    ):
        return False
    if re.search(r"\b(?:mon|tue|wed|thu|fri|sat|sun)(?:day)?\b", lowered) and re.search(
        r"\d+\s*(?:am|pm)", lowered
    ):
        return False
    if re.fullmatch(r"\d+\s*(?:am|pm)\s*[-–]\s*\d+\s*(?:am|pm)", lowered):
        return False
    if re.fullmatch(r"[A-Z0-9 &/+.-]{3,32}", text) and not any(word in lowered for word in FOOD_WORDS):
        return False
    words = re.findall(r"[A-Za-z][A-Za-z'-]+", text)
    if not words:
        return False
    return any(word in lowered for word in FOOD_WORDS)


def extract_items_from_text(text: str) -> list[str]:
    raw_lines = re.split(r"[\r\n]+", text)
    items: list[str] = []
    seen: set[str] = set()
    for raw_line in raw_lines:
        for candidate in re.split(r"\s{3,}|(?:\s+[•|]\s+)", raw_line):
            item = clean_item_line(candidate)
            key = item.lower()
            if key in seen or not looks_like_menu_item(item):
                continue
            seen.add(key)
            items.append(item)
            if len(items) >= 12:
                return items
    return items


def has_high_confidence_items(items: list[str]) -> bool:
    return len(items) >= 3


def quality_menu_items(items: list[str], source_name: str) -> list[str]:
    source_key = re.sub(r"[^a-z0-9]+", " ", source_name.lower()).strip()
    quality: list[str] = []
    seen: set[str] = set()
    for item in items:
        text = clean_item_line(item)
        lowered = text.lower()
        key = re.sub(r"[^a-z0-9]+", " ", lowered).strip()
        if key in seen:
            continue
        if source_key and (key == source_key or source_key in key):
            continue
        if re.search(
            r"\b(?:"
            r"address|applicable|delivery|facilities|food prices|hours|legal|menu|near you|"
            r"notice|packages|price|prices|request|services|subject|wedding"
            r")\b",
            lowered,
        ):
            continue
        words = re.findall(r"[A-Za-z][A-Za-z'-]*", text)
        if len(words) < 2 or len(words) > 10:
            continue
        if text.isupper() and len(words) <= 3 and re.search(r"\b(?:and|&)\b", lowered):
            continue
        if not any(word in lowered for word in FOOD_WORDS):
            continue
        seen.add(key)
        quality.append(text)
    return quality[:12]


def has_low_noise_quality(items: list[str]) -> bool:
    if not has_high_confidence_items(items):
        return False
    noisy_lines = 0
    for item in items[:8]:
        lowered = item.lower()
        if item[:1].islower():
            noisy_lines += 1
        if item.endswith("."):
            noisy_lines += 1
        if re.search(
            r"\b(?:approx|contains|dairy|gelatin|gelatine|gst|ingredients?|prices?|subjected|experience)\b",
            lowered,
        ):
            noisy_lines += 1
    return noisy_lines <= 1


def extract_from_source(url: str, timeout_s: int) -> MenuExtraction:
    try:
        if not robots_allows(url, timeout_s):
            return MenuExtraction(source_url=url, items=[], status="blocked_by_robots")
        fetched = fetch_url(url, timeout_s)
        content_type = fetched.content_type.lower()
        body_start = fetched.body[:20].lower()
        if "pdf" in content_type or body_start.startswith(b"%pdf"):
            text = pdf_text(fetched.body)
            images = []
        elif "html" in content_type or b"html" in body_start:
            text = html_text(fetched.body)
            images = html_images(fetched.url, fetched.body, f"{urlparse(fetched.url).scheme}://{urlparse(fetched.url).netloc}")
        else:
            return MenuExtraction(source_url=fetched.url, items=[], status="unsupported_content")
        return MenuExtraction(
            source_url=fetched.url,
            items=extract_items_from_text(text),
            status="extracted",
            images=images,
        )
    except Exception as exc:
        return MenuExtraction(source_url=url, items=[], status="fetch_failed", error=str(exc)[:160])


def find_menu_sources(home_url: str, timeout_s: int) -> list[str]:
    home_domain = urlparse(home_url).netloc.lower()
    if any(domain in home_domain for domain in EXCLUDED_SOURCE_DOMAINS):
        return []
    if not robots_allows(home_url, timeout_s):
        raise PermissionError("robots.txt disallows homepage access")
    fetched = fetch_url(home_url, timeout_s)
    candidates = html_links(fetched.url, fetched.body)
    same_domain = urlparse(fetched.url).netloc.lower().removeprefix("www.")
    # Keep extraction on the operator's own host. External delivery/social/directory pages
    # are not treated as an official menu source.
    candidates = [item for item in candidates if urlparse(item).netloc.lower().removeprefix("www.") == same_domain]
    ranked = sorted(
        candidates,
        key=lambda item: (
            0 if item.lower().split("?", 1)[0].endswith(".pdf") else 1,
            0 if urlparse(item).netloc.lower().removeprefix("www.") == same_domain else 1,
            len(item),
        ),
    )
    return ranked[:4]


def load_headers(sheet) -> dict[str, int]:
    return {normalize_spaces(cell.value): index for index, cell in enumerate(sheet[1], start=1)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=200)
    parser.add_argument("--max-websites", type=int, default=25)
    parser.add_argument("--timeout", type=int, default=12)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--apply-menu-sources", action="store_true")
    parser.add_argument("--source-only", action="store_true")
    parser.add_argument("--refresh-review", action="store_true", help="Review current nonblank menus from official-site sources without applying changes")
    parser.add_argument("--public-access-csv", default="data/singapore_fnb_verified.csv")
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()
    if args.refresh_review and args.apply:
        raise SystemExit("--refresh-review is review-only; do not combine it with --apply")

    workbook_path = Path(args.workbook).expanduser()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = load_headers(ws)
    required = ["Stall / Restaurant Name", "Google Website", "Menu Source", "Menu Items"]
    missing = [field for field in required if field not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    menu_col = headers.get("Menu")
    licence_col = headers.get("Licence No")
    if args.refresh_review and not licence_col:
        raise SystemExit("--refresh-review requires a Licence No column")
    public_licences: set[str] = set()
    if args.refresh_review:
        with Path(args.public_access_csv).expanduser().open(newline="", encoding="utf-8-sig") as source:
            public_licences = {
                re.sub(r"[^A-Z0-9]", "", (row.get("licence_number") or "").upper())
                for row in csv.DictReader(source)
            }
    all_workbook_licences: dict[str, int] = {}
    if args.refresh_review:
        for row_num in range(2, ws.max_row + 1):
            licence_key = re.sub(r"[^A-Z0-9]", "", normalize_spaces(ws.cell(row_num, licence_col).value).upper())
            if licence_key:
                all_workbook_licences[licence_key] = all_workbook_licences.get(licence_key, 0) + 1
    report: list[dict[str, Any]] = []
    websites_checked = 0
    rows_updated = 0
    source_rows_updated = 0
    skipped_non_public_or_ambiguous = 0
    skipped_third_party_ordering_sites = 0

    last_row = min(ws.max_row, args.start_row + args.max_rows - 1)
    for row_num in range(args.start_row, last_row + 1):
        website = normalize_spaces(ws.cell(row_num, headers["Google Website"]).value)
        menu_source = ws.cell(row_num, headers["Menu Source"]).value
        menu_items = ws.cell(row_num, headers["Menu Items"]).value
        menu_summary = ws.cell(row_num, menu_col).value if menu_col else None
        if args.refresh_review:
            licence = normalize_spaces(ws.cell(row_num, licence_col).value)
            licence_key = re.sub(r"[^A-Z0-9]", "", licence.upper())
            if not licence_key or licence_key not in public_licences or all_workbook_licences.get(licence_key) != 1:
                skipped_non_public_or_ambiguous += 1
                continue
            if is_blank(website):
                continue
            if is_third_party_ordering_url(website):
                skipped_third_party_ordering_sites += 1
                continue
        elif not website or not is_blank(menu_source) or not is_blank(menu_items):
            continue
        if not args.refresh_review and is_blank(website):
            continue
        if not args.refresh_review and not is_blank(menu_summary):
            continue
        if websites_checked >= args.max_websites:
            break
        websites_checked += 1

        name = normalize_spaces(ws.cell(row_num, headers["Stall / Restaurant Name"]).value)
        row_report: dict[str, Any] = {"row": row_num, "name": name, "website": website, "refresh_review": args.refresh_review}
        if args.refresh_review:
            row_report["source_class"] = "catalogue_linked_website_candidate"
        if args.refresh_review:
            row_report.update({
                "licence_no": normalize_spaces(ws.cell(row_num, licence_col).value),
                "address": normalize_spaces(ws.cell(row_num, headers["Address"]).value) if "Address" in headers else "",
                "postal_code": normalize_spaces(ws.cell(row_num, headers["Postal Code"]).value) if "Postal Code" in headers else "",
                "place_id": normalize_spaces(ws.cell(row_num, headers["Google Place ID"]).value) if "Google Place ID" in headers else "",
                "google_display_name": normalize_spaces(ws.cell(row_num, headers["Google Display Name"]).value) if "Google Display Name" in headers else "",
                "google_formatted_address": normalize_spaces(ws.cell(row_num, headers["Google Formatted Address"]).value) if "Google Formatted Address" in headers else "",
                "current_menu_source": normalize_spaces(menu_source),
                "current_menu_items": normalize_spaces(menu_items),
                "current_menu_summary": normalize_spaces(menu_summary),
            })
        try:
            website_host = urlparse(website).netloc.lower()
            if urlparse(website).scheme not in {"http", "https"} or not website_host:
                row_report.update({"status": "invalid_website_url"})
                report.append(row_report)
                continue
            sources = find_menu_sources(website, args.timeout)
        except (HTTPError, URLError, TimeoutError, OSError, ValueError) as exc:
            row_report.update({"status": "source_discovery_failed", "error": str(exc)[:160]})
            report.append(row_report)
            continue

        best: MenuExtraction | None = None
        for source in sources:
            extraction = extract_from_source(source, args.timeout)
            if best is None or len(extraction.items) > len(best.items):
                best = extraction
            if len(extraction.items) >= 3:
                break

        if not best:
            row_report.update({"status": "no_source"})
            report.append(row_report)
            continue

        quality_items = quality_menu_items(best.items, name)
        high_confidence = has_low_noise_quality(quality_items)
        image_candidates: list[dict[str, str]] = []
        for image in best.images or []:
            descriptive = f"{image['alt']} {image['title']}".strip()
            descriptive_key = re.sub(r"[^a-z0-9]+", " ", descriptive.lower()).strip()
            matched_dish = next((item for item in quality_items if re.sub(r"[^a-z0-9]+", " ", item.lower()).strip() in descriptive_key), "")
            if matched_dish:
                kind = "dish_image_candidate"
            elif any(term in descriptive_key for term in ("storefront", "store front", "outlet", "restaurant", "interior", "exterior", "dining room", "our cafe", "our café")):
                kind = "venue_image_candidate"
            else:
                kind = "unclassified_official_page_image"
            image_candidates.append({
                **image, "candidate_kind": kind, "associated_menu_item": matched_dish,
                "source_page_url": best.source_url, "rights_status": "permission_required",
                "display_approved": "false",
            })
        row_report.update(
            {
                "status": best.status,
                "source": best.source_url,
                "items_found": len(best.items),
                "quality_items_found": len(quality_items),
                "high_confidence": high_confidence,
                "items": best.items[:8],
                "quality_items": quality_items[:8],
                "image_candidates": image_candidates,
                "error": best.error,
                "observed_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            }
        )
        source_can_be_recorded = best.status == "extracted" and bool(best.source_url)
        if args.apply and args.apply_menu_sources and source_can_be_recorded:
            ws.cell(row_num, headers["Menu Source"]).value = f"online_menu_candidate: {best.source_url}"
            source_rows_updated += 1
        report.append(row_report)

    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "start_row": args.start_row,
                "last_scanned_row": last_row,
                "websites_checked": websites_checked,
                "skipped_non_public_or_ambiguous": skipped_non_public_or_ambiguous,
                "skipped_third_party_ordering_sites": skipped_third_party_ordering_sites,
                "rows_updated": rows_updated,
                "source_rows_updated": source_rows_updated,
                "apply": args.apply,
                "report": report,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
