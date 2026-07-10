#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it in the active Python environment, "
        "or run with the Codex workspace Python runtime used for workbook tooling."
    ) from exc


SOURCE_PREFIX = "category_seed"
REPLACEABLE_SOURCE_PREFIXES = ("", "online_menu_candidate:")


@dataclass(frozen=True)
class CategorySeed:
    label: str
    aliases: tuple[str, ...]
    items: tuple[str, ...]


CATEGORY_SEEDS: tuple[CategorySeed, ...] = (
    CategorySeed(
        label="chicken_rice",
        aliases=(
            "chicken rice",
            "chicken rice stall",
            "fitra chicken rice",
            "hainanese chicken rice",
            "sergeant hainanese chicken rice",
        ),
        items=(
            "Hainanese Chicken Rice",
            "Steamed Chicken Rice",
            "Roasted Chicken Rice",
            "Chicken Drumstick Rice",
            "Chicken Wing Rice",
            "Chicken Cutlet Rice",
            "Chicken Rice Set",
            "Chicken Soup",
        ),
    ),
    CategorySeed(
        label="western_food",
        aliases=(
            "western food",
            "western",
            "western cuisine",
            "western stall",
            "western food stall",
            "popular food",
        ),
        items=(
            "Chicken Chop",
            "Fish & Chips",
            "Grilled Chicken",
            "Chicken Cutlet",
            "Pork Chop",
            "Sirloin Steak",
            "Spaghetti",
            "Fries",
        ),
    ),
    CategorySeed(
        label="yong_tau_foo",
        aliases=("yong tau foo", "yong tau fu", "yong tau foo stall"),
        items=(
            "Yong Tau Foo Soup",
            "Dry Yong Tau Foo",
            "Laksa Yong Tau Foo",
            "Ampang Yong Tau Foo",
            "Bee Hoon",
            "Yellow Noodles",
            "Fried Beancurd",
            "Fish Paste Items",
        ),
    ),
    CategorySeed(
        label="beverages",
        aliases=(
            "beverage",
            "beverages",
            "drink stall",
            "drinks",
            "drinks stall",
            "drinkstall",
            "hot and cold",
            "hot & cold",
            "hot drinks and snacks",
            "hot drinks & snacks",
        ),
        items=(
            "Kopi",
            "Teh",
            "Iced Milo",
            "Lime Juice",
            "Sugar Cane Juice",
            "Soft Drinks",
            "Bottled Water",
            "Tea",
        ),
    ),
    CategorySeed(
        label="mixed_vegetable_rice",
        aliases=(
            "mixed vegetable rice",
            "mixed vegetables rice",
            "mixed veg rice",
            "mixed vege rice",
            "mixed rice",
            "mix rice",
            "mix vegetable rice",
            "mix vegetables rice",
            "mix veg rice",
            "mixed rice stall",
            "economic rice",
            "economical rice",
            "econ rice",
            "chinese mixed rice",
            "curry rice",
            "mvr",
        ),
        items=(
            "Cai Fan",
            "Fried Egg",
            "Curry Chicken",
            "Sweet and Sour Pork",
            "Braised Tofu",
            "Stir-Fried Vegetables",
            "Fried Fish",
            "Pork Chop",
        ),
    ),
    CategorySeed(
        label="nasi_padang",
        aliases=("nasi padang",),
        items=(
            "Beef Rendang",
            "Ayam Masak Merah",
            "Sayur Lodeh",
            "Sambal Goreng",
            "Curry Chicken",
            "Nasi Padang Set",
            "Begedil",
            "Sambal Egg",
        ),
    ),
    CategorySeed(
        label="indian_food",
        aliases=("indian food", "indian cuisine", "indian stall", "indian muslim food", "indian muslim stall"),
        items=(
            "Roti Prata",
            "Murtabak",
            "Nasi Briyani",
            "Mee Goreng",
            "Nasi Goreng",
            "Thosai",
            "Chicken Curry",
            "Teh Tarik",
        ),
    ),
    CategorySeed(
        label="seafood",
        aliases=("seafood", "seafood stall"),
        items=(
            "Chilli Crab",
            "Black Pepper Crab",
            "Cereal Prawn",
            "Sambal Stingray",
            "Lala",
            "Fish Head Curry",
            "Salted Egg Prawn",
            "Seafood Fried Rice",
        ),
    ),
    CategorySeed(
        label="ban_mian",
        aliases=("ban mian", "ban mee", "ban miam"),
        items=(
            "Soup Ban Mian",
            "Dry Ban Mian",
            "Mee Hoon Kueh",
            "You Mian",
            "U-Mian",
            "Fish Soup Ban Mian",
            "Pork Ban Mian",
            "Fried Dumpling",
        ),
    ),
    CategorySeed(
        label="fishball_noodle",
        aliases=("fishball noodle", "fishball noodles", "fish ball noodle", "fish ball noodles"),
        items=(
            "Fishball Noodle Dry",
            "Fishball Noodle Soup",
            "Mee Pok",
            "Mee Kia",
            "Fish Cake",
            "Minced Meat Noodle",
            "Laksa",
            "Fried Fishball",
        ),
    ),
    CategorySeed(
        label="fish_soup",
        aliases=("fish soup", "sliced fish soup", "teochew fish soup", "chai s original sliced fish soup"),
        items=(
            "Sliced Fish Soup",
            "Fried Fish Soup",
            "Fish Head Soup",
            "Seafood Soup",
            "Fish Bee Hoon Soup",
            "Tom Yam Fish Soup",
            "Rice",
            "Bittergourd Fish Soup",
        ),
    ),
    CategorySeed(
        label="roasted_delights",
        aliases=(
            "roasted delight",
            "roasted delights",
            "roasted food",
            "roasted duck",
            "roasted master",
            "roast delight",
        ),
        items=(
            "Roast Duck Rice",
            "Roast Pork Rice",
            "Char Siew Rice",
            "Soy Sauce Chicken Rice",
            "Roast Meat Noodles",
            "Roast Duck Noodles",
            "Three Treasure Rice",
            "Wanton Soup",
        ),
    ),
    CategorySeed(
        label="korean_food",
        aliases=("korean cuisine", "korean food", "korean", "japanese and korean cuisine", "japanese & korean cuisine"),
        items=(
            "Bibimbap",
            "Kimchi Fried Rice",
            "Tteokbokki",
            "Korean Fried Chicken",
            "Bulgogi",
            "Army Stew",
            "Kimchi Soup",
            "Ramyeon",
        ),
    ),
    CategorySeed(
        label="vegetarian",
        aliases=("vegetarian", "vegetarian food"),
        items=(
            "Vegetarian Bee Hoon",
            "Vegetarian Rice",
            "Thunder Tea Rice",
            "Mushroom Noodles",
            "Mock Meat Rice",
            "Stir-Fried Vegetables",
            "Vegetarian Hor Fun",
            "Tofu Rice",
        ),
    ),
    CategorySeed(
        label="japanese_food",
        aliases=("japanese cuisine", "japanese food", "japanese"),
        items=(
            "Chicken Katsu Curry",
            "Teriyaki Chicken Don",
            "Salmon Don",
            "Udon",
            "Ramen",
            "Bento",
            "Sushi",
            "Tempura",
        ),
    ),
    CategorySeed(
        label="mala",
        aliases=("mala xiang guo", "mala hotpot", "mala hot pot", "ma la xiang guo"),
        items=(
            "Mala Xiang Guo",
            "Mala Hotpot",
            "Mala Tang",
            "Spicy Dry Pot",
            "Mala Noodles",
            "Mala Chicken",
            "Mala Vegetables",
            "Mala Soup",
        ),
    ),
    CategorySeed(
        label="malay_food",
        aliases=("malay food", "malay cuisine", "malay stall", "muslim food", "muslim stall"),
        items=(
            "Nasi Lemak",
            "Mee Siam",
            "Mee Rebus",
            "Lontong",
            "Soto Ayam",
            "Mee Soto",
            "Nasi Goreng",
            "Ayam Penyet",
        ),
    ),
    CategorySeed(
        label="dim_sum",
        aliases=("dim sum",),
        items=(
            "Siew Mai",
            "Har Gow",
            "Char Siew Bao",
            "Lotus Leaf Rice",
            "Chee Cheong Fun",
            "Egg Tart",
            "Xiao Long Bao",
            "Prawn Dumpling",
        ),
    ),
    CategorySeed(
        label="herbal_soup",
        aliases=("herbal soup",),
        items=(
            "Herbal Chicken Soup",
            "Lotus Root Pork Rib Soup",
            "Old Cucumber Soup",
            "Black Chicken Soup",
            "Peanut Pork Rib Soup",
            "Watercress Soup",
            "Rice",
            "Steamed Egg",
        ),
    ),
    CategorySeed(
        label="thai_food",
        aliases=("thai food", "thai cuisine"),
        items=(
            "Pad Thai",
            "Green Curry Chicken",
            "Tom Yum Soup",
            "Basil Chicken Rice",
            "Pineapple Fried Rice",
            "Mango Sticky Rice",
            "Thai Omelette Rice",
            "Papaya Salad",
        ),
    ),
    CategorySeed(
        label="bak_kut_teh",
        aliases=("bak kut teh",),
        items=(
            "Pork Rib Soup",
            "Pig's Trotter",
            "Braised Tau Pok",
            "You Tiao",
            "Rice",
            "Preserved Vegetables",
            "Dough Fritters",
            "Salted Vegetables",
        ),
    ),
    CategorySeed(
        label="ayam_penyet",
        aliases=("ayam penyet",),
        items=(
            "Ayam Penyet",
            "Ayam Bakar",
            "Ikan Penyet",
            "Tahu Tempe",
            "Nasi Ayam Penyet",
            "Sambal",
            "Fried Egg",
            "Beef Rendang",
        ),
    ),
    CategorySeed(
        label="roti_prata",
        aliases=("roti prata",),
        items=(
            "Plain Prata",
            "Egg Prata",
            "Cheese Prata",
            "Mutton Murtabak",
            "Chicken Curry",
            "Fish Curry",
            "Teh Tarik",
            "Thosai",
        ),
    ),
    CategorySeed(
        label="wanton_noodle",
        aliases=("wanton noodle", "wanton noodles", "noodle house", "noodle stall"),
        items=(
            "Wanton Noodle Dry",
            "Wanton Noodle Soup",
            "Char Siew Noodle",
            "Fried Wanton",
            "Dumpling Noodle",
            "Roast Duck Noodle",
            "Chicken Feet Noodle",
            "Vegetable Noodle",
        ),
    ),
    CategorySeed(
        label="popiah",
        aliases=("popiah",),
        items=(
            "Popiah",
            "Kueh Pie Tee",
            "Turnip Roll",
            "Rojak",
            "Fresh Spring Roll",
            "Laksa",
        ),
    ),
    CategorySeed(
        label="bubble_tea",
        aliases=("bubble tea",),
        items=(
            "Pearl Milk Tea",
            "Brown Sugar Fresh Milk",
            "Classic Milk Tea",
            "Fruit Tea",
            "Green Tea",
            "Oolong Tea",
            "Cheese Tea",
            "Taro Milk Tea",
        ),
    ),
    CategorySeed(
        label="fruit_juice",
        aliases=("fruit juice", "fruit box"),
        items=(
            "Watermelon Juice",
            "Apple Juice",
            "Orange Juice",
            "Sugar Cane Juice",
            "Avocado Juice",
            "Mango Juice",
            "Carrot Juice",
            "Mixed Fruit Juice",
        ),
    ),
    CategorySeed(
        label="snacks",
        aliases=("snacks",),
        items=(
            "Curry Puff",
            "Spring Roll",
            "Fishball",
            "Fried Chicken Wing",
            "Fries",
            "Nuggets",
            "Samosa",
            "Fried Dumpling",
        ),
    ),
    CategorySeed(
        label="claypot_rice",
        aliases=("claypot rice", "claypot delight"),
        items=(
            "Chicken Claypot Rice",
            "Salted Fish Claypot Rice",
            "Chinese Sausage Claypot Rice",
            "Mushroom Claypot Rice",
            "Claypot Sesame Oil Chicken",
            "Claypot Tofu",
        ),
    ),
    CategorySeed(
        label="local_delights",
        aliases=("local delights",),
        items=(
            "Laksa",
            "Chicken Rice",
            "Nasi Lemak",
            "Mee Siam",
            "Curry Chicken",
            "Kaya Toast",
            "Mee Rebus",
            "Fried Bee Hoon",
        ),
    ),
    CategorySeed(
        label="noodles",
        aliases=("noodle", "noodles", "chinese noodle", "chinese noodles", "chinese noodle stall"),
        items=(
            "Fishball Noodle",
            "Wanton Noodle",
            "Ban Mian",
            "Lor Mee",
            "Prawn Noodle",
            "Minced Meat Noodle",
            "Mee Pok",
            "Mee Kia",
        ),
    ),
    CategorySeed(
        label="ban_mian_fish_soup",
        aliases=("ban mian fish soup", "ban mian / fish soup", "ban mian and fish soup", "ban mian & fish soup", "ban mian/fish soup"),
        items=(
            "Soup Ban Mian",
            "Dry Ban Mian",
            "Sliced Fish Soup",
            "Fried Fish Soup",
            "Mee Hoon Kueh",
            "You Mian",
            "Rice",
            "Fried Dumpling",
        ),
    ),
    CategorySeed(
        label="nasi_lemak",
        aliases=("nasi lemak", "superman nasi lemak"),
        items=(
            "Nasi Lemak",
            "Chicken Wing Nasi Lemak",
            "Fried Fish Nasi Lemak",
            "Otah",
            "Fried Egg",
            "Ikan Bilis",
            "Sambal",
            "Curry Chicken",
        ),
    ),
    CategorySeed(
        label="chinese_cooked_food",
        aliases=(
            "chinese food",
            "chinese cuisine",
            "chinese rice",
            "chinese cooked food",
            "chinese wok",
            "mini wok",
            "zi char",
        ),
        items=(
            "Sweet and Sour Pork Rice",
            "Ginger Onion Chicken Rice",
            "Fried Rice",
            "Hor Fun",
            "Mui Fan",
            "Claypot Tofu Rice",
            "Sambal Fried Rice",
            "Chinese Soup",
        ),
    ),
    CategorySeed(
        label="economic_bee_hoon",
        aliases=("economic bee hoon", "bee hoon"),
        items=(
            "Economic Bee Hoon",
            "Fried Bee Hoon",
            "Fried Noodles",
            "Fried Egg",
            "Chicken Wing",
            "Fish Cake",
            "Luncheon Meat",
            "Spring Roll",
        ),
    ),
    CategorySeed(
        label="prawn_noodle",
        aliases=("prawn noodle", "prawn noodles"),
        items=(
            "Prawn Noodle Soup",
            "Dry Prawn Noodle",
            "Pork Rib Prawn Noodle",
            "Big Prawn Noodle",
            "Prawn Mee",
            "Pig Tail Soup",
            "Fish Cake",
            "Bee Hoon",
        ),
    ),
    CategorySeed(
        label="porridge",
        aliases=("porridge", "teochew porridge", "ban heng teochew porridge"),
        items=(
            "Plain Porridge",
            "Minced Pork Porridge",
            "Chicken Porridge",
            "Fish Porridge",
            "Century Egg Porridge",
            "Braised Peanuts",
            "Salted Egg",
            "Preserved Vegetables",
        ),
    ),
    CategorySeed(
        label="duck_rice",
        aliases=("duck rice",),
        items=(
            "Braised Duck Rice",
            "Roast Duck Rice",
            "Duck Noodle",
            "Duck Kway Teow",
            "Braised Egg",
            "Tau Kwa",
            "Duck Soup",
            "Yam Rice",
        ),
    ),
    CategorySeed(
        label="dessert",
        aliases=(
            "dessert",
            "thejellyhearts",
            "llao llao",
            "gelatissimo",
            "cold stone creamery",
            "baskin robbins",
            "chocolate origin",
            "four seasons durians",
            "meruto cheesecake",
        ),
        items=(
            "Gelato",
            "Ice Cream",
            "Cheesecake",
            "Chocolate Cake",
            "Durian Dessert",
            "Yogurt",
            "Jelly Cheesecake",
            "Waffles",
        ),
    ),
    CategorySeed(
        label="bakery",
        aliases=(
            "bakery",
            "bake inc",
            "bread fresh",
            "breadwerks",
            "cake history",
            "butter studio",
            "so good bakery",
            "patisserie g",
            "patisserie glace",
            "homies bakery",
            "ah mah traditional egg sponge cake",
        ),
        items=(
            "Bread",
            "Croissant",
            "Butter Roll",
            "Sponge Cake",
            "Chocolate Cake",
            "Cheesecake",
            "Pastries",
            "Muffin",
        ),
    ),
    CategorySeed(
        label="cafe",
        aliases=(
            "cafe",
            "good news cafe",
            "hill street coffee shop",
            "st 82 coffee house",
            "coffee and toast",
            "local coffee people",
            "nanyang old coffee",
            "new king cafe",
            "new king cafe pte ltd",
            "nature cafe",
            "nested cafe",
        ),
        items=(
            "Kaya Toast",
            "Soft-Boiled Eggs",
            "Kopi",
            "Teh",
            "Nasi Lemak",
            "Mee Siam",
            "Curry Puff",
            "Sandwich",
        ),
    ),
    CategorySeed(
        label="food_court_operator",
        aliases=(
            "koufu",
            "koufu pte ltd",
            "kopitiam investment pte ltd",
            "food republic",
            "food republic pte ltd",
            "food loft",
            "food more",
            "food canopy",
            "banquet",
            "banquet kitchen",
            "my kampung",
            "s11",
            "broadway",
            "broadway food centre holdings",
        ),
        items=(
            "Chicken Rice",
            "Yong Tau Foo",
            "Ban Mian",
            "Western Food",
            "Mixed Vegetable Rice",
            "Korean Food",
            "Japanese Food",
            "Drinks",
        ),
    ),
    CategorySeed(
        label="supermarket_convenience",
        aliases=(
            "cold storage",
            "cold storage singapore 1983",
            "ntuc",
            "ntuc fairprice",
            "ntuc fairprice co operative",
            "fairprice",
            "fairprice finest",
            "fairprice shop",
            "giant",
            "sheng siong",
            "sheng siong supermarket",
            "hao mart",
            "prime supermarket",
            "shell select",
            "spc",
            "u stars supermarket",
            "dondondonki",
            "caltex",
        ),
        items=(
            "Ready Meals",
            "Sandwiches",
            "Sushi",
            "Salads",
            "Bakery Items",
            "Hot Snacks",
            "Beverages",
            "Desserts",
        ),
    ),
    CategorySeed(
        label="canteen_catering",
        aliases=(
            "staff canteen",
            "staff cafeteria",
            "canteen",
            "school",
            "school canteen",
            "ntwu canteen",
            "in house catering",
            "iss catering services",
            "compass group singapore",
            "sodexo singapore",
            "sports catering services",
            "pro 3 institutional catering",
            "the wok people",
            "ntuc foodfare",
            "ntuc foodfare co operative",
        ),
        items=(
            "Chicken Rice",
            "Mixed Vegetable Rice",
            "Nasi Lemak",
            "Fishball Noodle",
            "Western Food",
            "Vegetarian Rice",
            "Drinks",
            "Snacks",
        ),
    ),
    CategorySeed(
        label="ramen_japanese_noodles",
        aliases=("ajisen ramen", "takagi ramen", "kajiken"),
        items=(
            "Tonkotsu Ramen",
            "Shoyu Ramen",
            "Miso Ramen",
            "Mazemen",
            "Gyoza",
            "Karaage",
            "Chashu Rice",
            "Egg Ramen",
        ),
    ),
    CategorySeed(
        label="sushi",
        aliases=("itacho sushi", "umisushi", "samurice"),
        items=(
            "Salmon Sushi",
            "Tuna Sushi",
            "California Roll",
            "Tamago Sushi",
            "Teriyaki Chicken Bento",
            "Salmon Don",
            "Onigiri",
            "Miso Soup",
        ),
    ),
    CategorySeed(
        label="mexican",
        aliases=("guzman y gomez",),
        items=(
            "Burrito",
            "Burrito Bowl",
            "Tacos",
            "Nachos",
            "Quesadilla",
            "Churros",
            "Guacamole",
            "Mexican Fries",
        ),
    ),
    CategorySeed(
        label="pizza",
        aliases=("canadian pizza", "the wicked garlic"),
        items=(
            "Hawaiian Pizza",
            "Pepperoni Pizza",
            "Margherita Pizza",
            "Pasta",
            "Garlic Bread",
            "Chicken Wings",
            "Lasagna",
            "Mushroom Soup",
        ),
    ),
    CategorySeed(
        label="tea_herbal_drinks",
        aliases=("herbal tea", "hockhua herbal tea", "oriental herbal tea", "share tea", "super tea", "sweet talk", "cha tra mue"),
        items=(
            "Herbal Tea",
            "Luo Han Guo",
            "Chrysanthemum Tea",
            "Pearl Milk Tea",
            "Thai Milk Tea",
            "Fruit Tea",
            "Grass Jelly Drink",
            "Winter Melon Tea",
        ),
    ),
)


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return clean(value) in {"", "-", "—", "N/A", "NA", "None", "null"}


def normalize(value: Any) -> str:
    text = clean(value).lower().replace("&", " and ")
    text = re.sub(r"\b(?:pte\.?\s*ltd\.?|private\s+limited|limited|ltd\.?|llp|inc\.?)\b", " ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def load_headers(sheet) -> dict[str, int]:
    return {clean(cell.value): index for index, cell in enumerate(sheet[1], start=1)}


def source_is_replaceable(value: Any) -> bool:
    source = clean(value)
    return source == "" or source.startswith(REPLACEABLE_SOURCE_PREFIXES[1])


def seed_for_name(name: str) -> CategorySeed | None:
    normalized_name = normalize(name)
    matches = [
        seed
        for seed in CATEGORY_SEEDS
        if normalized_name in {normalize(alias) for alias in seed.aliases}
    ]
    return matches[0] if len(matches) == 1 else None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100000)
    parser.add_argument("--max-updates", type=int, default=1000)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = load_headers(ws)
    required = ["Stall / Restaurant Name", "Signature Dish", "Menu", "Menu Items", "Menu Source"]
    missing = [field for field in required if field not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    updates: list[dict[str, Any]] = []
    last_row = min(ws.max_row, args.start_row + args.max_rows - 1)
    for row_num in range(args.start_row, last_row + 1):
        if len(updates) >= args.max_updates:
            break
        if not is_blank(ws.cell(row_num, headers["Menu"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu Items"]).value):
            continue
        if not source_is_replaceable(ws.cell(row_num, headers["Menu Source"]).value):
            continue
        seed = seed_for_name(clean(ws.cell(row_num, headers["Stall / Restaurant Name"]).value))
        if not seed:
            continue

        item_text = ", ".join(seed.items)
        update = {
            "row": row_num,
            "name": ws.cell(row_num, headers["Stall / Restaurant Name"]).value,
            "source": f"{SOURCE_PREFIX}: {seed.label}",
            "items": list(seed.items),
        }
        updates.append(update)
        if args.apply:
            if is_blank(ws.cell(row_num, headers["Signature Dish"]).value):
                ws.cell(row_num, headers["Signature Dish"]).value = seed.items[0]
            ws.cell(row_num, headers["Menu"]).value = item_text
            ws.cell(row_num, headers["Menu Items"]).value = item_text
            ws.cell(row_num, headers["Menu Source"]).value = f"{SOURCE_PREFIX}: {seed.label}"

    backup_path = None
    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "apply": args.apply,
                "backup": str(backup_path) if backup_path else None,
                "updated": len(updates),
                "updates": updates[:100],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
