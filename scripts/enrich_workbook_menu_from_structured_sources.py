#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


LEGAL_SUFFIX_RE = re.compile(
    r"\b(?:pte\.?\s*ltd\.?|private\s+limited|limited|ltd\.?|llp|plc|inc\.?|co\.?)\b",
    re.IGNORECASE,
)
GENERIC_ITEM_RE = re.compile(
    r"\b(?:"
    r"affordable|array|authentic|best|casual|choice|choices|classic|classics|concept|"
    r"culture|culinary|customer|customers|delicious|dining|experience|flavour|flavours|"
    r"fresh|freshly|heart|hearty|ingredients|lifestyle|major|menu|natural|offering|"
    r"options|place|quality|range|selection|signature steaks|specialities|specialty|"
    r"taste|tradition|variety|wide|well[- ]brewed"
    r")\b",
    re.IGNORECASE,
)
FOOD_WORDS = {
    "ayam",
    "banh",
    "bao",
    "bbq",
    "beef",
    "bento",
    "biryani",
    "bread",
    "broth",
    "bun",
    "burger",
    "cake",
    "carbonara",
    "cheese",
    "chicken",
    "chirashi",
    "chocolate",
    "coffee",
    "congee",
    "cookie",
    "cream",
    "croissant",
    "curry",
    "dim sum",
    "don",
    "donut",
    "duck",
    "dumpling",
    "egg",
    "fish",
    "fries",
    "gyoza",
    "hor fun",
    "ice cream",
    "kaya",
    "kebab",
    "laksa",
    "la mian",
    "mala",
    "mee",
    "muffin",
    "mushroom",
    "noodle",
    "oyster",
    "pasta",
    "pizza",
    "pork",
    "prawn",
    "puff",
    "ramen",
    "rice",
    "robata",
    "robata-yaki",
    "robatayaki",
    "roll",
    "salad",
    "salmon",
    "sandwich",
    "satay",
    "seafood",
    "soup",
    "sushi",
    "taco",
    "tart",
    "tea",
    "toast",
    "tofu",
    "tonkatsu",
    "waffle",
    "wing",
    "yakiniku",
}
TRIGGER_PATTERNS = [
    re.compile(
        r"\b(?:including|such as|favorites? such as|favourites? such as|menu features|"
        r"menu includes|features|serving|serves|offers|signature dishes? (?:like|include|includes)|"
        r"signature items? (?:like|include|includes)|signature offerings? (?:like|include|includes)|"
        r"best[- ]selling dish(?:es)?(?:,| include| includes)?|try(?: out)?(?: the)?)\s+([^.;]+)",
        re.IGNORECASE,
    ),
]
STOP_RE = re.compile(
    r"\b(?:all|and more|each|for yourself|made|perfect|prepared|served|that|to match|where|which|while|with its)\b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class SourceCandidate:
    name: str
    normalized_name: str
    mall: str
    normalized_mall: str
    unit: str
    source_url: str
    items: tuple[str, ...]
    confidence: str


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return clean(value) in {"", "-", "—", "N/A", "NA", "None", "null"}


def normalize_name(value: Any) -> str:
    text = clean(value).lower().replace("&", " and ")
    text = LEGAL_SUFFIX_RE.sub(" ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return clean(text)


def normalize_unit(value: Any) -> str:
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def strip_item_prefix(value: str) -> str:
    text = clean(value).strip(" ,;:-")
    text = re.sub(
        r"^(?:all[- ]time|best[- ]selling|house|popular|signature)?\s*"
        r"(?:classics?|dishes|favorites?|favourites?|items|offerings|specials?)\s+"
        r"(?:such as|like|include|includes)?\s+",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"^(?:the|their|our|a|an)\s+", "", text, flags=re.IGNORECASE)
    return clean(text).strip(" ,;:-")


def contains_food_word(value: str) -> bool:
    lowered = value.lower()
    return any(word in lowered for word in FOOD_WORDS)


def valid_item(value: str) -> bool:
    text = strip_item_prefix(value)
    if not (3 <= len(text) <= 64):
        return False
    first_alpha = re.search(r"[A-Za-z]", text)
    if not first_alpha or not first_alpha.group(0).isupper():
        return False
    if GENERIC_ITEM_RE.search(text) and not contains_food_word(text):
        return False
    if GENERIC_ITEM_RE.search(text):
        return False
    if re.search(r"\b(?:customer|menu|restaurant|store|outlet|venue)\b", text, re.IGNORECASE):
        return False
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", text)
    if len(words) < 2 or len(words) > 8:
        return False
    return contains_food_word(text)


def split_item_fragment(fragment: str) -> list[str]:
    fragment = STOP_RE.split(fragment, maxsplit=1)[0]
    fragment = re.sub(r"\([^)]*\)", " ", fragment)
    parts = re.split(r",|;|\bor\b|\band\b(?!\s*(?:chips|cheese|egg|toast)\b)", fragment, flags=re.IGNORECASE)
    return [strip_item_prefix(part) for part in parts]


def extract_items(description: str) -> tuple[str, ...]:
    items: list[str] = []
    seen: set[str] = set()
    for pattern in TRIGGER_PATTERNS:
        for match in pattern.finditer(description):
            for item in split_item_fragment(match.group(1)):
                key = normalize_name(item)
                if key and key not in seen and valid_item(item):
                    seen.add(key)
                    items.append(item)
    for match in re.finditer(
        r"\b([A-Z][A-Za-z&'/-]+(?:\s+[A-Z][A-Za-z&'/-]+){0,5}\s+"
        r"(?:Pasta|Pizza|Ramen|Rice|Burger|Toast|Curry|Chicken|Fish|Donut|Tart|Sushi|Bento|Soup))\b",
        description,
    ):
        item = strip_item_prefix(match.group(1))
        key = normalize_name(item)
        if key and key not in seen and valid_item(item):
            seen.add(key)
            items.append(item)
    return tuple(items[:10])


def source_url_for_record(record: dict[str, Any], fallback_path: Path) -> str:
    for key in ("page_url", "website", "url", "source_url"):
        value = clean(record.get(key))
        if value:
            return value
    return str(fallback_path)


def load_source_candidates(source_dir: Path) -> list[SourceCandidate]:
    candidates: list[SourceCandidate] = []
    for path in sorted(source_dir.glob("*.json")):
        with path.open(encoding="utf-8") as handle:
            payload = json.load(handle)
        records = payload if isinstance(payload, list) else payload.get("stores") or payload.get("items") or []
        if not isinstance(records, list):
            continue
        for record in records:
            if not isinstance(record, dict):
                continue
            description = clean(record.get("description") or record.get("shortdescription"))
            raw_items = extract_items(description)
            source_name_key = normalize_name(name := clean(record.get("name") or record.get("jcr:title")))
            items = tuple(
                item
                for item in raw_items
                if normalize_name(item) != source_name_key and source_name_key not in normalize_name(item)
            )
            if not items:
                continue
            mall = clean(record.get("mall") or record.get("building") or record.get("location_note"))
            if not name:
                continue
            candidates.append(
                SourceCandidate(
                    name=name,
                    normalized_name=normalize_name(name),
                    mall=mall,
                    normalized_mall=normalize_name(mall),
                    unit=normalize_unit(record.get("unit_number")),
                    source_url=source_url_for_record(record, path),
                    items=items,
                    confidence="source_description",
                )
            )
    return candidates


def build_candidate_indexes(
    candidates: list[SourceCandidate],
) -> tuple[dict[tuple[str, str], list[SourceCandidate]], dict[str, SourceCandidate]]:
    by_name_mall: dict[tuple[str, str], list[SourceCandidate]] = defaultdict(list)
    by_name: dict[str, list[SourceCandidate]] = defaultdict(list)
    for candidate in candidates:
        if candidate.normalized_name and candidate.normalized_mall:
            by_name_mall[(candidate.normalized_name, candidate.normalized_mall)].append(candidate)
        if candidate.normalized_name:
            by_name[candidate.normalized_name].append(candidate)

    unique_name: dict[str, SourceCandidate] = {}
    for normalized_name, group in by_name.items():
        item_sets = {candidate.items for candidate in group}
        if len(group) == 1 or len(item_sets) == 1:
            unique_name[normalized_name] = group[0]
    return by_name_mall, unique_name


def find_candidate(
    row_name: str,
    row_building: str,
    row_unit: str,
    by_name_mall: dict[tuple[str, str], list[SourceCandidate]],
    unique_name: dict[str, SourceCandidate],
) -> SourceCandidate | None:
    normalized_name = normalize_name(row_name)
    normalized_building = normalize_name(row_building)
    normalized_unit = normalize_unit(row_unit)
    if not normalized_name:
        return None

    if normalized_building:
        for (candidate_name, candidate_mall), group in by_name_mall.items():
            if candidate_name != normalized_name:
                continue
            if candidate_mall and (
                candidate_mall in normalized_building or normalized_building in candidate_mall
            ):
                if normalized_unit:
                    unit_matches = [candidate for candidate in group if candidate.unit == normalized_unit]
                    if unit_matches:
                        return unit_matches[0]
                return group[0]

    return unique_name.get(normalized_name)


def load_headers(sheet) -> dict[str, int]:
    return {clean(cell.value): index for index, cell in enumerate(sheet[1], start=1)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--source-dir", default="data/malls")
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100000)
    parser.add_argument("--max-updates", type=int, default=500)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    source_dir = Path(args.source_dir)
    candidates = load_source_candidates(source_dir)
    by_name_mall, unique_name = build_candidate_indexes(candidates)

    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = load_headers(ws)
    required = [
        "Stall / Restaurant Name",
        "Menu",
        "Menu Items",
        "Menu Source",
        "Signature Dish",
        "Building Name",
        "Hawker Centre / Building",
        "Unit No.",
    ]
    missing = [field for field in required if field not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    updates: list[dict[str, Any]] = []
    last_row = min(ws.max_row, args.start_row + args.max_rows - 1)
    for row_num in range(args.start_row, last_row + 1):
        if len(updates) >= args.max_updates:
            break
        if not is_blank(ws.cell(row_num, headers["Menu"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu Items"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu Source"]).value):
            continue

        row_name = ws.cell(row_num, headers["Stall / Restaurant Name"]).value
        row_building = ws.cell(row_num, headers["Building Name"]).value or ws.cell(
            row_num, headers["Hawker Centre / Building"]
        ).value
        row_unit = ws.cell(row_num, headers["Unit No."]).value
        candidate = find_candidate(str(row_name or ""), str(row_building or ""), str(row_unit or ""), by_name_mall, unique_name)
        if not candidate:
            continue

        menu_text = ", ".join(candidate.items)
        update = {
            "row": row_num,
            "name": clean(row_name),
            "matched_source_name": candidate.name,
            "source": candidate.source_url,
            "items": list(candidate.items),
            "confidence": candidate.confidence,
        }
        updates.append(update)
        if args.apply:
            ws.cell(row_num, headers["Menu"]).value = menu_text
            ws.cell(row_num, headers["Menu Items"]).value = menu_text
            ws.cell(row_num, headers["Menu Source"]).value = f"{candidate.confidence}: {candidate.source_url}"
            if is_blank(ws.cell(row_num, headers["Signature Dish"]).value):
                ws.cell(row_num, headers["Signature Dish"]).value = candidate.items[0]

    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "source_candidates": len(candidates),
                "start_row": args.start_row,
                "last_scanned_row": last_row,
                "updates": len(updates),
                "apply": args.apply,
                "sample": updates[:30],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
