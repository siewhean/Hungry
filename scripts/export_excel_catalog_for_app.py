#!/usr/bin/env python3
"""Export the Excel master workbook into a compact app runtime catalog."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dish_taxonomy import lookup_dish_taxonomy
from validate_excel_workbook import DEFAULT_WORKBOOK, clean


DEFAULT_OUTPUT = Path("src/lib/services/excel/excelCatalog.generated.json")
MAX_MENU_ITEMS_PER_VENUE = 0
DEFAULT_COORDINATE = {"latitude": 1.3521, "longitude": 103.8198}
BAD_IMAGE_HOSTS = (
    "images.unsplash.com",
    "source.unsplash.com",
    "picsum.photos",
    "wikimedia.org",
    "food.grab.com/static/images/logo-grabfood",
    "food.grab.com/static/images/quick-add",
    "plus-white.svg",
    "placeholder",
)

BUILDING_COLUMNS = [
    "id",
    "name",
    "type",
    "latitude",
    "longitude",
    "floors",
    "venueCount",
    "isClosedToday",
    "heroDishId",
    "planningArea",
    "imageUrl",
    "logoUrl",
]

VENUE_COLUMNS = [
    "id",
    "buildingId",
    "name",
    "unitNumber",
    "floor",
    "isOpen",
    "topDishId",
    "operatingHours",
    "planningArea",
    "imageUrl",
    "logoUrl",
    "buildingName",
    "address",
    "postalCode",
    "latitude",
    "longitude",
    "rating",
    "reviewCount",
    "cuisineTags",
    "dietaryFlags",
]

DISH_COLUMNS = [
    "id",
    "venueId",
    "buildingId",
    "name",
    "foodTypeTag",
    "imageUrl",
    "price",
    "priceRange",
    "tasteTags",
    "tasteVector",
    "dietaryFlags",
    "mustOrder",
]


def slugify(value: Any, fallback: str) -> str:
    text = clean(value)
    if not text:
        return fallback
    slug = re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")
    return slug[:96] or fallback


def row_value(row: dict[str, Any], key: str) -> Any:
    value = row.get(key)
    return None if value == "" else value


def text_value(row: dict[str, Any], key: str) -> str:
    value = clean(row_value(row, key))
    return "" if value is None else str(value).strip()


def number_value(row: dict[str, Any], key: str) -> float | None:
    value = row_value(row, key)
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number == number else None


def bool_value(row: dict[str, Any], key: str) -> bool:
    value = row_value(row, key)
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def numeric_value(value: Any) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number == number else None


def good_url(value: Any, *, allow_wikimedia: bool = False) -> str:
    url = text_value({"url": value}, "url")
    if not url or "," in url or not re.match(r"^https?://", url, flags=re.I):
        return ""
    lowered = url.lower()
    bad_hosts = BAD_IMAGE_HOSTS if not allow_wikimedia else BAD_IMAGE_HOSTS[:3]
    if any(host in lowered for host in bad_hosts):
        return ""
    return url


def title_case_area(value: str) -> str:
    if not value:
        return "Singapore"
    return " ".join(part[:1].upper() + part[1:].lower() for part in value.split())


def infer_building_type(row: dict[str, Any]) -> str:
    source = " ".join(
        [
            text_value(row, "Stall Type"),
            text_value(row, "Google Primary Type"),
            text_value(row, "Building Name"),
            text_value(row, "Hawker Centre / Building"),
        ]
    ).lower()
    if "hawker" in source:
        return "hawker_centre"
    if "food court" in source or "kopitiam" in source:
        return "food_court"
    if "cafe" in source:
        return "cafe"
    if "bar" in source or "bistro" in source:
        return "bar_bistro"
    if "mall" in source or "shopping" in source:
        return "mall"
    if "restaurant" in source:
        return "restaurant"
    if "standalone" in source:
        return "standalone"
    return "other"


def split_tags(value: str) -> list[str]:
    if not value:
        return []
    stripped = value.strip()
    if stripped.startswith("[") and stripped.endswith("]"):
        try:
            parsed = json.loads(stripped)
        except json.JSONDecodeError:
            parsed = None
        if isinstance(parsed, list):
            raw = [str(item) for item in parsed]
        else:
            raw = re.split(r"[,;/|]+", value)
    else:
        raw = re.split(r"[,;/|]+", value)
    seen: set[str] = set()
    tags: list[str] = []
    for item in raw:
        tag = item.strip()
        if not tag:
            continue
        key = tag.lower()
        if key in seen:
            continue
        seen.add(key)
        tags.append(tag)
    return tags


DIETARY_MAP = {
    "halal": "halal",
    "vegetarian": "vegetarian",
    "vegan": "vegan",
    "pork free": "pork_free",
    "pork_free": "pork_free",
    "nut free": "nut_free",
    "nut_free": "nut_free",
    "gluten free": "gluten_free",
    "gluten_free": "gluten_free",
}


def dietary_flags(value: str) -> list[str]:
    flags: list[str] = []
    lowered = value.lower().replace("-", " ")
    for needle, flag in DIETARY_MAP.items():
        if needle in lowered and flag not in flags:
            flags.append(flag)
    return flags


def infer_taste_tags(cuisine: str, dish_name: str) -> list[str]:
    taxonomy_entry = lookup_dish_taxonomy(dish_name)
    if taxonomy_entry:
        return list(taxonomy_entry.taste_tags)
    text = f"{cuisine} {dish_name}".lower()
    tags: list[str] = []
    for needle, tag in [
        ("chicken rice", "Comforting"),
        ("rice", "Hearty"),
        ("noodle", "Slurpable"),
        ("soup", "Comforting"),
        ("curry", "Rich"),
        ("laksa", "Lemak"),
        ("mala", "Spicy"),
        ("spicy", "Spicy"),
        ("coffee", "Aromatic"),
        ("tea", "Aromatic"),
        ("dessert", "Sweet"),
        ("cake", "Sweet"),
        ("juice", "Refreshing"),
        ("western", "Hearty"),
        ("seafood", "Umami"),
        ("japanese", "Balanced"),
        ("korean", "Bold"),
        ("indian", "Spiced"),
        ("thai", "Tangy"),
        ("vegetarian", "Fresh"),
    ]:
        if needle in text and tag not in tags:
            tags.append(tag)
    return tags[:4] or ["Local"]


def infer_taste_vector(cuisine: str, dish_name: str) -> dict[str, int]:
    taxonomy_entry = lookup_dish_taxonomy(dish_name)
    if taxonomy_entry:
        return dict(taxonomy_entry.taste_vector)
    text = f"{cuisine} {dish_name}".lower()
    vector = {"sweet": 4, "spicy": 3, "sour": 3, "salty": 6, "rich": 5, "aromatic": 6}
    if any(word in text for word in ("dessert", "cake", "sweet", "mango", "chocolate")):
        vector.update({"sweet": 8, "salty": 3, "rich": 6})
    if any(word in text for word in ("spicy", "mala", "curry", "sambal", "laksa", "thai")):
        vector.update({"spicy": 7, "aromatic": 8})
    if any(word in text for word in ("soup", "porridge", "fish soup")):
        vector.update({"rich": 3, "salty": 5})
    if any(word in text for word in ("coffee", "tea", "kopi")):
        vector.update({"sweet": 5, "salty": 2, "aromatic": 9})
    return vector


def infer_food_type_tag(
    cuisine_tags: list[str],
    dish_name: str,
    *,
    explicit: str = "",
) -> str | None:
    """Prefer source-authored types, then the structured dish taxonomy."""

    if explicit:
        return explicit
    taxonomy_entry = lookup_dish_taxonomy(dish_name)
    if taxonomy_entry:
        return taxonomy_entry.cuisine_category
    return cuisine_tags[0] if cuisine_tags else None


PRICE_RE = re.compile(
    r"\((?:S\$|\$)?\s*(\d+(?:\.\d+)?)\s*(?:[–-]\s*(?:S\$|\$)?\s*(\d+(?:\.\d+)?))?\)",
    flags=re.I,
)


def parse_price(item: str, fallback_min: float | None, fallback_max: float | None) -> tuple[float, str]:
    match = PRICE_RE.search(item)
    if match:
        price_min = float(match.group(1))
        price_max = float(match.group(2)) if match.group(2) else price_min
    else:
        price_min = fallback_min or 0.0
        price_max = fallback_max or price_min
    if price_max and price_max != price_min:
        return price_min, f"S${price_min:.2f} - S${price_max:.2f}"
    return price_min, f"S${price_min:.2f}"


def clean_dish_name(item: str) -> str:
    text = PRICE_RE.sub("", item).strip()
    text = re.sub(r"^\d+\s*[.)]\s*", "", text)
    text = re.sub(r"\s+", " ", text).strip(" -")
    return text[:120]


def split_menu_items(menu: str) -> list[str]:
    if not menu:
        return []
    if ";" in menu:
        parts = menu.split(";")
    else:
        parts = re.split(r",(?=\s*(?:\d+[.)]?\s+|[A-Z0-9]))", menu)
    cleaned: list[str] = []
    seen: set[str] = set()
    for part in parts:
        item = part.strip()
        if not item:
            continue
        name = clean_dish_name(item)
        key = name.lower()
        if len(name) < 2 or key in seen:
            continue
        seen.add(key)
        cleaned.append(item)
    return cleaned


def normalize_lookup_key(value: Any) -> str:
    text = clean(value)
    if text is None:
        return ""
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def venue_family_key(value: Any) -> str:
    key = normalize_lookup_key(value)
    if key.startswith("ven sfa "):
        return key.removeprefix("ven ")
    return key


def normalize_stall_lookup_key(value: Any) -> str:
    text = clean(value)
    if text is None:
        return ""
    normalized = str(text)
    normalized = re.sub(r"\s+-\s+.+$", "", normalized)
    normalized = re.sub(r"\([^)]*\)", " ", normalized)
    key = normalize_lookup_key(normalized)
    key = re.sub(r"\b(?:pte|private)\s+ltd\b", " ", key)
    key = re.sub(r"\b(?:ltd|llp|inc|co|company)\b", " ", key)
    key = re.sub(r"\b(?:singapore|sg)\b", " ", key)
    key = re.sub(r"\s+", " ", key).strip()
    return key


def stall_lookup_keys(value: Any) -> list[str]:
    exact_key = normalize_lookup_key(value)
    brand_key = normalize_stall_lookup_key(value)
    keys: list[str] = []
    for key in (exact_key, brand_key):
        if key and key not in keys:
            keys.append(key)
    return keys


def clean_evidence_dish_name(value: Any) -> str:
    text = clean(value)
    if text is None:
        return ""
    dish_name = re.sub(r"^\s*[A-Z]?\d+\s*[.)-]\s*", "", str(text), flags=re.I)
    dish_name = re.sub(r"\s+", " ", dish_name).strip(" -")
    return dish_name


def stable_url_choice(urls: list[str], seed: str) -> str:
    if not urls:
        return ""
    digest = hashlib.sha1(seed.encode("utf-8")).hexdigest()
    return urls[int(digest[:8], 16) % len(urls)]


def load_source_image_evidence(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Source Image Evidence" not in workbook.sheetnames:
        return {
            "rows": 0,
            "directRows": 0,
            "byDish": {},
            "byStall": {},
        }

    sheet = workbook["Source Image Evidence"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    by_dish: defaultdict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    by_stall: defaultdict[str, list[str]] = defaultdict(list)
    row_count = 0
    direct_count = 0

    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        row_count += 1
        row = {str(headers[index]): values[index] if index < len(values) else None for index in range(len(headers))}
        if text_value(row, "url_kind").lower() != "direct_image":
            continue
        image_url = good_url(row_value(row, "media_or_source_url"))
        stall_keys = stall_lookup_keys(row_value(row, "stall_name"))
        dish_key = normalize_lookup_key(clean_evidence_dish_name(row_value(row, "dish_name")))
        if not image_url or not stall_keys:
            continue
        direct_count += 1
        for stall_key in stall_keys:
            if dish_key and not any(item["url"] == image_url for item in by_dish[(stall_key, dish_key)]):
                by_dish[(stall_key, dish_key)].append(
                    {"url": image_url, "price": numeric_value(row_value(row, "price"))}
                )
            if image_url not in by_stall[stall_key]:
                by_stall[stall_key].append(image_url)

    return {
        "rows": row_count,
        "directRows": direct_count,
        "byDish": dict(by_dish),
        "byStall": dict(by_stall),
    }


def load_sheet_rows(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        row = {str(headers[index]): values[index] if index < len(values) else None for index in range(len(headers))}
        row["_row_number"] = row_number
        rows.append(row)
    return rows


def load_named_sheet_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        row = {str(headers[index]): values[index] if index < len(headers) else None for index in range(len(headers))}
        row["_row_number"] = row_number
        rows.append(row)
    return rows


def build_catalog(rows: list[dict[str, Any]], workbook_path: Path, max_menu_items: int) -> dict[str, Any]:
    source_image_evidence = load_source_image_evidence(workbook_path)
    source_images_by_dish: dict[tuple[str, str], list[dict[str, Any]]] = source_image_evidence["byDish"]
    source_images_by_stall: dict[str, list[str]] = source_image_evidence["byStall"]
    buildings_by_key: dict[str, dict[str, Any]] = {}
    building_venue_counts: defaultdict[str, int] = defaultdict(int)
    building_hero_dishes: dict[str, str] = {}
    venues: list[dict[str, Any]] = []
    dishes: list[dict[str, Any]] = []
    dish_images_from_dish_column = 0
    dish_images_from_source_evidence = 0
    dish_images_from_venue_image = 0
    represented_business_ids: set[str] = set()
    represented_venue_families: set[str] = set()
    represented_venue_keys: set[str] = set()
    supabase_venues = load_named_sheet_rows(workbook_path, "Supabase Venues")
    supabase_dishes = load_named_sheet_rows(workbook_path, "Supabase Dishes")
    supabase_venues_by_family: dict[str, dict[str, Any]] = {}
    for venue_row in supabase_venues:
        source_venue_id = text_value(venue_row, "id")
        family_key = venue_family_key(source_venue_id)
        if not family_key:
            continue
        existing_row = supabase_venues_by_family.get(family_key)
        row_has_logo = bool(good_url(row_value(venue_row, "logo_url"), allow_wikimedia=True))
        row_has_image = bool(good_url(row_value(venue_row, "image_url")))
        existing_has_logo = bool(
            existing_row and good_url(row_value(existing_row, "logo_url"), allow_wikimedia=True)
        )
        existing_has_image = bool(existing_row and good_url(row_value(existing_row, "image_url")))
        if (
            existing_row is None
            or (row_has_logo and not existing_has_logo)
            or (row_has_image and not existing_has_image)
        ):
            supabase_venues_by_family[family_key] = venue_row
    supabase_dishes_by_venue: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for dish_row in supabase_dishes:
        if not bool_value(dish_row, "menu_approved_for_display"):
            continue
        venue_id = text_value(dish_row, "venue_id")
        dish_name = text_value(dish_row, "name")
        if not venue_id or not dish_name:
            continue
        supabase_dishes_by_venue[venue_id].append(dish_row)

    for row in rows:
        row_number = int(row["_row_number"])
        name = text_value(row, "Stall / Restaurant Name") or text_value(row, "Google Display Name")
        if not name:
            continue
        raw_business_id = text_value(row, "Business ID")
        business_id = normalize_lookup_key(raw_business_id)
        if business_id:
            represented_business_ids.add(business_id)
            represented_venue_families.add(venue_family_key(raw_business_id))
        area = title_case_area(text_value(row, "Area") or text_value(row, "Location") or text_value(row, "Region"))
        postal = text_value(row, "Postal Code")
        unit_number = text_value(row, "Unit No.") or "-"
        represented_venue_keys.add(
            normalize_lookup_key("|".join([name, postal, unit_number, text_value(row, "Address")]))
        )
        floor = unit_number.split("-")[0] if unit_number.startswith("#") and "-" in unit_number else "L1"
        building_name = (
            text_value(row, "Building Name")
            or text_value(row, "Hawker Centre / Building")
            or "Standalone"
        )
        building_key_source = f"{building_name}|{area}|{postal if building_name.lower() != 'standalone' else row_number}"
        building_id = "excel-building-" + slugify(building_key_source, f"row-{row_number}")
        venue_seed = text_value(row, "Business ID") or text_value(row, "Licence No") or f"{name}|{postal}|{unit_number}|{row_number}"
        venue_id = "excel-venue-" + slugify(venue_seed, f"row-{row_number}")
        latitude = number_value(row, "Latitude") or number_value(row, "Google Latitude")
        longitude = number_value(row, "Longitude") or number_value(row, "Google Longitude")
        coordinate = {
            "latitude": latitude if latitude is not None else DEFAULT_COORDINATE["latitude"],
            "longitude": longitude if longitude is not None else DEFAULT_COORDINATE["longitude"],
        }
        logo_url = good_url(row_value(row, "Logo URL"), allow_wikimedia=True)
        venue_image_url = good_url(row_value(row, "Image URL"))
        supabase_venue_row = supabase_venues_by_family.get(venue_family_key(raw_business_id)) if raw_business_id else None
        if supabase_venue_row:
            logo_url = logo_url or good_url(row_value(supabase_venue_row, "logo_url"), allow_wikimedia=True)
            venue_image_url = venue_image_url or good_url(row_value(supabase_venue_row, "image_url"))
        dish_image_url = good_url(row_value(row, "Dishes Image URL"))
        current_stall_lookup_keys = stall_lookup_keys(name)
        cuisine = text_value(row, "Cuisine")
        cuisine_tags = split_tags(cuisine)
        flags = dietary_flags(text_value(row, "Dietary Options"))
        fallback_min = number_value(row, "Price Min (SGD)")
        fallback_max = number_value(row, "Price Max (SGD)")
        menu_items = split_menu_items(text_value(row, "Menu Items"))
        if not menu_items:
            signature = text_value(row, "Signature Dish") or text_value(row, "Menu")
            if signature:
                menu_items = [signature]
        if max_menu_items > 0:
            menu_items = menu_items[:max_menu_items]

        supabase_dish_rows: list[dict[str, Any]] = []
        if raw_business_id:
            supabase_dish_rows.extend(supabase_dishes_by_venue.get(raw_business_id, []))
            supabase_dish_rows.extend(supabase_dishes_by_venue.get(f"ven-{raw_business_id}", []))
        supabase_dishes_by_name: dict[str, dict[str, Any]] = {}
        for dish_row in supabase_dish_rows:
            dish_lookup_key = normalize_lookup_key(text_value(dish_row, "name"))
            if not dish_lookup_key:
                continue
            existing_dish_row = supabase_dishes_by_name.get(dish_lookup_key)
            existing_has_image = bool(
                existing_dish_row
                and bool_value(existing_dish_row, "media_approved_for_display")
                and good_url(row_value(existing_dish_row, "image_url"))
            )
            next_has_image = bool(
                bool_value(dish_row, "media_approved_for_display")
                and good_url(row_value(dish_row, "image_url"))
            )
            if existing_dish_row is None or (next_has_image and not existing_has_image):
                supabase_dishes_by_name[dish_lookup_key] = dish_row

        venue_dish_ids: list[str] = []
        for index, item in enumerate(menu_items):
            dish_name = clean_dish_name(item)
            if not dish_name:
                continue
            price, price_range = parse_price(item, fallback_min, fallback_max)
            dish_id = f"{venue_id}-dish-{index + 1}-{slugify(dish_name, 'item')}"
            dish_key = normalize_lookup_key(dish_name)
            exact_source_evidence = next(
                (
                    source_images_by_dish[(stall_key, dish_key)]
                    for stall_key in current_stall_lookup_keys
                    if (stall_key, dish_key) in source_images_by_dish
                ),
                [],
            )
            representative_source_images = next(
                (
                    source_images_by_stall[stall_key]
                    for stall_key in current_stall_lookup_keys
                    if stall_key in source_images_by_stall
                ),
                [],
            )
            source_image_seed = "|".join([*current_stall_lookup_keys, dish_key, str(index)])
            exact_source_price = exact_source_evidence[0].get("price") if exact_source_evidence else None
            supabase_dish_row = supabase_dishes_by_name.get(dish_key)
            supabase_dish_image_url = (
                good_url(row_value(supabase_dish_row, "image_url"))
                if supabase_dish_row and bool_value(supabase_dish_row, "media_approved_for_display")
                else ""
            )
            source_image_url = exact_source_evidence[0]["url"] if exact_source_evidence else stable_url_choice(
                representative_source_images,
                source_image_seed,
            )
            image_url = supabase_dish_image_url or dish_image_url or source_image_url or venue_image_url
            if supabase_dish_image_url or dish_image_url:
                dish_images_from_dish_column += 1
            elif source_image_url:
                dish_images_from_source_evidence += 1
            elif venue_image_url:
                dish_images_from_venue_image += 1
            if exact_source_price is not None:
                price = float(exact_source_price)
                price_range = f"S${price:.2f}"
            elif supabase_dish_row:
                supabase_price = number_value(supabase_dish_row, "price")
                if supabase_price is not None and supabase_price > 0:
                    price = supabase_price
                    price_range = text_value(supabase_dish_row, "price_range") or f"S${price:.2f}"
            dish = {
                "id": dish_id,
                "venueId": venue_id,
                "buildingId": building_id,
                "name": dish_name,
                "foodTypeTag": infer_food_type_tag(cuisine_tags, dish_name),
                "imageUrl": image_url,
                "distanceMeters": 0,
                "price": price,
                "priceRange": price_range,
                "tasteTags": infer_taste_tags(cuisine, dish_name),
                "tasteVector": infer_taste_vector(cuisine, dish_name),
                "dietaryFlags": flags,
                "mustOrder": index == 0,
            }
            dishes.append({key: value for key, value in dish.items() if value is not None})
            venue_dish_ids.append(dish_id)

        top_dish_id = venue_dish_ids[0] if venue_dish_ids else ""
        venue = {
            "id": venue_id,
            "buildingId": building_id,
            "name": name,
            "unitNumber": unit_number,
            "floor": floor,
            "isOpen": True,
            "topDishId": top_dish_id,
            "operatingHours": text_value(row, "Opening Hours") or "Hours not provided",
            "deliveryLinks": [],
            "planningArea": area,
            "imageUrl": venue_image_url,
            "logoUrl": logo_url,
            "buildingName": building_name,
            "address": text_value(row, "Address") or text_value(row, "Google Formatted Address"),
            "postalCode": postal,
            "latitude": coordinate["latitude"],
            "longitude": coordinate["longitude"],
            "rating": number_value(row, "Google Rating") or 0.0,
            "reviewCount": int(number_value(row, "Google Review Count") or 0),
            "cuisineTags": cuisine_tags,
            "dietaryFlags": flags,
        }
        venues.append(venue)
        building_venue_counts[building_id] += 1
        if top_dish_id and building_id not in building_hero_dishes:
            building_hero_dishes[building_id] = top_dish_id
        if building_id not in buildings_by_key:
            buildings_by_key[building_id] = {
                "id": building_id,
                "name": building_name if building_name.lower() != "standalone" else name,
                "type": infer_building_type(row),
                "latitude": coordinate["latitude"],
                "longitude": coordinate["longitude"],
                "floors": [floor],
                "venueCount": 0,
                "isClosedToday": False,
                "heroDishId": top_dish_id,
                "planningArea": area,
                "imageUrl": venue_image_url,
                "logoUrl": logo_url,
            }
        else:
            floors = buildings_by_key[building_id]["floors"]
            if floor not in floors:
                floors.append(floor)

    def supabase_venue_sort_key(row: dict[str, Any]) -> tuple[str, int]:
        source_venue_id = text_value(row, "id")
        return (venue_family_key(source_venue_id), 1 if source_venue_id.startswith("ven-") else 0)

    for row in sorted(supabase_venues, key=supabase_venue_sort_key):
        source_venue_id = text_value(row, "id")
        name = text_value(row, "name")
        if not source_venue_id or not name or not bool_value(row, "menu_approved_for_display"):
            continue
        source_venue_family = venue_family_key(source_venue_id)
        business_id = normalize_lookup_key(text_value(row, "business_id"))
        business_family = venue_family_key(text_value(row, "business_id"))
        unit_number = text_value(row, "unit_number") or "-"
        postal = text_value(row, "postal_code")
        address = text_value(row, "address")
        venue_key = normalize_lookup_key("|".join([name, postal, unit_number, address]))
        if (
            source_venue_family in represented_venue_families
            or (business_id and business_id in represented_business_ids)
            or (business_family and business_family in represented_venue_families)
            or venue_key in represented_venue_keys
        ):
            continue

        approved_dish_rows = supabase_dishes_by_venue.get(source_venue_id, [])
        if not approved_dish_rows:
            continue

        area = title_case_area(text_value(row, "planning_area"))
        floor = text_value(row, "floor") or (
            unit_number.split("-")[0] if unit_number.startswith("#") and "-" in unit_number else "L1"
        )
        source_building_id = text_value(row, "building_id")
        building_name = "Standalone"
        building_id = "excel-building-supabase-" + slugify(
            source_building_id or f"{address}|{postal}|{area}",
            f"venue-{len(venues) + 1}",
        )
        venue_id = source_venue_id
        latitude = number_value(row, "latitude")
        longitude = number_value(row, "longitude")
        coordinate = {
            "latitude": latitude if latitude is not None else DEFAULT_COORDINATE["latitude"],
            "longitude": longitude if longitude is not None else DEFAULT_COORDINATE["longitude"],
        }
        logo_url = good_url(row_value(row, "logo_url"), allow_wikimedia=True)
        venue_image_url = (
            good_url(row_value(row, "image_url")) if bool_value(row, "media_approved_for_display") else ""
        )
        cuisine_tags = split_tags(text_value(row, "cuisine_tags"))
        flags = dietary_flags(text_value(row, "dietary_tags"))
        venue_dish_ids: list[str] = []

        for index, dish_row in enumerate(approved_dish_rows):
            dish_name = text_value(dish_row, "name")
            price = number_value(dish_row, "price") or 0.0
            price_range = text_value(dish_row, "price_range") or (f"S${price:.2f}" if price > 0 else "")
            dish_image_url = (
                good_url(row_value(dish_row, "image_url"))
                if bool_value(dish_row, "media_approved_for_display")
                else ""
            )
            if dish_image_url:
                dish_images_from_dish_column += 1
            dish_id = text_value(dish_row, "id") or f"{venue_id}-dish-{index + 1}-{slugify(dish_name, 'item')}"
            dish = {
                "id": dish_id,
                "venueId": venue_id,
                "buildingId": building_id,
                "name": dish_name,
                "foodTypeTag": infer_food_type_tag(
                    cuisine_tags,
                    dish_name,
                    explicit=text_value(dish_row, "food_type_tag"),
                ),
                "imageUrl": dish_image_url,
                "distanceMeters": 0,
                "price": price,
                "priceRange": price_range,
                "tasteTags": infer_taste_tags(",".join(cuisine_tags), dish_name),
                "tasteVector": infer_taste_vector(",".join(cuisine_tags), dish_name),
                "dietaryFlags": flags,
                "mustOrder": index == 0,
            }
            dishes.append({key: value for key, value in dish.items() if value is not None})
            venue_dish_ids.append(dish_id)

        top_dish_id = venue_dish_ids[0] if venue_dish_ids else ""
        venue = {
            "id": venue_id,
            "buildingId": building_id,
            "name": name,
            "unitNumber": unit_number,
            "floor": floor,
            "isOpen": True,
            "topDishId": top_dish_id,
            "operatingHours": text_value(row, "operating_hours") or "Hours not provided",
            "deliveryLinks": [],
            "planningArea": area,
            "imageUrl": venue_image_url,
            "logoUrl": logo_url,
            "buildingName": building_name,
            "address": address,
            "postalCode": postal,
            "latitude": coordinate["latitude"],
            "longitude": coordinate["longitude"],
            "rating": number_value(row, "google_rating") or 0.0,
            "reviewCount": int(number_value(row, "google_review_count") or 0),
            "cuisineTags": cuisine_tags,
            "dietaryFlags": flags,
        }
        venues.append(venue)
        if business_id:
            represented_business_ids.add(business_id)
        represented_venue_families.add(source_venue_family)
        if business_family:
            represented_venue_families.add(business_family)
        represented_venue_keys.add(venue_key)
        building_venue_counts[building_id] += 1
        if top_dish_id and building_id not in building_hero_dishes:
            building_hero_dishes[building_id] = top_dish_id
        if building_id not in buildings_by_key:
            buildings_by_key[building_id] = {
                "id": building_id,
                "name": building_name if building_name.lower() != "standalone" else name,
                "type": "standalone",
                "latitude": coordinate["latitude"],
                "longitude": coordinate["longitude"],
                "floors": [floor],
                "venueCount": 0,
                "isClosedToday": False,
                "heroDishId": top_dish_id,
                "planningArea": area,
                "imageUrl": venue_image_url,
                "logoUrl": logo_url,
            }
        else:
            floors = buildings_by_key[building_id]["floors"]
            if floor not in floors:
                floors.append(floor)

    for building_id, building in buildings_by_key.items():
        building["venueCount"] = building_venue_counts[building_id]
        building["heroDishId"] = building_hero_dishes.get(building_id, building.get("heroDishId") or "")

    workbook_bytes = workbook_path.read_bytes()
    return {
        "metadata": {
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "sourceWorkbook": str(workbook_path),
            "sourceSha256": hashlib.sha256(workbook_bytes).hexdigest(),
            "sourceRows": len(rows),
            "maxMenuItemsPerVenue": max_menu_items,
            "buildingCount": len(buildings_by_key),
            "venueCount": len(venues),
            "dishCount": len(dishes),
            "dishImageCount": sum(1 for dish in dishes if dish.get("imageUrl")),
            "dishImagesFromDishColumn": dish_images_from_dish_column,
            "dishImagesFromSourceEvidence": dish_images_from_source_evidence,
            "dishImagesFromVenueImage": dish_images_from_venue_image,
            "sourceImageEvidenceRows": source_image_evidence["rows"],
            "sourceDirectImageRows": source_image_evidence["directRows"],
            "sourceImageEvidenceDishKeys": len(source_images_by_dish),
            "sourceImageEvidenceStallKeys": len(source_images_by_stall),
        },
        "buildingColumns": BUILDING_COLUMNS,
        "buildings": compact_rows(list(buildings_by_key.values()), BUILDING_COLUMNS),
        "venueColumns": VENUE_COLUMNS,
        "venues": compact_rows(venues, VENUE_COLUMNS),
        "dishColumns": DISH_COLUMNS,
        "dishes": compact_rows(dishes, DISH_COLUMNS),
    }


def compact_rows(rows: list[dict[str, Any]], columns: list[str]) -> list[list[Any]]:
    compacted: list[list[Any]] = []
    for row in rows:
        values: list[Any] = []
        for column in columns:
            value = row.get(column)
            values.append("" if value is None else value)
        compacted.append(values)
    return compacted


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--max-menu-items", type=int, default=MAX_MENU_ITEMS_PER_VENUE)
    args = parser.parse_args()

    rows = load_sheet_rows(args.workbook)
    catalog = build_catalog(rows, args.workbook, args.max_menu_items)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(catalog["metadata"], indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
