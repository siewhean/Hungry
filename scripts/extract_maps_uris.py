#!/usr/bin/env python3
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "Hungry F&B Database.xlsx"
OUTPUT_PATH = ROOT / "data" / "uris_to_check.json"
SHEET_TAB = "Sheet1"

def clean(value) -> str:
    return "" if value is None else str(value).strip()

def main():
    print("Loading Excel workbook...")
    workbook = load_workbook(EXCEL_PATH, read_only=True)
    worksheet = workbook[SHEET_TAB]

    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean(v) for v in header_row]

    uri_idx = headers.index("Google Maps URI") if "Google Maps URI" in headers else -1

    if uri_idx == -1:
        print("No Google Maps URI column.")
        return

    targets = []

    for idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        uri = clean(row[uri_idx]) if uri_idx < len(row) else ""
        if uri:
            targets.append({
                "row_index": idx,
                "uri": uri
            })

    print(f"Found {len(targets)} URIs to check.")

    # Dump to JSON
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(targets, f, indent=2)

if __name__ == "__main__":
    main()
