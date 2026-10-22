#!/usr/bin/env python3
"""Extract reviewable F&B data candidates from permitted public first-party pages.

This tool is intentionally proposal-only. It never writes to the workbook, a
database, or a third-party account. It has no API-key support and makes only
bounded HTTPS GET requests to public URLs already present in the workbook or
an explicit per-licence source map.
"""

from __future__ import annotations

import argparse
import hashlib
import http.client
import ipaddress
import json
import re
import socket
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.robotparser
import urllib.request
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data/Hungry F&B Database.Public Access 2026-10-17 Verified Menu Catalogue Refresh 40.xlsx"
USER_AGENT = "HungryPublicFnbResearch/1.0"
MAX_REQUEST_CAP = 5000
MAX_BODY_BYTES = 2_000_000
MAX_EVIDENCE_CHARS = 360
MAX_ITEMS_PER_PAGE = 250
BLOCKED_HOST_SUFFIXES = (
    "google.com",
    "google.com.sg",
    "googleapis.com",
    "grab.com",
    "foodpanda.sg",
    "foodpanda.com",
    "oddle.me",
    "deliveroo.com.sg",
    "mapbox.com",
    "onemap.gov.sg",
)
BLOCKED_QUERY_KEYS = {"key", "api_key", "apikey", "access_token", "token", "client_secret", "subscription-key"}
FIELD_HEADERS = {
    "F&B Name": "F&B Name",
    "SFA Licence": "SFA Licence",
    "Full Address": "Full Address",
    "Postal Code": "Postal Code",
    "Cuisine": "Cuisine",
    "Menu / Price": "Menu / Price",
    "Opening Hours / Closed Days": "Opening Hours / Closed Days",
    "Contact Details": "Contact Details",
    "Website": "Website",
    "Menu URL": "Menu URL",
    "Latitude": "Latitude",
    "Longitude": "Longitude",
    "Rating / Review Count": "Rating / Review Count",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (list, tuple)):
        parts = [clean_text(item) for item in value]
        return ", ".join(part for part in parts if part) or None
    if isinstance(value, dict):
        return clean_text(value.get("name") or value.get("@id") or value.get("url"))
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def format_value(value: Any) -> str | None:
    if isinstance(value, (dict, list)):
        if isinstance(value, dict) and (value.get("url") or value.get("name") or value.get("@id")):
            return clean_text(value.get("url") or value.get("name") or value.get("@id"))
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return clean_text(value)


def resolve_image_url(value: Any, page_url: str) -> str | None:
    if isinstance(value, list):
        value = next((item for item in value if item), None)
    if isinstance(value, dict):
        value = value.get("url") or value.get("contentUrl") or value.get("thumbnailUrl")
    text = clean_text(value)
    if not text:
        return None
    absolute = urllib.parse.urljoin(page_url, text)
    return absolute if urllib.parse.urlsplit(absolute).scheme == "https" else None


def schema_types(node: dict[str, Any]) -> set[str]:
    raw = node.get("@type", [])
    if isinstance(raw, str):
        raw = [raw]
    return {str(item).rsplit("/", 1)[-1].lower() for item in raw if item}


def walk_json(value: Any):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk_json(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_json(child)


class PageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.scripts: list[str] = []
        self.links: list[dict[str, str]] = []
        self.metas: dict[str, str] = {}
        self.title_parts: list[str] = []
        self._in_jsonld = False
        self._in_title = False
        self._script_parts: list[str] = []
        self._active_anchor: int | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attr = {key.lower(): value or "" for key, value in attrs}
        if tag == "script" and attr.get("type", "").lower() == "application/ld+json":
            self._in_jsonld = True
            self._script_parts = []
        elif tag == "title":
            self._in_title = True
        elif tag == "meta":
            key = (attr.get("property") or attr.get("name") or "").lower()
            if key and attr.get("content"):
                self.metas[key] = attr["content"]
        elif tag == "link" and attr.get("href"):
            self.links.append({"rel": attr.get("rel", "").lower(), "href": attr["href"], "text": ""})
        elif tag == "a" and attr.get("href"):
            self.links.append({"rel": "anchor", "href": attr["href"], "text": ""})
            self._active_anchor = len(self.links) - 1

    def handle_endtag(self, tag: str) -> None:
        if tag == "script" and self._in_jsonld:
            self.scripts.append("".join(self._script_parts))
            self._script_parts = []
            self._in_jsonld = False
        elif tag == "title":
            self._in_title = False
        elif tag == "a":
            self._active_anchor = None

    def handle_data(self, data: str) -> None:
        if self._in_jsonld:
            self._script_parts.append(data)
        if self._in_title:
            self.title_parts.append(data)
        if self._active_anchor is not None:
            self.links[self._active_anchor]["text"] += data


def parse_page(html: str, page_url: str) -> dict[str, Any]:
    parser = PageParser()
    parser.feed(html)
    schema: list[Any] = []
    for raw in parser.scripts:
        try:
            schema.append(json.loads(raw))
        except (json.JSONDecodeError, TypeError):
            continue
    nodes = [node for root in schema for node in walk_json(root)]
    entities = [node for node in nodes if schema_types(node) & {"restaurant", "foodestablishment", "localbusiness", "cafe", "bakery"}]
    entity = next((item for item in entities if item.get("name")), entities[0] if entities else {})
    proposals: list[dict[str, Any]] = []

    def add(field: str, value: Any, evidence: str, confidence: str = "medium") -> None:
        value_text = format_value(value)
        if value_text:
            proposals.append({"field": field, "value": value_text, "evidence": evidence[:MAX_EVIDENCE_CHARS], "confidence": confidence})

    add("F&B Name", entity.get("name"), "schema.org entity name", "high")
    add("Cuisine", entity.get("servesCuisine"), "schema.org servesCuisine", "medium")
    add("Website", entity.get("url") or parser.metas.get("og:url"), "schema.org url / og:url", "medium")
    add("Menu URL", entity.get("menu") or entity.get("hasMenu"), "schema.org menu / hasMenu", "medium")
    add("Opening Hours / Closed Days", entity.get("openingHours") or entity.get("openingHoursSpecification"), "schema.org openingHours", "medium")
    contact_links = [urllib.parse.unquote(item["href"].split(":", 1)[1]) for item in parser.links if item["href"].lower().startswith(("tel:", "mailto:"))]
    add("Contact Details", entity.get("telephone") or contact_links, "schema.org telephone or official page tel/mailto link", "medium")
    if not entity.get("openingHours") and not entity.get("openingHoursSpecification"):
        for meta_key in ("business:hours", "og:hours", "opening_hours"):
            if parser.metas.get(meta_key):
                add("Opening Hours / Closed Days", parser.metas[meta_key], f"HTML {meta_key}", "low")
                break

    address = entity.get("address")
    if isinstance(address, dict):
        address_text = ", ".join(filter(None, [clean_text(address.get(k)) for k in ("streetAddress", "addressLocality", "postalCode")]))
        add("Full Address", address_text, "schema.org PostalAddress", "medium")
        add("Postal Code", address.get("postalCode"), "schema.org PostalAddress.postalCode", "medium")
    geo = entity.get("geo")
    if isinstance(geo, dict):
        add("Latitude", geo.get("latitude"), "schema.org GeoCoordinates; address/building point", "medium")
        add("Longitude", geo.get("longitude"), "schema.org GeoCoordinates; address/building point", "medium")
    rating = entity.get("aggregateRating")
    if isinstance(rating, dict):
        score = clean_text(rating.get("ratingValue"))
        count = clean_text(rating.get("reviewCount") or rating.get("ratingCount"))
        if score:
            add("Rating / Review Count", f"{score} / {count or 'review count not published'}", "schema.org aggregateRating; first-party page", "medium")

    menu_items: list[dict[str, Any]] = []
    for node in nodes:
        types = schema_types(node)
        if "menuitem" not in types and not (node.get("name") and ("offers" in node or "nutrition" in node)):
            continue
        name = clean_text(node.get("name"))
        if not name or name.lower() in {"menu", "food menu"}:
            continue
        offers = node.get("offers")
        offer = offers[0] if isinstance(offers, list) and offers else offers
        if not isinstance(offer, dict):
            offer = {}
        menu_items.append({
            "name": name,
            "description": clean_text(node.get("description")),
            "price": clean_text(offer.get("price")),
            "price_currency": clean_text(offer.get("priceCurrency")),
            "evidence": "schema.org MenuItem/Offer",
            "rights_status": "source rights unverified; review before retention/display",
        })
        item_image = resolve_image_url(node.get("image"), page_url)
        if item_image:
            menu_items[-1]["image_url"] = item_image
        if len(menu_items) >= MAX_ITEMS_PER_PAGE:
            break

    page_links: list[dict[str, str]] = []
    for item in parser.links:
        absolute = urllib.parse.urljoin(page_url, item["href"])
        parts = urllib.parse.urlsplit(absolute)
        if parts.scheme not in {"https", "http"}:
            continue
        if item["rel"] == "anchor" and re.search(r"menu|food|contact|about|location|branch", absolute, re.I):
            label = re.sub(r"\s+", " ", item.get("text", "")).strip()[:100]
            page_links.append({"url": absolute, "label": label})
            if re.search(r"menu|food", f"{absolute} {label}", re.I):
                add("Menu URL", absolute, f"official page link labeled {label or 'menu'}", "low")
        elif item["rel"] == "canonical":
            add("Website", absolute, "HTML canonical link", "medium")

    images = []
    for node in nodes:
        if "menuitem" in schema_types(node) and node.get("image"):
            item_image = resolve_image_url(node.get("image"), page_url)
            if item_image:
                images.append({"url": item_image, "scope": "menu item", "rights_status": "unverified"})
    og_image = parser.metas.get("og:image")
    if og_image:
        images.append({"url": urllib.parse.urljoin(page_url, og_image), "scope": "page/venue image", "rights_status": "unverified"})

    title = clean_text(" ".join(parser.title_parts))
    if title:
        add("Page Title", title, "HTML title", "low")
    for key, field in (("og:site_name", "Website Brand"), ("og:description", "Page Description")):
        add(field, parser.metas.get(key), f"HTML {key}", "low")

    return {
        "proposals": proposals,
        "menu_items": menu_items,
        "image_candidates": images,
        "related_page_links": page_links,
        "schema_entity_types": sorted(set().union(*(schema_types(node) for node in entities))) if entities else [],
        "page_title": title,
    }


def normalized_host(url: str) -> str | None:
    try:
        parts = urllib.parse.urlsplit(url)
        if parts.scheme.lower() != "https" or not parts.hostname or parts.username or parts.password or parts.port not in (None, 443):
            return None
        host = parts.hostname.rstrip(".").lower().encode("idna").decode("ascii")
        return host
    except (ValueError, UnicodeError):
        return None


def blocked_host(host: str) -> bool:
    google_host = bool(re.search(r"(?:^|\.)google\.[a-z.]+$", host))
    return google_host or any(host == suffix or host.endswith("." + suffix) for suffix in BLOCKED_HOST_SUFFIXES)


def has_credential_query(url: str) -> bool:
    query_keys = {key.lower() for key, _ in urllib.parse.parse_qsl(urllib.parse.urlsplit(url).query, keep_blank_values=True)}
    return bool(query_keys & BLOCKED_QUERY_KEYS)


def resolve_public_ip(host: str) -> str | None:
    try:
        addresses = {record[4][0] for record in socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)}
    except OSError:
        return None
    if not addresses:
        return None
    for address in addresses:
        try:
            parsed = ipaddress.ip_address(address)
        except ValueError:
            return None
        if not parsed.is_global:
            return None
    return sorted(addresses)[0]


class PinnedHTTPSConnection(http.client.HTTPSConnection):
    """Connect to a prevalidated public IP while retaining TLS hostname checks."""

    def __init__(self, host: str, pinned_ip: str, **kwargs: Any) -> None:
        self.pinned_ip = pinned_ip
        super().__init__(host, **kwargs)

    def connect(self) -> None:
        sock = socket.create_connection((self.pinned_ip, self.port), self.timeout, self.source_address)
        if self._tunnel_host:
            self.sock = sock
            self._tunnel()
        self.sock = self._context.wrap_socket(sock, server_hostname=self.host)


class PinnedHTTPSHandler(urllib.request.HTTPSHandler):
    def __init__(self, pinned_ip: str) -> None:
        super().__init__(context=ssl.create_default_context())
        self.pinned_ip = pinned_ip

    def https_open(self, req):
        return self.do_open(lambda host, **kwargs: PinnedHTTPSConnection(host, self.pinned_ip, **kwargs), req, context=self._context)


class SameHostRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        old_host = normalized_host(req.full_url)
        new_host = normalized_host(newurl)
        if not old_host or old_host != new_host:
            raise urllib.error.HTTPError(newurl, code, "cross-host or non-HTTPS redirect blocked", headers, fp)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


class Fetcher:
    def __init__(self, max_requests: int, delay: float, timeout: float, host_policy: dict[str, dict[str, Any]], user_agent: str = USER_AGENT) -> None:
        self.max_requests = max_requests
        self.delay = delay
        self.timeout = timeout
        self.host_policy = host_policy
        self.user_agent = user_agent
        self.requests = 0
        self.successful = 0
        self.skipped: list[dict[str, Any]] = []
        self.last_request_by_host: dict[str, float] = {}
        self.robots: dict[str, urllib.robotparser.RobotFileParser | None] = {}

    def _request(self, url: str) -> tuple[int, str, bytes, str]:
        if self.requests >= self.max_requests:
            raise RuntimeError("request budget exhausted")
        host = normalized_host(url)
        if not host:
            raise ValueError("only public HTTPS URLs are accepted")
        if blocked_host(host):
            raise PermissionError("source host is excluded by the no-spend/reuse policy")
        if host not in self.host_policy:
            raise PermissionError("host has no reviewed collection-permission entry")
        if has_credential_query(url):
            raise PermissionError("credential-like query parameter blocked")
        pinned_ip = resolve_public_ip(host)
        if not pinned_ip:
            raise PermissionError("host did not resolve exclusively to public IP addresses")
        wait = self.delay - (time.monotonic() - self.last_request_by_host.get(host, 0))
        if wait > 0:
            time.sleep(wait)
        request = urllib.request.Request(url, headers={"User-Agent": self.user_agent, "Accept": "text/html,application/xhtml+xml,application/ld+json;q=0.9,*/*;q=0.1"})
        self.requests += 1
        self.last_request_by_host[host] = time.monotonic()
        try:
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}), SameHostRedirectHandler(), PinnedHTTPSHandler(pinned_ip))
            with opener.open(request, timeout=self.timeout) as response:
                body = response.read(MAX_BODY_BYTES + 1)
                if len(body) > MAX_BODY_BYTES:
                    raise ValueError("response exceeded 2 MB safety limit")
                final_url = response.geturl()
                if normalized_host(final_url) != host:
                    raise PermissionError("cross-host redirect was blocked")
                self.successful += 1
                return response.status, response.headers.get("Content-Type", ""), body, final_url
        except urllib.error.HTTPError as error:
            if error.code in {402, 429}:
                raise RuntimeError(f"remote server returned {error.code}; stopping to avoid charges/rate-limit pressure") from error
            raise

    def can_fetch(self, url: str) -> tuple[bool, str]:
        host = normalized_host(url)
        if not host:
            return False, "non-HTTPS or invalid URL"
        if blocked_host(host):
            return False, "Google/maps and third-party delivery platforms are excluded"
        if host not in self.host_policy:
            return False, "host has no reviewed collection-permission entry"
        if host not in self.robots:
            robots_url = f"https://{host}/robots.txt"
            parser = urllib.robotparser.RobotFileParser(robots_url)
            try:
                status, _, body, _ = self._request(robots_url)
                if status >= 400:
                    parser = None
                else:
                    parser.parse(body.decode("utf-8", errors="replace").splitlines())
            except urllib.error.HTTPError as error:
                if error.code in {404, 410}:
                    parser = None
                else:
                    parser = False
            except RuntimeError as error:
                if "budget exhausted" in str(error) or "returned 429" in str(error) or "returned 402" in str(error):
                    raise
                parser = False
            except (OSError, ValueError, PermissionError):
                parser = False
            self.robots[host] = parser
        parser = self.robots[host]
        if parser is False:
            return False, "robots.txt could not be checked; fail closed"
        if parser and not parser.can_fetch(self.user_agent, url):
            return False, "disallowed by robots.txt"
        return True, "robots policy allowed or no robots file published"

    def fetch_page(self, url: str) -> dict[str, Any]:
        allowed, reason = self.can_fetch(url)
        if not allowed:
            raise PermissionError(reason)
        status, content_type, body, final_url = self._request(url)
        if status < 200 or status >= 300:
            raise ValueError(f"HTTP status {status}")
        if content_type and not re.search(r"(?:text/html|application/xhtml\+xml|application/ld\+json)", content_type, re.I):
            raise ValueError(f"non-HTML content type: {content_type}")
        html = body.decode("utf-8", errors="replace")
        return {"http_status": status, "content_type": content_type, "body_sha256": hashlib.sha256(body).hexdigest(), "final_url": final_url, **parse_page(html, final_url)}


def get_row_urls(row: dict[str, Any], explicit_urls: list[str]) -> list[str]:
    urls = []
    for value in [row.get("Website"), row.get("Menu URL"), *explicit_urls]:
        if not value:
            continue
        for candidate in re.split(r"[\s|;,]+", str(value)):
            candidate = candidate.strip()
            if candidate.lower().startswith("https://") and candidate not in urls:
                urls.append(candidate)
    # Keep the run narrow: at most two pages per outlet, each host is separately rate-limited.
    return urls[:2]


def load_source_map(path: Path | None) -> dict[str, list[str]]:
    if not path:
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("source map must be a JSON object keyed by SFA licence")
    result = {}
    for licence, urls in data.items():
        if isinstance(urls, str):
            urls = [urls]
        if isinstance(urls, list):
            result[str(licence)] = [str(url) for url in urls if isinstance(url, str)]
    return result


def load_host_policy(path: Path | None) -> dict[str, dict[str, Any]]:
    if not path:
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    hosts = raw.get("hosts", raw) if isinstance(raw, dict) else None
    if not isinstance(hosts, dict):
        raise ValueError("host policy must be a JSON object or contain a hosts object")
    policy: dict[str, dict[str, Any]] = {}
    for configured_host, entry in hosts.items():
        host = str(configured_host).strip().lower().rstrip(".")
        if not isinstance(entry, dict) or entry.get("collection_permitted") is not True or blocked_host(host):
            continue
        source_class = str(entry.get("source_class") or "").strip()
        no_cost_basis = str(entry.get("no_cost_basis") or "").strip()
        evidence_url = str(entry.get("terms_or_license_url") or "")
        basis = str(entry.get("permission_basis") or "").strip()
        reviewed_at = str(entry.get("reviewed_at") or "").strip()
        if (
            source_class not in {"official_website", "open_dataset"}
            or entry.get("no_request_fee") is not True
            or not no_cost_basis
            or not normalized_host(evidence_url)
            or has_credential_query(evidence_url)
            or not basis
            or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", reviewed_at)
        ):
            continue
        policy[host] = {
            "source_class": source_class,
            "no_request_fee": True,
            "no_cost_basis": no_cost_basis,
            "terms_or_license_url": evidence_url,
            "permission_basis": basis,
            "reviewed_at": reviewed_at,
        }
    return policy


def read_targets(workbook_path: Path, limit: int, offset: int, source_map: dict[str, list[str]]) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Public F&B Directory" not in workbook.sheetnames:
        raise ValueError("workbook has no Public F&B Directory sheet")
    sheet = workbook["Public F&B Directory"]
    values = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(values)]
    rows = []
    for excel_row, values_row in enumerate(values, start=2):
        row = {headers[index]: value for index, value in enumerate(values_row) if index < len(headers) and headers[index]}
        licence = clean_text(row.get("SFA Licence"))
        if not licence:
            continue
        explicit = source_map.get(licence, [])
        urls = get_row_urls(row, explicit)
        missing_fields = [
            field for field in ("Cuisine", "Menu / Price", "Opening Hours / Closed Days", "Contact Details", "Menu URL", "Latitude", "Longitude", "Rating / Review Count")
            if not clean_text(row.get(field))
        ]
        if not urls or not missing_fields:
            continue
        rows.append({"excel_row": excel_row, "sfa_licence": licence, "name": clean_text(row.get("F&B Name")), "registered_name": clean_text(row.get("Registered Business Name")), "address": clean_text(row.get("Full Address")), "postal_code": clean_text(row.get("Postal Code")), "existing": {field: clean_text(row.get(field)) for field in FIELD_HEADERS if row.get(field) not in (None, "")}, "missing_fields": missing_fields, "urls": urls})
    workbook.close()
    return rows[offset : offset + limit if limit else None]


def safe_output_dir(output_dir: Path | None) -> Path:
    if output_dir:
        return output_dir.resolve()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return ROOT / "data/source_snapshots" / f"free_source_harvest_{stamp}"


def run(options: argparse.Namespace) -> dict[str, Any]:
    started_at = utc_now()
    workbook_path = options.workbook.resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    source_map = load_source_map(options.source_map)
    host_policy = load_host_policy(options.host_policy)
    targets = read_targets(workbook_path, options.limit, options.offset, source_map)
    fetcher = Fetcher(options.max_requests, options.delay_seconds, options.timeout_seconds, host_policy)
    output_dir = safe_output_dir(options.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    candidates_path = output_dir / "candidates.jsonl"
    skipped_path = output_dir / "skipped.jsonl"
    extracted_outlets = 0
    candidate_field_count = 0
    menu_item_count = 0
    image_candidate_count = 0
    page_count = 0
    stopped_reason = None

    with candidates_path.open("w", encoding="utf-8") as candidates_file, skipped_path.open("w", encoding="utf-8") as skipped_file:
        for target in targets:
            outlet_pages = []
            for url in target["urls"]:
                if page_count >= options.max_pages:
                    stopped_reason = "page budget exhausted"
                    break
                try:
                    result = fetcher.fetch_page(url)
                    page_count += 1
                    record = {
                        "schema_version": 1,
                        "captured_at": utc_now(),
                        "sfa_licence": target["sfa_licence"],
                        "excel_row": target["excel_row"],
                        "outlet_name": target["name"],
                        "registered_name": target["registered_name"],
                        "target_address": target["address"],
                        "target_postal_code": target["postal_code"],
                        "source_url": url,
                        "final_url": result["final_url"],
                        "source_type": "public first-party/explicitly mapped page; source rights remain site-specific",
                        "host_policy": host_policy[normalized_host(url)],
                        "robots_result": "allowed",
                        "http_status": result["http_status"],
                        "content_type": result["content_type"],
                        "body_sha256": result["body_sha256"],
                        "schema_entity_types": result["schema_entity_types"],
                        "target_missing_fields": target["missing_fields"],
                        "proposals": result["proposals"],
                        "menu_items": result["menu_items"],
                        "image_candidates": result["image_candidates"],
                        "related_page_links": result["related_page_links"],
                        "page_title": result["page_title"],
                        "identity_review_required": True,
                        "automatic_match": False,
                    }
                    candidates_file.write(json.dumps(record, ensure_ascii=False) + "\n")
                    outlet_pages.append(record)
                    candidate_field_count += len(record["proposals"])
                    menu_item_count += len(record["menu_items"])
                    image_candidate_count += len(record["image_candidates"])
                except Exception as error:  # one bad source must not discard the batch
                    skipped_file.write(json.dumps({"captured_at": utc_now(), "sfa_licence": target["sfa_licence"], "excel_row": target["excel_row"], "url": url, "reason": str(error)[:300]}, ensure_ascii=False) + "\n")
                    fetcher.skipped.append({"sfa_licence": target["sfa_licence"], "url": url, "reason": str(error)[:300]})
                    if "budget exhausted" in str(error) or "returned 429" in str(error) or "returned 402" in str(error):
                        stopped_reason = str(error)
                        break
            if outlet_pages:
                extracted_outlets += 1
            if stopped_reason:
                break

    summary = {
        "schema_version": 1,
        "status": "completed_with_candidates" if candidate_field_count or menu_item_count else "completed_no_candidates",
        "started_at": started_at,
        "finished_at": utc_now(),
        "workbook": str(workbook_path),
        "mode": "proposal-only; workbook and databases are read-only",
        "uses_paid_apis": False,
        "network_policy": {
            "protocol": "HTTPS GET only",
            "robots_txt": "required; fail closed when unavailable except HTTP 404/410",
            "per_host_delay_seconds": options.delay_seconds,
            "max_requests_including_robots": options.max_requests,
            "max_pages": options.max_pages,
            "max_response_bytes": MAX_BODY_BYTES,
            "blocked_domains": list(BLOCKED_HOST_SUFFIXES),
            "reviewed_host_allowlist_required": True,
            "allowed_source_classes": ["official_website", "open_dataset"],
            "each_host_attests_no_request_fee": True,
            "reviewed_hosts": sorted(host_policy),
            "no_login_or_captcha_bypass": True,
            "no_google_places_or_maps_requests": True,
            "no_mapbox_or_onemap_requests": True,
        },
        "targets_selected": len(targets),
        "outlets_with_source_pages": extracted_outlets,
        "pages_fetched": page_count,
        "requests_including_robots": fetcher.requests,
        "successful_http_requests": fetcher.successful,
        "field_candidates": candidate_field_count,
        "menu_item_candidates": menu_item_count,
        "image_candidates": image_candidate_count,
        "skipped_count": len(fetcher.skipped),
        "stopped_reason": stopped_reason,
        "artifacts": {"candidates": str(candidates_path), "skipped": str(skipped_path)},
        "limitations": [
            "Only existing workbook URLs and exact-licence URLs supplied through --source-map are visited; this script does not search Google or scrape search results.",
            "Every extracted value remains a candidate until exact outlet identity, freshness, source rights, and field fit are reviewed.",
            "First-party rating metadata is reported as the source's own rating; no Google/Grab/foodpanda ratings are collected.",
            "Public image URLs are review-only; public visibility does not establish reuse rights.",
            "The extractor does not remove closed outlets or modify the workbook.",
        ],
    }
    (output_dir / "run-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--source-map", type=Path, help="JSON object mapping exact SFA Licence to permitted source URL(s)")
    parser.add_argument("--host-policy", type=Path, help="JSON file with reviewed terms/license permission per exact source host")
    parser.add_argument("--output-dir", type=Path, help="Output folder; defaults to data/source_snapshots/free_source_harvest_<time>")
    parser.add_argument("--limit", type=int, default=25, help="Maximum outlets per run; 0 means all remaining eligible outlets (request cap still applies)")
    parser.add_argument("--offset", type=int, default=0, help="Skip this many eligible outlets for repeatable batches")
    parser.add_argument("--max-requests", type=int, default=50, help=f"Hard cap including robots.txt (1-{MAX_REQUEST_CAP})")
    parser.add_argument("--max-pages", type=int, default=25, help="Maximum content pages to fetch in one run")
    parser.add_argument("--delay-seconds", type=float, default=3.0, help="Minimum delay between requests to each host")
    parser.add_argument("--timeout-seconds", type=float, default=15.0)
    options = parser.parse_args(argv)
    if options.limit < 0 or options.offset < 0:
        parser.error("--limit and --offset cannot be negative")
    if not 1 <= options.max_requests <= MAX_REQUEST_CAP:
        parser.error(f"--max-requests must be from 1 to {MAX_REQUEST_CAP}")
    if options.max_pages < 1:
        parser.error("--max-pages must be at least 1")
    if options.delay_seconds < 1.0:
        parser.error("--delay-seconds must be at least 1.0")
    if options.timeout_seconds < 1.0:
        parser.error("--timeout-seconds must be at least 1.0")
    return options


def main(argv: list[str] | None = None) -> int:
    try:
        summary = run(parse_args(argv))
    except Exception as error:
        print(f"harvest failed: {error}", file=sys.stderr)
        return 1
    print(json.dumps({key: summary[key] for key in ("status", "targets_selected", "outlets_with_source_pages", "pages_fetched", "requests_including_robots", "field_candidates", "menu_item_candidates", "image_candidates", "skipped_count", "stopped_reason", "artifacts")}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
