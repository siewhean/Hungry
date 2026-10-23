#!/usr/bin/env python3
r"""Load the canonical workbook catalog snapshot into buildings/venues/dishes.

Reads the "Supabase Venues" and "Supabase Dishes" sheets of the Hungry F&B
workbook, validates them against the catalog raw release manifest, derives one
buildings row per referenced building_id (the workbook has no buildings sheet),
and either:

  --dry-run  validate + derive in memory and print a JSON report, including a
             local estimate of catalog_raw_publish_quality_findings(); no files
             and no database writes.
  --apply    write each table to a private temp CSV file and load them with
             psql `\copy ... from '<file>'` in ONE transaction, ordered
             buildings -> venues -> dishes. Data never passes through psql's
             script parser (see the injection note below); the temp files are
             deleted afterwards. --csv-dir optionally writes them for inspection.

Workbook rows are loaded as-is. The loader never drops, merges, or rewrites
venue/dish rows to satisfy the publish quality gate; it only performs type
coercion required by the table types (see coerce_* helpers).

Safety:
  * The DB URL comes from --db-url-file or SUPABASE_DB_URL and is never printed;
    credentials reach psql through PG* environment variables, not argv.
  * Only localhost/127.0.0.1/::1 are accepted unless the host is listed in
    HUNGRY_SUPABASE_ALLOWED_HOSTS (comma-separated).
  * Refuses to load into non-empty tables unless --replace. --replace deletes
    only these three tables and refuses if any other table still references
    their rows (so no ON DELETE CASCADE/SET NULL reaches other tables).
  * The dishes INSERT guard trigger (enforce_merchant_dish_write_guard) allows
    postgres/supabase_admin/service_role; the transaction asserts current_user
    is one of those instead of disabling triggers (no session_replication_role).
  * Injection: workbook values are never written into the psql script. Inline
    `COPY ... FROM STDIN` blocks would let a value containing a `\.` line end
    the COPY early and have the following lines run as psql meta-commands
    (e.g. `\!` shell). The script only contains fixed SQL plus `\copy` lines
    naming loader-created temp files; psql streams those files to the server
    over the COPY protocol, so their contents are parsed only as CSV data.
    Every non-NULL field is CSV-quoted; NUL bytes are rejected in validate().
    subprocess is always called with list args (no shell).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import re
import subprocess
import sys
import tempfile
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator
from urllib.parse import unquote, urlparse

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "Hungry F&B Database.xlsx"
DEFAULT_MANIFEST = ROOT / "supabase" / "publish" / "catalog_raw_release.manifest.json"
VENUES_SHEET = "Supabase Venues"
DISHES_SHEET = "Supabase Dishes"
SOURCE_SHEET = "Sheet1"  # SFA/Google source rows keyed by Business ID

DB_URL_ENV = "SUPABASE_DB_URL"
ALLOWED_HOSTS_ENV = "HUNGRY_SUPABASE_ALLOWED_HOSTS"
LOCAL_HOSTS = frozenset({"localhost", "127.0.0.1", "::1"})
# Roles the dishes insert guard trigger exempts; see enforce_merchant_dish_write_guard().
TRIGGER_EXEMPT_ROLES = ("postgres", "supabase_admin", "service_role")

# Singapore bounding box used to ignore obviously wrong venue coordinates when
# averaging a building position, and the island centroid used only as a last
# resort (flagged with source_confidence 0) when no venue has coordinates.
SG_LAT_RANGE = (1.15, 1.50)
SG_LON_RANGE = (103.55, 104.10)
SG_CENTROID = (1.3521, 103.8198)

BUILDING_SOURCE_TYPE = "workbook_derived"
UNASSIGNED_BUILDING_ID = "building-unassigned"
PLACEHOLDER_VALUES = frozenset(
    {"", "standalone", "n/a", "na", "none", "null", "unknown", "-", "tbc", "tba", "nil"}
)

VENUE_COLUMNS = [
    "id", "building_id", "business_id", "stall_id", "licence_no", "name", "unit_number",
    "floor", "hygiene_grade", "operating_hours", "address", "postal_code", "planning_area",
    "latitude", "longitude", "google_place_id", "google_rating", "google_review_count",
    "source_type", "source_confidence", "last_verified_at", "updated_at",
    "media_approved_for_display", "menu_approved_for_display", "media_rights_status",
    "menu_rights_status", "image_url", "logo_url", "menu_source_url", "cuisine_tags",
    "dietary_tags",
]
# Workbook dish columns; food_type_tag has no dishes column and is not loaded.
DISH_SHEET_COLUMNS = [
    "id", "venue_id", "building_id", "name", "description", "price", "price_range",
    "image_url", "media_approved_for_display", "menu_approved_for_display",
    "media_rights_status", "menu_rights_status", "menu_source_url", "source_type",
    "food_type_tag", "updated_at",
]
DISH_UNMAPPED = ("food_type_tag",)
DISH_COLUMNS = [c for c in DISH_SHEET_COLUMNS if c not in DISH_UNMAPPED]
BUILDING_COLUMNS = [
    "id", "name", "building_type", "latitude", "longitude", "address", "postal_code",
    "planning_area", "venue_count", "source_type", "source_confidence",
]
# NOT NULL columns with a '' default: workbook nulls are written as ''.
VENUE_EMPTY_STRING_COLUMNS = {"unit_number", "operating_hours"}
DISH_EMPTY_STRING_COLUMNS = {"price_range"}
# NOT NULL without a usable default; a null here is loaded as '' so the quality
# gate (blank_*_name) reports it instead of the load aborting.
NAME_COLUMNS = {"name"}
BOOL_COLUMNS = {"media_approved_for_display", "menu_approved_for_display"}
ARRAY_COLUMNS = {"cuisine_tags", "dietary_tags"}

STALL_TYPE_TO_BUILDING_TYPE = {
    "restaurant": "restaurant",
    "chain outlet": "chain_outlet",
    "cafe": "cafe",
    "bar / bistro": "bar_bistro",
    "hawker stall": "hawker_centre",
    "food court stall": "food_court",
    "zi char": "restaurant",
    "other": "other",
}
BUILDING_TYPE_KEYWORDS = [
    ("hawker_centre", re.compile(r"hawker|food centre|market (and|&) food|food village", re.I)),
    ("food_court", re.compile(r"food ?court|food junction|koufu|kopitiam food hall", re.I)),
    ("kopitiam", re.compile(r"kopitiam|coffee ?shop|eating house", re.I)),
    ("mall", re.compile(r"\bmall\b|shopping centre|\bplaza\b", re.I)),
]

# Python ports of the SQL helpers in
# supabase/migrations/20260724113000_catalog_raw_publish_quality_gate.sql.
MOCK_ID_RE = re.compile(r"(^|[-_:])(mock|test|fixture|demo|sample)([-_:]|$)", re.I)
NON_CONSUMER_RULES = [
    ("convenience_store", re.compile(
        r"(^|[^a-z0-9])(7[- ]?eleven|cheers|fairprice xpress|shell select|caltex star mart|spc.?go)([^a-z0-9]|$)")),
    ("education_site", re.compile(
        r"(school|primary school|secondary school|junior college|polytechnic|university|kindergarten|child.?care|student care)")),
    ("private_or_restricted_site", re.compile(
        r"(staff canteen|employee canteen|workers? canteen|army camp|military camp|air base|naval base|police station|prison)")),
    ("non_consumer_business", re.compile(
        r"(food processing|manufactur|warehouse|cold store|distribution centre|distribution center|factory|industrial)")),
]
NORMALIZE_RE = re.compile(r"[\s!-/:-@\[-`{-~]+")

log = logging.getLogger("load_workbook_catalog_snapshot")


class LoaderError(RuntimeError):
    pass


# --------------------------------------------------------------------------- DB target


def read_db_url(db_url_file: str | None) -> str:
    if db_url_file:
        path = Path(db_url_file).expanduser()
        if not path.is_file():
            raise LoaderError(f"--db-url-file not found: {path}")
        value = path.read_text(encoding="utf-8").strip()
    else:
        value = os.environ.get(DB_URL_ENV, "").strip()
    if not value:
        raise LoaderError(f"no database URL: pass --db-url-file or set {DB_URL_ENV}")
    return value


def allowed_hosts() -> set[str]:
    extra = os.environ.get(ALLOWED_HOSTS_ENV, "")
    return set(LOCAL_HOSTS) | {h.strip().lower() for h in extra.split(",") if h.strip()}


def pg_env_from_url(db_url: str) -> dict[str, str]:
    """Validate the target host and translate the URL into PG* env vars.

    Error messages name the host only; the URL/password are never echoed.
    """
    parsed = urlparse(db_url)
    if parsed.scheme not in {"postgres", "postgresql"}:
        raise LoaderError("database URL must use the postgres:// or postgresql:// scheme")
    host = (parsed.hostname or "").lower()
    if not host:
        raise LoaderError("database URL has no host")
    if host not in allowed_hosts():
        raise LoaderError(
            f"refusing non-local database host {host}; add it to {ALLOWED_HOSTS_ENV} "
            "(comma-separated) to confirm the intended target"
        )
    env = {
        "PGHOST": host,
        "PGPORT": str(parsed.port or 5432),
        "PGDATABASE": unquote(parsed.path.lstrip("/")) or "postgres",
        "PGCONNECT_TIMEOUT": "10",
        "PGAPPNAME": "load_workbook_catalog_snapshot",
    }
    if parsed.username:
        env["PGUSER"] = unquote(parsed.username)
    if parsed.password:
        env["PGPASSWORD"] = unquote(parsed.password)
    for part in parsed.query.split("&"):
        if part.startswith("sslmode="):
            env["PGSSLMODE"] = part.split("=", 1)[1]
    return env


# --------------------------------------------------------------------------- coercion


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def coerce_postal(value: Any) -> tuple[str | None, bool]:
    """Excel stores some postal codes as numbers, dropping the leading zero."""
    if value is None or value == "":
        return None, False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value)).zfill(6), True
    return str(value).strip() or None, False


def coerce_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return int(round(float(value)))


def coerce_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return float(value)


def coerce_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()[:10] or None


def coerce_timestamp(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip() or None


def coerce_bool(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "t", "1", "yes", "y"}:
        return True
    if text in {"false", "f", "0", "no", "n"}:
        return False
    raise LoaderError(f"cannot parse boolean value {value!r}")


def coerce_tags(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        items = value
    else:
        text = str(value).strip()
        if text.startswith("["):
            items = json.loads(text)
        else:
            items = [part for part in re.split(r"[,;|]", text)]
    return [str(item).strip() for item in items if str(item).strip()]


def pg_array_literal(items: list[str]) -> str:
    escaped = ['"' + item.replace("\\", "\\\\").replace('"', '\\"') + '"' for item in items]
    return "{" + ",".join(escaped) + "}"


def is_placeholder(value: str | None) -> bool:
    return value is None or value.strip().lower() in PLACEHOLDER_VALUES


def normalize_identity(value: str | None) -> str:
    return NORMALIZE_RE.sub("", (value or "").strip().lower())


def non_consumer_reason(text: str) -> str | None:
    lowered = text.lower()
    for reason, pattern in NON_CONSUMER_RULES:
        if pattern.search(lowered):
            return reason
    return None


# --------------------------------------------------------------------------- workbook


def iter_sheet(workbook: Any, sheet: str, expected: list[str]) -> Iterator[dict[str, Any]]:
    if sheet not in workbook.sheetnames:
        raise LoaderError(f"workbook is missing sheet {sheet!r}")
    rows = workbook[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
    missing = [c for c in expected if c not in header]
    if missing:
        raise LoaderError(f"sheet {sheet!r} is missing columns: {missing}")
    extra = [c for c in header if c and c not in expected]
    if extra:
        log.warning("sheet %r has unmapped columns (ignored): %s", sheet, extra)
    index = {name: i for i, name in enumerate(header)}
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        yield {name: (row[i] if i < len(row) else None) for name, i in index.items()}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_workbook_rows(path: Path) -> tuple[list[dict], list[dict], dict[str, dict], Counter]:
    try:
        import openpyxl  # imported lazily so the guard helpers work without data deps
    except ImportError as exc:  # pragma: no cover - environment specific
        raise LoaderError("openpyxl is required (use the data venv)") from exc

    stats: Counter = Counter()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        venues: list[dict] = []
        for raw in iter_sheet(workbook, VENUES_SHEET, VENUE_COLUMNS):
            row: dict[str, Any] = {}
            for col in VENUE_COLUMNS:
                value = raw.get(col)
                if col == "postal_code":
                    value, padded = coerce_postal(value)
                    stats["venue_postal_zero_padded"] += int(padded)
                elif col in {"latitude", "longitude", "google_rating", "source_confidence"}:
                    value = coerce_float(value)
                elif col == "google_review_count":
                    value = coerce_int(value)
                elif col == "last_verified_at":
                    value = coerce_date(value)
                elif col == "updated_at":
                    value = coerce_timestamp(value)
                elif col in BOOL_COLUMNS:
                    value = coerce_bool(value)
                elif col in ARRAY_COLUMNS:
                    value = coerce_tags(value)
                else:
                    value = clean_text(value)
                    if value is None and (col in VENUE_EMPTY_STRING_COLUMNS or col in NAME_COLUMNS):
                        stats[f"venue_null_{col}_as_empty"] += 1
                        value = ""
                row[col] = value
            venues.append(row)

        dishes: list[dict] = []
        for raw in iter_sheet(workbook, DISHES_SHEET, DISH_SHEET_COLUMNS):
            row = {}
            for col in DISH_COLUMNS:
                value = raw.get(col)
                if col == "price":
                    value = coerce_float(value)
                    if value is None:
                        stats["dish_null_price_as_zero"] += 1
                        value = 0.0
                elif col == "updated_at":
                    value = coerce_timestamp(value)
                elif col in BOOL_COLUMNS:
                    value = coerce_bool(value)
                else:
                    value = clean_text(value)
                    if value is None and (col in DISH_EMPTY_STRING_COLUMNS or col in NAME_COLUMNS):
                        stats[f"dish_null_{col}_as_empty"] += 1
                        value = ""
                row[col] = value
            dishes.append(row)

        source_rows: dict[str, dict] = {}
        if SOURCE_SHEET in workbook.sheetnames:
            wanted = ["Business ID", "Hawker Centre / Building", "Building Name", "Stall Type"]
            for raw in iter_sheet(workbook, SOURCE_SHEET, wanted):
                business_id = clean_text(raw.get("Business ID"))
                if business_id:
                    source_rows[business_id] = {
                        "building_name": clean_text(raw.get("Building Name")),
                        "centre_name": clean_text(raw.get("Hawker Centre / Building")),
                        "stall_type": clean_text(raw.get("Stall Type")),
                    }
    finally:
        workbook.close()
    return venues, dishes, source_rows, stats


# --------------------------------------------------------------------------- validation


def validate(venues: list[dict], dishes: list[dict], manifest: dict | None) -> list[str]:
    errors: list[str] = []
    venue_ids = [v["id"] for v in venues]
    dish_ids = [d["id"] for d in dishes]
    if any(not vid for vid in venue_ids):
        errors.append("venue rows with blank id")
    if any(not did for did in dish_ids):
        errors.append("dish rows with blank id")
    if len(set(venue_ids)) != len(venue_ids):
        errors.append(f"duplicate venue ids: {len(venue_ids) - len(set(venue_ids))}")
    if len(set(dish_ids)) != len(dish_ids):
        errors.append(f"duplicate dish ids: {len(dish_ids) - len(set(dish_ids))}")
    if any(not v["building_id"] for v in venues):
        errors.append("venue rows with blank building_id")
    if any(not d["building_id"] for d in dishes):
        errors.append("dish rows with blank building_id")
    if any(not v["floor"] for v in venues):
        errors.append("venue rows with blank floor (venues.floor is NOT NULL)")
    known = set(venue_ids)
    orphans = sum(1 for d in dishes if d["venue_id"] not in known)
    if orphans:
        errors.append(f"dishes referencing unknown venue ids: {orphans}")
    bad_grade = sum(1 for v in venues if v["hygiene_grade"] not in {"A", "B", "C", "D"})
    if bad_grade:
        errors.append(f"venues with hygiene_grade outside A-D: {bad_grade}")
    for row in (*venues, *dishes):
        for value in row.values():
            if isinstance(value, str) and re.search(r"(^|\n)\\\.(\r?\n|$)", value):
                errors.append(f"row {row['id']} contains a COPY end-of-data marker line")
            if isinstance(value, str) and "\x00" in value:
                errors.append(f"row {row['id']} contains a NUL byte (not storable in Postgres text)")
    if manifest:
        if len(venues) != manifest.get("expectedVenueTotal"):
            errors.append(
                f"venue total {len(venues)} != manifest {manifest.get('expectedVenueTotal')}"
            )
        if len(dishes) != manifest.get("expectedDishTotal"):
            errors.append(
                f"dish total {len(dishes)} != manifest {manifest.get('expectedDishTotal')}"
            )
    return errors


# --------------------------------------------------------------------------- buildings


def slug_to_name(slug: str) -> str:
    words = [w for w in slug.split("-") if w]
    return " ".join(w.capitalize() for w in words)


def name_from_building_id(building_id: str) -> tuple[str | None, str | None]:
    """Return (name, postal) encoded in readable building ids, if any."""
    match = re.fullmatch(r"rw-bld-(.+?)(?:-(\d{6}))?", building_id)
    if match:
        slug = re.sub(r"^part-of-", "", match.group(1))
        return slug_to_name(slug), match.group(2)
    match = re.fullmatch(r"building-(.+)", building_id)
    if match and building_id != UNASSIGNED_BUILDING_ID:
        return slug_to_name(match.group(1)), None
    return None, None


def most_common(values: Iterable[str | None], skip_placeholders: bool = True) -> str | None:
    counter = Counter(v for v in values if v and not (skip_placeholders and is_placeholder(v)))
    if not counter:
        return None
    # Deterministic tie-break: highest count, then lexical order.
    return sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]


def infer_building_type(name: str, source_names: list[str], stall_types: list[str]) -> tuple[str, str]:
    for building_type, pattern in BUILDING_TYPE_KEYWORDS:
        if pattern.search(name):
            return building_type, "name_keyword"
    named = [n for n in source_names if n]
    if named and all(n.strip().lower() == "standalone" for n in named):
        return "standalone", "source_standalone"
    mapped = [STALL_TYPE_TO_BUILDING_TYPE.get(s.strip().lower()) for s in stall_types if s]
    mapped = [m for m in mapped if m]
    if mapped:
        return Counter(mapped).most_common(1)[0][0], "source_stall_type_majority"
    return "other", "default_other"


def in_singapore(lat: float | None, lon: float | None) -> bool:
    return (
        lat is not None and lon is not None
        and SG_LAT_RANGE[0] <= lat <= SG_LAT_RANGE[1]
        and SG_LON_RANGE[0] <= lon <= SG_LON_RANGE[1]
    )


def derive_buildings(
    venues: list[dict], dishes: list[dict], source_rows: dict[str, dict]
) -> tuple[list[dict], dict[str, Any]]:
    by_building: dict[str, list[dict]] = defaultdict(list)
    for venue in venues:
        by_building[venue["building_id"]].append(venue)
    venue_by_id = {v["id"]: v for v in venues}
    # Dishes may reference a building that no venue uses; position those from
    # the dish's own venue so the FK resolves without touching dish rows.
    dish_only: dict[str, list[dict]] = defaultdict(list)
    for dish in dishes:
        if dish["building_id"] not in by_building:
            venue = venue_by_id.get(dish["venue_id"])
            if venue:
                dish_only[dish["building_id"]].append(venue)

    name_sources: Counter = Counter()
    type_sources: Counter = Counter()
    coord_sources: Counter = Counter()
    buildings: list[dict] = []
    for building_id in sorted(set(by_building) | set(dish_only)):
        members = by_building.get(building_id) or dish_only[building_id]
        joined = [source_rows[v["business_id"]] for v in members if v.get("business_id") in source_rows]
        source_names = [j["building_name"] or j["centre_name"] for j in joined]
        address = most_common(v["address"] for v in members)
        postal = most_common(v["postal_code"] for v in members)
        id_name, id_postal = name_from_building_id(building_id)

        if building_id == UNASSIGNED_BUILDING_ID:
            name, name_source = "Unassigned location", "unassigned_placeholder"
        elif id_name:
            name, name_source = id_name, "building_id_slug"
        elif most_common(source_names):
            name, name_source = most_common(source_names), "source_sheet_building_name"
        elif address:
            name, name_source = address, "venue_address_mode"
        elif postal or id_postal:
            name, name_source = f"Singapore {postal or id_postal}", "postal_code"
        else:
            name, name_source = most_common(v["name"] for v in members) or building_id, "venue_name_mode"
        name_sources[name_source] += 1

        building_type, type_source = infer_building_type(
            name, source_names, [j["stall_type"] for j in joined]
        )
        if building_id == UNASSIGNED_BUILDING_ID:
            building_type, type_source = "other", "default_other"
        type_sources[type_source] += 1

        coords = [(v["latitude"], v["longitude"]) for v in members if in_singapore(v["latitude"], v["longitude"])]
        if coords:
            lat = sum(c[0] for c in coords) / len(coords)
            lon = sum(c[1] for c in coords) / len(coords)
            confidence = 0.5
            coord_sources["venue_mean" if building_id in by_building else "dish_venue_mean"] += 1
        else:
            lat, lon = SG_CENTROID
            confidence = 0.0
            coord_sources["sg_centroid_fallback"] += 1

        buildings.append({
            "id": building_id,
            "name": name,
            "building_type": building_type,
            "latitude": round(lat, 8),
            "longitude": round(lon, 8),
            "address": address,
            "postal_code": postal or id_postal,
            "planning_area": most_common(v["planning_area"] for v in members),
            "venue_count": len(by_building.get(building_id, [])),
            "source_type": BUILDING_SOURCE_TYPE,
            "source_confidence": confidence,
            "_name_source": name_source,
        })
    summary = {
        "buildings": len(buildings),
        "dish_only_buildings": len(dish_only),
        "name_sources": dict(name_sources),
        "building_type_sources": dict(type_sources),
        "building_types": dict(Counter(b["building_type"] for b in buildings)),
        "coordinate_sources": dict(coord_sources),
    }
    return buildings, summary


# --------------------------------------------------------------------------- gate estimate


def estimate_quality_findings(buildings: list[dict], venues: list[dict], dishes: list[dict]) -> dict[str, Any]:
    """Local estimate of public.catalog_raw_publish_quality_findings().

    manual_exclusion needs catalog_raw_excluded_venues and is not estimated.
    Normalization approximates the SQL [[:space:][:punct:]] class with ASCII.
    """
    building_name = {b["id"]: b["name"] for b in buildings}
    building_source = {b["id"]: b["_name_source"] for b in buildings}
    findings: dict[str, list] = defaultdict(list)
    loader_caused_non_consumer: list = []
    loader_caused_by_source: Counter = Counter()

    for b in buildings:
        if MOCK_ID_RE.search(b["id"]):
            findings["mock_building_id"].append([b["id"], b["name"]])
    for v in venues:
        if MOCK_ID_RE.search(v["id"]):
            findings["mock_venue_id"].append([v["id"], v["name"]])
        if not v["name"].strip():
            findings["blank_venue_name"].append([v["id"], v["name"]])
        combined = " ".join(p for p in (v["name"], building_name.get(v["building_id"])) if p)
        reason = non_consumer_reason(combined)
        if reason:
            findings["non_consumer_venue"].append([v["id"], v["name"], building_name.get(v["building_id"]), reason])
            if non_consumer_reason(v["name"]) is None:
                loader_caused_non_consumer.append([v["id"], v["name"], building_name.get(v["building_id"]), reason])
                loader_caused_by_source[building_source.get(v["building_id"])] += 1
    for d in dishes:
        if MOCK_ID_RE.search(d["id"]):
            findings["mock_dish_id"].append([d["id"], d["name"]])
        if not d["name"].strip():
            findings["blank_dish_name"].append([d["id"], d["venue_id"]])

    venue_groups: dict[tuple, list] = defaultdict(list)
    for v in venues:
        if v["name"].strip() and v["unit_number"].strip():
            key = (v["building_id"], normalize_identity(v["name"]), normalize_identity(v["unit_number"]))
            venue_groups[key].append(v)
    dup_venue_rows = 0
    for group in venue_groups.values():
        if len(group) > 1:
            dup_venue_rows += len(group)
            findings["semantic_duplicate_venue"].append(
                [[g["id"] for g in group][:4], group[0]["name"], group[0]["unit_number"], group[0]["building_id"]]
            )
    dish_groups: dict[tuple, list] = defaultdict(list)
    for d in dishes:
        if d["name"].strip():
            dish_groups[(d["venue_id"], normalize_identity(d["name"]), d["price"])].append(d)
    dup_dish_rows = 0
    for group in dish_groups.values():
        if len(group) > 1:
            dup_dish_rows += len(group)
            findings["semantic_duplicate_dish"].append(
                [[g["id"] for g in group][:4], group[0]["name"], group[0]["venue_id"]]
            )

    dup_venue_unassigned = sum(1 for f in findings["semantic_duplicate_venue"] if f[3] == UNASSIGNED_BUILDING_ID)
    return {
        "counts": {code: len(rows) for code, rows in sorted(findings.items())},
        "examples": {code: rows[:5] for code, rows in sorted(findings.items())},
        "semantic_duplicate_rows": {"venue": dup_venue_rows, "dish": dup_dish_rows},
        "semantic_duplicate_venue_groups_in_unassigned_building": dup_venue_unassigned,
        "non_consumer_caused_by_derived_building_name": {
            "count": len(loader_caused_non_consumer),
            "by_name_source": dict(loader_caused_by_source),
            "examples": loader_caused_non_consumer[:5],
        },
        "not_estimated": ["manual_exclusion (needs catalog_raw_excluded_venues)"],
    }


def placeholder_name_report(buildings: list[dict]) -> dict[str, Any]:
    generic = re.compile(r"^(local food court|food court|coffee ?shop|kopitiam|unassigned|singapore \d{6})", re.I)
    flagged = [b for b in buildings if generic.search(b["name"]) or b["name"] == b["id"]]
    return {
        "generic_or_placeholder_names": len(flagged),
        "examples": [[b["id"], b["name"], b["venue_count"]] for b in flagged[:5]],
        "unassigned_building_venue_count": next(
            (b["venue_count"] for b in buildings if b["id"] == UNASSIGNED_BUILDING_ID), 0
        ),
    }


# --------------------------------------------------------------------------- COPY stream


def csv_field(value: Any) -> str:
    if value is None:
        return ""  # unquoted empty = NULL in COPY CSV
    if isinstance(value, bool):
        text = "t" if value else "f"
    elif isinstance(value, list):
        text = pg_array_literal(value)
    else:
        text = str(value)
    return '"' + text.replace('"', '""') + '"'  # quoted, so '' stays an empty string


def csv_lines(rows: Iterable[dict], columns: list[str]) -> Iterator[str]:
    for row in rows:
        yield ",".join(csv_field(row.get(col)) for col in columns) + "\n"


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def build_preamble(replace: bool) -> str:
    exempt = ", ".join(sql_literal(r) for r in TRIGGER_EXEMPT_ROLES)
    lines = [
        "\\set ON_ERROR_STOP 1",
        "set client_encoding = 'UTF8';",
        "begin;",
        "set local statement_timeout = 0;",
        "set local lock_timeout = '15s';",
        "do $$ begin",
        f"  if current_user not in ({exempt}) then",
        "    raise exception 'loader must run as a role exempt from enforce_merchant_dish_write_guard (got %)', current_user;",
        "  end if;",
        "end $$;",
        "lock table public.buildings, public.venues, public.dishes in share row exclusive mode;",
    ]
    if replace:
        # Delete only these three tables, and only when nothing else references
        # their rows: ON DELETE CASCADE / SET NULL must not reach other tables.
        lines += [
            "do $$",
            "declare r record; n bigint;",
            "begin",
            "  for r in",
            "    select c.conrelid::regclass as child, a.attname as col",
            "    from pg_constraint c",
            "    join pg_attribute a on a.attrelid = c.conrelid and a.attnum = c.conkey[1]",
            "    where c.contype = 'f'",
            "      and c.confrelid in ('public.buildings'::regclass, 'public.venues'::regclass, 'public.dishes'::regclass)",
            "      and c.conrelid not in ('public.buildings'::regclass, 'public.venues'::regclass, 'public.dishes'::regclass)",
            "  loop",
            "    execute format('select count(*) from %s where %I is not null', r.child, r.col) into n;",
            "    if n > 0 then",
            "      raise exception '--replace refused: % rows in % reference catalog tables', n, r.child;",
            "    end if;",
            "  end loop;",
            "end $$;",
            "delete from public.dishes;",
            "delete from public.venues;",
            "delete from public.buildings;",
        ]
    else:
        lines += [
            "do $$ begin",
            "  if exists (select 1 from public.buildings) or exists (select 1 from public.venues)",
            "     or exists (select 1 from public.dishes) then",
            "    raise exception 'buildings/venues/dishes are not empty; rerun with --replace to reload';",
            "  end if;",
            "end $$;",
        ]
    return "\n".join(lines) + "\n"


def build_postamble(n_buildings: int, n_venues: int, n_dishes: int) -> str:
    return (
        "do $$ begin\n"
        f"  if (select count(*) from public.buildings) <> {n_buildings}\n"
        f"     or (select count(*) from public.venues) <> {n_venues}\n"
        f"     or (select count(*) from public.dishes) <> {n_dishes} then\n"
        "    raise exception 'post-load row counts do not match the workbook';\n"
        "  end if;\n"
        "end $$;\n"
        "commit;\n"
    )


SAFE_COPY_PATH = re.compile(r"^[A-Za-z0-9_./ -]+$")


def write_copy_file(path: Path, columns: list[str], rows: Iterable[dict]) -> None:
    """Write header-less COPY CSV data; values never enter the psql script."""
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.writelines(csv_lines(rows, columns))


def copy_block(table: str, columns: list[str], path: Path) -> str:
    # The path is loader-created (mkdtemp), never data-derived; still refuse
    # anything that could break out of the quoted \copy filename.
    if not SAFE_COPY_PATH.match(str(path)):
        raise LoaderError(f"unsafe temp path for \\copy: {path!r}")
    return f"\\copy public.{table} ({', '.join(columns)}) from {sql_literal(str(path))} with (format csv)\n"


def load_script(replace: bool, paths: dict[str, Path], counts: dict[str, int]) -> Iterator[str]:
    """Yield the full psql script: one transaction, buildings -> venues -> dishes.

    Contains only fixed SQL and \\copy lines naming temp files; no row data.
    """
    yield build_preamble(replace)
    yield copy_block("buildings", BUILDING_COLUMNS, paths["buildings"])
    yield copy_block("venues", VENUE_COLUMNS, paths["venues"])
    yield copy_block("dishes", DISH_COLUMNS, paths["dishes"])
    yield build_postamble(counts["buildings"], counts["venues"], counts["dishes"])


def load_via_temp_files(
    env: dict[str, str], buildings: list[dict], venues: list[dict], dishes: list[dict], replace: bool,
) -> None:
    with tempfile.TemporaryDirectory(prefix="hungry-catalog-load-") as tmp:  # mode 0700
        tmp_dir = Path(tmp)
        paths: dict[str, Path] = {}
        for name, columns, rows in (
            ("buildings", BUILDING_COLUMNS, buildings),
            ("venues", VENUE_COLUMNS, venues),
            ("dishes", DISH_COLUMNS, dishes),
        ):
            paths[name] = tmp_dir / f"{name}.csv"
            write_copy_file(paths[name], columns, rows)
            log.info("staged %s (%d bytes)", name, paths[name].stat().st_size)
        counts = {"buildings": len(buildings), "venues": len(venues), "dishes": len(dishes)}
        run_psql_stream(env, load_script(replace, paths, counts))


def run_psql_stream(env: dict[str, str], chunks: Iterable[str]) -> None:
    proc = subprocess.Popen(
        ["psql", "-X", "-q", "-v", "ON_ERROR_STOP=1", "-f", "-"],
        stdin=subprocess.PIPE,
        env={**os.environ, **env},
        text=True,
        encoding="utf-8",
    )
    assert proc.stdin is not None
    try:
        for chunk in chunks:
            proc.stdin.write(chunk)
    except BrokenPipeError:
        pass  # psql stopped early (ON_ERROR_STOP); its exit code reports why
    finally:
        try:
            proc.stdin.close()
        except BrokenPipeError:
            pass
    if proc.wait() != 0:
        raise LoaderError(f"psql exited with code {proc.returncode}; transaction rolled back")


def run_psql_query(env: dict[str, str], sql: str) -> str:
    result = subprocess.run(
        ["psql", "-X", "-At", "-v", "ON_ERROR_STOP=1", "-c", sql],
        env={**os.environ, **env}, capture_output=True, text=True, check=False,
    )
    if result.returncode != 0:
        raise LoaderError(f"psql query failed: {result.stderr.strip()}")
    return result.stdout.strip()


def write_csvs(csv_dir: Path, buildings: list[dict], venues: list[dict], dishes: list[dict]) -> None:
    csv_dir.mkdir(parents=True, exist_ok=True)
    for name, columns, rows in (
        ("buildings", BUILDING_COLUMNS, buildings),
        ("venues", VENUE_COLUMNS, venues),
        ("dishes", DISH_COLUMNS, dishes),
    ):
        path = csv_dir / f"{name}.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            handle.write(",".join(columns) + "\n")
            handle.writelines(csv_lines(rows, columns))
        log.info("wrote %s", path)


# --------------------------------------------------------------------------- main


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true", help="validate and report only; no files, no DB writes")
    mode.add_argument("--apply", action="store_true", help="load into the database in one transaction")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--db-url-file", help=f"file containing the DB URL (else ${DB_URL_ENV})")
    parser.add_argument("--replace", action="store_true",
                        help="delete existing buildings/venues/dishes rows first (refused if other tables reference them)")
    parser.add_argument("--csv-dir", type=Path, help="optionally also write buildings/venues/dishes CSVs here")
    parser.add_argument("--skip-checksum", action="store_true", help="do not compare the workbook to the manifest checksum")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", stream=sys.stderr)
    args = parse_args(argv)
    timings: dict[str, float] = {}

    # Resolve and guard the DB target before any expensive work.
    pg_env: dict[str, str] | None = None
    if args.apply:
        pg_env = pg_env_from_url(read_db_url(args.db_url_file))
        log.info("target host %s:%s database %s", pg_env["PGHOST"], pg_env["PGPORT"], pg_env["PGDATABASE"])

    manifest = json.loads(args.manifest.read_text(encoding="utf-8")) if args.manifest.is_file() else None
    if not args.workbook.is_file():
        raise LoaderError(f"workbook not found: {args.workbook}")
    if manifest and not args.skip_checksum:
        started = time.monotonic()
        checksum = sha256_file(args.workbook)
        timings["checksum_s"] = round(time.monotonic() - started, 2)
        if checksum != manifest.get("sourceChecksumSha256"):
            raise LoaderError("workbook checksum does not match the release manifest")
        log.info("workbook checksum matches manifest %s", manifest.get("manifestVersion"))

    started = time.monotonic()
    venues, dishes, source_rows, coercions = load_workbook_rows(args.workbook)
    timings["read_workbook_s"] = round(time.monotonic() - started, 2)
    log.info("read %d venues, %d dishes, %d source rows", len(venues), len(dishes), len(source_rows))

    errors = validate(venues, dishes, manifest)
    if errors:
        for error in errors:
            log.error("validation: %s", error)
        raise LoaderError(f"{len(errors)} validation error(s); nothing loaded")

    started = time.monotonic()
    buildings, building_summary = derive_buildings(venues, dishes, source_rows)
    timings["derive_buildings_s"] = round(time.monotonic() - started, 2)

    report: dict[str, Any] = {
        "mode": "apply" if args.apply else "dry-run",
        "manifest_version": manifest.get("manifestVersion") if manifest else None,
        "rows": {"buildings": len(buildings), "venues": len(venues), "dishes": len(dishes)},
        "coercions": dict(coercions),
        "unmapped_workbook_columns": {"Supabase Dishes": list(DISH_UNMAPPED), "Supabase Venues": []},
        "building_derivation": building_summary,
        "building_name_quality": placeholder_name_report(buildings),
    }
    if args.dry_run:
        started = time.monotonic()
        report["quality_gate_estimate"] = estimate_quality_findings(buildings, venues, dishes)
        timings["estimate_s"] = round(time.monotonic() - started, 2)

    if args.csv_dir:
        write_csvs(args.csv_dir, buildings, venues, dishes)

    if args.apply and pg_env is not None:
        started = time.monotonic()
        load_via_temp_files(pg_env, buildings, venues, dishes, args.replace)
        timings["load_s"] = round(time.monotonic() - started, 2)
        counts = run_psql_query(
            pg_env,
            "select (select count(*) from public.buildings) || ',' || (select count(*) from public.venues)"
            " || ',' || (select count(*) from public.dishes) || ',' || pg_database_size(current_database())",
        ).split(",")
        report["loaded"] = {
            "buildings": int(counts[0]), "venues": int(counts[1]), "dishes": int(counts[2]),
            "pg_database_size_bytes": int(counts[3]),
        }
        log.info("load committed")

    report["timings"] = timings
    print(json.dumps(report, indent=2, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except LoaderError as exc:
        logging.getLogger("load_workbook_catalog_snapshot").error("%s", exc)
        sys.exit(2)
