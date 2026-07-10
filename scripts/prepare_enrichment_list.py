#!/usr/bin/env python3
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "Hungry F&B Database.xlsx"
SCRAPE_RESULTS_PATH = ROOT / "data" / "scrape_results.json"
OUTPUT_PATH = ROOT / "data" / "stalls_to_enrich.json"

def clean(val):
    if val is None:
        return ""
    return str(val).strip()

def main():
    if not SCRAPE_RESULTS_PATH.exists():
        print(f"Scrape results not found at {SCRAPE_RESULTS_PATH}. Run scrape_maps.js first.")
        return

    print("Loading scrape results...")
    with open(SCRAPE_RESULTS_PATH, "r", encoding="utf-8") as f:
        scrape_results = json.load(f)

    # Build map of row_index -> result
    res_map = {r["row_index"]: r for r in scrape_results}

    print("Loading Excel workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    headers = [clean(cell.value) for cell in ws[1]]
    name_idx = headers.index("Stall / Restaurant Name")
    building_idx = headers.index("Hawker Centre / Building")
    address_idx = headers.index("Address")
    lat_idx = headers.index("Latitude")
    lng_idx = headers.index("Longitude")
    uri_idx = headers.index("Google Maps URI")
    menu_idx = headers.index("Menu")
    img_idx = headers.index("Dishes Image URL")

    # Group valid unique CIDs
    unique_stalls = {}
    total_rows = 0
    active_rows = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        total_rows += 1

        # Check if row is valid and open according to scraper results
        res = res_map.get(idx)
        if not res:
            continue

        is_valid = res.get("isValid") == "Yes"
        is_closed = res.get("isClosed") == "Yes"

        if is_valid and not is_closed:
            active_rows += 1
            uri = clean(row[uri_idx])
            if not uri:
                continue

            # Keep the first occurrence of each unique CID
            if uri not in unique_stalls:
                unique_stalls[uri] = {
                    "row_index": idx,
                    "uri": uri,
                    "stall_name": clean(row[name_idx]),
                    "building": clean(row[building_idx]),
                    "address": clean(row[address_idx]),
                    "latitude": row[lat_idx],
                    "longitude": row[lng_idx],
                    "existing_menu": clean(row[menu_idx]),
                    "existing_images": clean(row[img_idx])
                }

    print(f"Total rows in Excel: {total_rows}")
    print(f"Active valid rows: {active_rows}")
    print(f"Unique active stalls (by CID): {len(unique_stalls)}")

    # Save to file
    stalls_list = list(unique_stalls.values())
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(stalls_list, f, indent=2)

    print(f"Saved {len(stalls_list)} unique active stalls to {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
