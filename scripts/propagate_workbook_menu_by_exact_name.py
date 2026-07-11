#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it in the active Python environment, "
        "or run with the Codex workspace Python runtime used for workbook tooling."
    ) from exc


LEGAL_SUFFIX_RE = re.compile(
    r"\b(?:pte\.?\s*ltd\.?|private\s+limited|limited|ltd\.?|llp|inc\.?|co\.?)\b",
    re.IGNORECASE,
)
TRUSTED_PREFIXES = (
    "source_description:",
    "curated_seed:",
    "online_price_menu:",
    "official_online_seed:",
    "online_reference_seed:",
)
GENERIC_NAMES = {
    "bakery",
    "cafe",
    "chicken rice",
    "coffee",
    "coffee shop",
    "drinks",
    "food court",
    "juice",
    "mixed rice",
    "restaurant",
    "seafood",
    "tea",
    "vegetarian",
    "western food",
}


@dataclass(frozen=True)
class TrustedMenu:
    source_row: int
    signature: str
    menu: str
    menu_items: str
    menu_source: str


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return clean(value) in {"", "-", "—", "N/A", "NA", "None", "null"}


def normalize_name(value: Any) -> str:
    text = clean(value).lower().replace("&", " and ")
    text = LEGAL_SUFFIX_RE.sub(" ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def load_headers(sheet) -> dict[str, int]:
    return {clean(cell.value): index for index, cell in enumerate(sheet[1], start=1)}


def is_generic_name(normalized_name: str) -> bool:
    return normalized_name in GENERIC_NAMES or len(normalized_name) < 5 or len(normalized_name.split()) < 2


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100000)
    parser.add_argument("--max-updates", type=int, default=500)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet]
    headers = load_headers(ws)
    required = ["Stall / Restaurant Name", "Signature Dish", "Menu", "Menu Items", "Menu Source"]
    missing = [field for field in required if field not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    trusted_by_name: dict[str, list[TrustedMenu]] = defaultdict(list)
    for row_num in range(2, ws.max_row + 1):
        normalized_name = normalize_name(ws.cell(row_num, headers["Stall / Restaurant Name"]).value)
        menu_items = clean(ws.cell(row_num, headers["Menu Items"]).value)
        menu_source = clean(ws.cell(row_num, headers["Menu Source"]).value)
        if not normalized_name or not menu_items or not menu_source.startswith(TRUSTED_PREFIXES):
            continue
        trusted_by_name[normalized_name].append(
            TrustedMenu(
                source_row=row_num,
                signature=clean(ws.cell(row_num, headers["Signature Dish"]).value),
                menu=clean(ws.cell(row_num, headers["Menu"]).value),
                menu_items=menu_items,
                menu_source=menu_source,
            )
        )

    unique_trusted: dict[str, TrustedMenu] = {}
    for normalized_name, menus in trusted_by_name.items():
        menu_item_sets = {menu.menu_items for menu in menus}
        if len(menu_item_sets) == 1:
            unique_trusted[normalized_name] = menus[0]

    updates: list[dict[str, Any]] = []
    last_row = min(ws.max_row, args.start_row + args.max_rows - 1)
    for row_num in range(args.start_row, last_row + 1):
        normalized_name = normalize_name(ws.cell(row_num, headers["Stall / Restaurant Name"]).value)
        if is_generic_name(normalized_name) or normalized_name not in unique_trusted:
            continue
        if not is_blank(ws.cell(row_num, headers["Menu"]).value):
            continue
        if not is_blank(ws.cell(row_num, headers["Menu Items"]).value):
            continue
        trusted = unique_trusted[normalized_name]
        if row_num == trusted.source_row:
            continue
        update = {
            "row": row_num,
            "name": ws.cell(row_num, headers["Stall / Restaurant Name"]).value,
            "source_row": trusted.source_row,
            "signature": trusted.signature,
            "menu_items": trusted.menu_items,
            "menu_source": trusted.menu_source,
        }
        updates.append(update)
        if args.apply:
            if is_blank(ws.cell(row_num, headers["Signature Dish"]).value):
                ws.cell(row_num, headers["Signature Dish"]).value = trusted.signature
            ws.cell(row_num, headers["Menu"]).value = trusted.menu or trusted.menu_items
            ws.cell(row_num, headers["Menu Items"]).value = trusted.menu_items
            ws.cell(row_num, headers["Menu Source"]).value = (
                f"name_propagated: row {trusted.source_row}; {trusted.menu_source}"
            )
        if len(updates) >= args.max_updates:
            break

    backup_path = None
    if args.apply:
        if args.backup:
            stamp = datetime.now().strftime("%Y%m%d%H%M%S")
            backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
            shutil.copy2(workbook_path, backup_path)
        wb.save(workbook_path)

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "sheet": args.sheet,
                "apply": args.apply,
                "backup": str(backup_path) if backup_path else None,
                "updated": len(updates),
                "updates": updates,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
