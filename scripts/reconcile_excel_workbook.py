#!/usr/bin/env python3
"""Create a deterministic, workbook-only reconciliation candidate.

The utility never contacts Supabase and refuses to overwrite its input. Records
that cannot be reconciled without choosing between conflicting facts are kept in
dedicated quarantine sheets instead of being discarded or guessed.
"""

from __future__ import annotations

import argparse
import os
import re
from collections import OrderedDict, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from data_display_policy import approved_rights_status, is_placeholder_url, truthy
from export_excel_catalog_for_app import (
    CatalogValidationError,
    source_match_keys,
    validate_export_inputs,
)
from validate_excel_workbook import DEFAULT_WORKBOOK, dedupe_key


SOURCE_SHEET = "Sheet1"
VENUE_SHEET = "Supabase Venues"
DISH_SHEET = "Supabase Dishes"
REPORT_SHEET = "Reconciliation Report"
QUARANTINE_SHEETS = {
    SOURCE_SHEET: "Quarantine Sheet1",
    VENUE_SHEET: "Quarantine Venues",
    DISH_SHEET: "Quarantine Dishes",
}
SOURCE_MEDIA_FIELDS = ("Image URL", "Dishes Image URL", "Logo URL")
SOURCE_MENU_FIELDS = ("Menu", "Menu Items", "Menu Source", "Signature Dish")
UNVERIFIED_SOURCE_MENU_PREFIXES = (
    "cuisine category seed",
    "category_seed:",
    "claude ai fallback",
    "online_reference_seed:",
    "online_menu_candidate:",
    "source_description:",
    "grabfood scrape",
    "foodpanda",
    "deliveroo",
)
CACHE_MEDIA_FIELDS = {
    VENUE_SHEET: ("image_url", "logo_url"),
    DISH_SHEET: ("image_url",),
}
QUARANTINE_METADATA = ("_reason_codes", "_action", "_source_row_number")


class ReconciliationError(ValueError):
    """Raised when no safe deterministic candidate can be produced."""


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def comparison_value(value: Any) -> tuple[str, Any] | None:
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return ("datetime", value.isoformat())
    if isinstance(value, date):
        return ("date", value.isoformat())
    if isinstance(value, bool):
        return ("bool", value)
    if isinstance(value, (int, float, Decimal)):
        return ("number", str(value))
    return ("text", clean_text(value))


def read_sheet_rows(workbook: Any, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if sheet_name not in workbook.sheetnames:
        raise ReconciliationError(f"Workbook is missing required sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_values is None:
        raise ReconciliationError(f"Workbook sheet {sheet_name!r} has no header row")
    headers = [clean_text(value) for value in header_values]
    if not any(headers):
        raise ReconciliationError(f"Workbook sheet {sheet_name!r} has no headers")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(not is_blank(value) for value in values):
            continue
        row = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        row["_row_number"] = row_number
        rows.append(row)
    return [header for header in headers if header], rows


def read_optional_sheet_rows(workbook: Any, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if sheet_name not in workbook.sheetnames:
        return [], []
    return read_sheet_rows(workbook, sheet_name)


def merge_compatible_rows(
    rows: list[dict[str, Any]],
    headers: list[str],
) -> tuple[dict[str, Any], list[str]]:
    merged = {header: rows[0].get(header) for header in headers}
    merged["_row_number"] = rows[0]["_row_number"]
    conflicts: list[str] = []
    for header in headers:
        distinct: OrderedDict[tuple[str, Any], Any] = OrderedDict()
        for row in rows:
            value = row.get(header)
            normalized = comparison_value(value)
            if normalized is not None and normalized not in distinct:
                distinct[normalized] = value
        if len(distinct) > 1:
            conflicts.append(header)
        elif distinct:
            merged[header] = next(iter(distinct.values()))
        else:
            merged[header] = None
    return merged, conflicts


def quarantine_copy(
    row: dict[str, Any],
    headers: list[str],
    reasons: Iterable[str],
    action: str,
) -> dict[str, Any]:
    copied = {header: row.get(header) for header in headers}
    copied["_reason_codes"] = ";".join(dict.fromkeys(reasons))
    copied["_action"] = action
    copied["_source_row_number"] = row.get("_row_number")
    return copied


def sanitize_display_fields(
    row: dict[str, Any],
    headers: list[str],
    sheet_name: str,
) -> tuple[dict[str, Any], dict[str, Any] | None, list[str]]:
    sanitized = dict(row)
    original = dict(row)
    reasons: list[str] = []
    media_fields = (
        SOURCE_MEDIA_FIELDS if sheet_name == SOURCE_SHEET else CACHE_MEDIA_FIELDS[sheet_name]
    )
    for field in media_fields:
        if is_placeholder_url(sanitized.get(field)):
            sanitized[field] = None
            reasons.append(f"placeholder_media:{field}")

    if sheet_name == SOURCE_SHEET:
        menu_source = clean_text(sanitized.get("Menu Source")).casefold()
        if menu_source.startswith(UNVERIFIED_SOURCE_MENU_PREFIXES):
            for field in SOURCE_MENU_FIELDS:
                if field in sanitized:
                    sanitized[field] = None
            reasons.append("unverified_source_menu")

    if sheet_name in (VENUE_SHEET, DISH_SHEET):
        for kind, approved_field, rights_field in (
            ("menu", "menu_approved_for_display", "menu_rights_status"),
            ("media", "media_approved_for_display", "media_rights_status"),
        ):
            if truthy(sanitized.get(approved_field)) and not approved_rights_status(
                sanitized.get(rights_field)
            ):
                sanitized[approved_field] = False
                reasons.append(f"approval_rights_conflict:{kind}")

    if not reasons:
        return sanitized, None, []
    return sanitized, quarantine_copy(original, headers, reasons, "sanitized_copy"), reasons


def reconcile_source_rows(
    rows: list[dict[str, Any]],
    headers: list[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    groups: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
    quarantined: list[dict[str, Any]] = []
    stats: defaultdict[str, int] = defaultdict(int)
    for row in rows:
        key = dedupe_key(row)
        if not key:
            quarantined.append(
                quarantine_copy(row, headers, ["ambiguous_source_identity"], "removed")
            )
            stats["ambiguous_source_rows"] += 1
            continue
        groups.setdefault(key, []).append(row)

    reconciled: list[dict[str, Any]] = []
    for key, group in groups.items():
        merged, conflict_fields = merge_compatible_rows(group, headers)
        if conflict_fields:
            reason = f"conflicting_duplicate:{','.join(conflict_fields)}"
            quarantined.extend(
                quarantine_copy(row, headers, [reason], "removed") for row in group
            )
            stats["conflicting_source_groups"] += 1
            stats["conflicting_source_rows"] += len(group)
            continue
        if len(group) > 1:
            stats["merged_source_groups"] += 1
            stats["merged_source_rows_removed"] += len(group) - 1
        sanitized, quarantine, reasons = sanitize_display_fields(merged, headers, SOURCE_SHEET)
        reconciled.append(sanitized)
        if quarantine:
            quarantined.append(quarantine)
            stats["sanitized_source_rows"] += 1
            stats["cleared_source_media_fields"] += sum(
                reason.startswith("placeholder_media:") for reason in reasons
            )
    return reconciled, quarantined, dict(stats)


def reconcile_cache_rows(
    rows: list[dict[str, Any]],
    headers: list[str],
    sheet_name: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    groups: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    quarantined: list[dict[str, Any]] = []
    stats: defaultdict[str, int] = defaultdict(int)
    for row in rows:
        row_id = clean_text(row.get("id"))
        if not row_id:
            quarantined.append(quarantine_copy(row, headers, ["missing_cache_id"], "removed"))
            stats["missing_id_rows"] += 1
            continue
        groups[row_id].append(row)

    reconciled: list[dict[str, Any]] = []
    for row_id in sorted(groups):
        group = groups[row_id]
        merged, conflict_fields = merge_compatible_rows(group, headers)
        if conflict_fields:
            reason = f"conflicting_duplicate:{','.join(conflict_fields)}"
            quarantined.extend(
                quarantine_copy(row, headers, [reason], "removed") for row in group
            )
            stats["conflicting_groups"] += 1
            stats["conflicting_rows"] += len(group)
            continue
        if len(group) > 1:
            stats["merged_groups"] += 1
            stats["merged_rows_removed"] += len(group) - 1
        sanitized, quarantine, reasons = sanitize_display_fields(merged, headers, sheet_name)
        reconciled.append(sanitized)
        if quarantine:
            quarantined.append(quarantine)
            stats["sanitized_rows"] += 1
            stats["cleared_media_fields"] += sum(
                reason.startswith("placeholder_media:") for reason in reasons
            )
            stats["cleared_invalid_approvals"] += sum(
                reason.startswith("approval_rights_conflict:") for reason in reasons
            )
    return reconciled, quarantined, dict(stats)


def quarantine_orphan_dishes(
    dishes: list[dict[str, Any]],
    headers: list[str],
    venue_ids: set[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    kept: list[dict[str, Any]] = []
    quarantined: list[dict[str, Any]] = []
    for dish in dishes:
        venue_id = clean_text(dish.get("venue_id"))
        if not venue_id or venue_id not in venue_ids:
            quarantined.append(
                quarantine_copy(dish, headers, ["orphan_dish_reference"], "removed")
            )
        else:
            kept.append(dish)
    return kept, quarantined


def quarantine_ambiguous_source_matches(
    rows: list[dict[str, Any]],
    headers: list[str],
    venues: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    indexes: defaultdict[str, set[str]] = defaultdict(set)
    for venue in venues:
        venue_id = clean_text(venue.get("id"))
        for key in source_match_keys(venue, cached=True):
            if venue_id:
                indexes[key].add(venue_id)

    kept: list[dict[str, Any]] = []
    quarantined: list[dict[str, Any]] = []
    for row in rows:
        candidates: set[str] = set()
        for key in source_match_keys(row, cached=False):
            candidates.update(indexes.get(key, set()))
        if len(candidates) > 1:
            quarantined.append(
                quarantine_copy(
                    row,
                    headers,
                    [f"ambiguous_cache_match:{','.join(sorted(candidates))}"],
                    "removed",
                )
            )
        else:
            kept.append(row)
    return kept, quarantined


def worksheet_rows(workbook: Any, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    headers, rows = read_optional_sheet_rows(workbook, sheet_name)
    return headers, rows


def replace_sheet(
    workbook: Any,
    sheet_name: str,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> None:
    if sheet_name in workbook.sheetnames:
        index = workbook.sheetnames.index(sheet_name)
        workbook.remove(workbook[sheet_name])
        sheet = workbook.create_sheet(sheet_name, index)
    else:
        sheet = workbook.create_sheet(sheet_name)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def combine_quarantine_rows(
    original_headers: list[str],
    existing_headers: list[str],
    existing_rows: list[dict[str, Any]],
    new_rows: list[dict[str, Any]],
) -> tuple[list[str], list[dict[str, Any]]]:
    headers = list(original_headers)
    for header in existing_headers:
        if header and header not in headers and header not in QUARANTINE_METADATA:
            headers.append(header)
    headers.extend(header for header in QUARANTINE_METADATA if header not in headers)
    return headers, [*existing_rows, *new_rows]


def report_rows(metrics: list[tuple[str, Any, str]]) -> list[dict[str, Any]]:
    return [{"metric": metric, "value": value, "notes": notes} for metric, value, notes in metrics]


def reconcile_workbook(workbook_path: Path, output_path: Path) -> dict[str, int]:
    source_path = workbook_path.expanduser().resolve()
    destination_path = output_path.expanduser().resolve()
    if source_path == destination_path:
        raise ReconciliationError(
            "--output must differ from --workbook; input overwrite is forbidden"
        )
    if not source_path.exists():
        raise ReconciliationError(f"Workbook does not exist: {source_path}")

    workbook = load_workbook(source_path, read_only=False, data_only=False)
    try:
        source_headers, source_rows = read_sheet_rows(workbook, SOURCE_SHEET)
        venue_headers, venue_rows = read_sheet_rows(workbook, VENUE_SHEET)
        dish_headers, dish_rows = read_sheet_rows(workbook, DISH_SHEET)
        existing_quarantine = {
            sheet: worksheet_rows(workbook, quarantine_sheet)
            for sheet, quarantine_sheet in QUARANTINE_SHEETS.items()
        }

        clean_source, quarantine_source, source_stats = reconcile_source_rows(
            source_rows, source_headers
        )
        clean_venues, quarantine_venues, venue_stats = reconcile_cache_rows(
            venue_rows, venue_headers, VENUE_SHEET
        )
        clean_dishes, quarantine_dishes, dish_stats = reconcile_cache_rows(
            dish_rows, dish_headers, DISH_SHEET
        )

        clean_source, ambiguous_source = quarantine_ambiguous_source_matches(
            clean_source, source_headers, clean_venues
        )
        quarantine_source.extend(ambiguous_source)
        venue_ids = {clean_text(row.get("id")) for row in clean_venues if clean_text(row.get("id"))}
        clean_dishes, orphan_dishes = quarantine_orphan_dishes(
            clean_dishes, dish_headers, venue_ids
        )
        quarantine_dishes.extend(orphan_dishes)

        validation_source = [
            {**row, "_row_number": index}
            for index, row in enumerate(clean_source, start=2)
        ]
        validation_venues = [
            {**row, "_row_number": index}
            for index, row in enumerate(clean_venues, start=2)
        ]
        validation_dishes = [
            {**row, "_row_number": index}
            for index, row in enumerate(clean_dishes, start=2)
        ]
        try:
            validate_export_inputs(validation_source, validation_venues, validation_dishes)
        except CatalogValidationError as exc:
            raise ReconciliationError(f"Reconciled candidate failed export preflight: {exc}") from exc

        replace_sheet(workbook, SOURCE_SHEET, source_headers, clean_source)
        replace_sheet(workbook, VENUE_SHEET, venue_headers, clean_venues)
        replace_sheet(workbook, DISH_SHEET, dish_headers, clean_dishes)

        quarantine_payload = {
            SOURCE_SHEET: (source_headers, quarantine_source),
            VENUE_SHEET: (venue_headers, quarantine_venues),
            DISH_SHEET: (dish_headers, quarantine_dishes),
        }
        for source_sheet, quarantine_sheet in QUARANTINE_SHEETS.items():
            original_headers, new_rows = quarantine_payload[source_sheet]
            existing_headers, existing_rows = existing_quarantine[source_sheet]
            quarantine_headers, all_rows = combine_quarantine_rows(
                original_headers,
                existing_headers,
                existing_rows,
                new_rows,
            )
            replace_sheet(workbook, quarantine_sheet, quarantine_headers, all_rows)

        metrics = [
            ("policy", "deterministic_workbook_only", "No Supabase or external writes"),
            ("input_sheet1_rows", len(source_rows), "Before reconciliation"),
            ("output_sheet1_rows", len(clean_source), "Candidate rows"),
            ("quarantined_sheet1_rows", len(quarantine_source), "New review records"),
            (
                "merged_source_groups",
                source_stats.get("merged_source_groups", 0),
                "Compatible duplicates",
            ),
            (
                "conflicting_source_groups",
                source_stats.get("conflicting_source_groups", 0),
                "Entire groups quarantined",
            ),
            (
                "ambiguous_source_rows",
                source_stats.get("ambiguous_source_rows", 0),
                "Missing stable identity",
            ),
            (
                "ambiguous_cache_match_rows",
                len(ambiguous_source),
                "Sheet1 rows matching multiple venues",
            ),
            (
                "sanitized_source_rows",
                source_stats.get("sanitized_source_rows", 0),
                "Originals copied to quarantine",
            ),
            ("input_venue_rows", len(venue_rows), "Before reconciliation"),
            ("output_venue_rows", len(clean_venues), "Candidate rows"),
            ("quarantined_venue_rows", len(quarantine_venues), "New review records"),
            (
                "merged_venue_groups",
                venue_stats.get("merged_groups", 0),
                "Compatible duplicate ids",
            ),
            (
                "conflicting_venue_groups",
                venue_stats.get("conflicting_groups", 0),
                "Entire groups quarantined",
            ),
            (
                "missing_venue_id_rows",
                venue_stats.get("missing_id_rows", 0),
                "Removed from candidate",
            ),
            (
                "sanitized_venue_rows",
                venue_stats.get("sanitized_rows", 0),
                "Originals copied to quarantine",
            ),
            ("input_dish_rows", len(dish_rows), "Before reconciliation"),
            ("output_dish_rows", len(clean_dishes), "Candidate rows"),
            ("quarantined_dish_rows", len(quarantine_dishes), "New review records"),
            (
                "merged_dish_groups",
                dish_stats.get("merged_groups", 0),
                "Compatible duplicate ids",
            ),
            (
                "conflicting_dish_groups",
                dish_stats.get("conflicting_groups", 0),
                "Entire groups quarantined",
            ),
            (
                "missing_dish_id_rows",
                dish_stats.get("missing_id_rows", 0),
                "Removed from candidate",
            ),
            (
                "sanitized_dish_rows",
                dish_stats.get("sanitized_rows", 0),
                "Originals copied to quarantine",
            ),
            ("orphan_dish_rows", len(orphan_dishes), "Removed from candidate"),
            (
                "cleared_placeholder_media_fields",
                source_stats.get("cleared_source_media_fields", 0)
                + venue_stats.get("cleared_media_fields", 0)
                + dish_stats.get("cleared_media_fields", 0),
                "Original values retained in quarantine",
            ),
            (
                "cleared_invalid_display_approvals",
                venue_stats.get("cleared_invalid_approvals", 0)
                + dish_stats.get("cleared_invalid_approvals", 0),
                "Approval was true without allowed rights",
            ),
            ("export_preflight", "passed", "Hardened structural preflight"),
        ]
        replace_sheet(
            workbook,
            REPORT_SHEET,
            ["metric", "value", "notes"],
            report_rows(metrics),
        )

        destination_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = destination_path.with_name(
            f".{destination_path.name}.{os.getpid()}.tmp{destination_path.suffix}"
        )
        if temp_path.exists():
            temp_path.unlink()
        try:
            workbook.save(temp_path)
            probe = load_workbook(temp_path, read_only=True, data_only=True)
            probe.close()
            temp_path.replace(destination_path)
        except BaseException:
            if temp_path.exists():
                temp_path.unlink()
            raise
    finally:
        workbook.close()

    return {metric: value for metric, value, _notes in metrics if isinstance(value, int)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        metrics = reconcile_workbook(args.workbook, args.output)
    except ReconciliationError as exc:
        print(f"reconciliation_failed={exc}")
        return 2
    print(f"reconciled_workbook={args.output}")
    for metric, value in metrics.items():
        print(f"{metric}={value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
