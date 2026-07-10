#!/usr/bin/env python3
"""Prunes permanently closed and invalid Google Maps URI rows from the Excel database.

This is an O(N) rebuild operation for openpyxl safety on large datasets.
"""

from __future__ import annotations

import os
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "Hungry F&B Database.xlsx"
SCRAPE_RESULTS_PATH = ROOT / "data" / "scrape_results.json"
SHEET_TAB = "Sheet1"

def clean(value) -> str:
    return "" if value is None else str(value).strip()

def main():
    if not SCRAPE_RESULTS_PATH.exists():
        print(f"Scrape results file not found at {SCRAPE_RESULTS_PATH}. Please run scrape_maps.js first.")
        return

    print("Loading scrape results...")
    with open(SCRAPE_RESULTS_PATH, "r", encoding="utf-8") as f:
        scrape_results = json.load(f)

    # Compile set of row indices to remove
    indices_to_remove = set()
    closed_count = 0
    invalid_count = 0

    for item in scrape_results:
        row_idx = item.get("row_index")
        is_closed = item.get("isClosed") == "Yes"
        is_invalid = item.get("isValid") == "No"

        if is_closed or is_invalid:
            indices_to_remove.add(row_idx)
            if is_closed:
                closed_count += 1
            if is_invalid:
                invalid_count += 1

    print(f"Loaded {len(scrape_results)} scraped results.")
    print(f"Found {len(indices_to_remove)} rows to remove (Permanently Closed: {closed_count}, Invalid URI: {invalid_count}).")

    if not indices_to_remove:
        print("No rows need to be removed. Exiting.")
        return

    print("Backing up workbook...")
    backup_path = EXCEL_PATH.with_name(f"Hungry F&B Database.backup-prune.xlsx")
    try:
        import shutil
        shutil.copy(EXCEL_PATH, backup_path)
        print(f"Saved backup copy to {backup_path}")
    except Exception as e:
        print(f"Warning: could not create backup copy: {e}")

    print("Loading Excel workbook...")
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_TAB]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean(v) for v in header_row]

    # Create new workbook to write valid rows O(N)
    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = SHEET_TAB
    new_ws.append(list(header_row))

    retained_count = 0
    removed_count = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if idx in indices_to_remove:
            removed_count += 1
        else:
            new_ws.append(list(row))
            retained_count += 1

    print(f"Rebuilding Excel sheet: kept {retained_count} valid rows, pruned {removed_count} rows.")
    new_wb.save(EXCEL_PATH)
    print("Excel database successfully saved.")

if __name__ == "__main__":
    main()
