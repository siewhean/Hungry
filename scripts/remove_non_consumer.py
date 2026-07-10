#!/usr/bin/env python3
from __future__ import annotations

import os
import re
import json
import urllib.parse
import urllib.request
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "Hungry F&B Database.xlsx"
SHEET_ID = "18R6h57sXZPuYyFkZhnvUr81mEwf4ZCynbzGPaLXr9ZM"
SHEET_TAB = "Sheet1"

NON_CONSUMER_KEYWORDS = [
    r"\bpte\.?\s*ltd\.?\b",
    r"\bprivate\s+limited\b",
    r"\bschool\b",
    r"\bcanteen\b",
    r"\bcatering\b",
    r"\bdistributor\b",
]
PATTERN = re.compile("|".join(NON_CONSUMER_KEYWORDS), re.IGNORECASE)

def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()

def fetch_access_token_from_refresh() -> str:
    client_id = clean(os.getenv("GOOGLE_OAUTH_CLIENT_ID"))
    client_secret = clean(os.getenv("GOOGLE_OAUTH_CLIENT_SECRET"))
    refresh_token = clean(os.getenv("GOOGLE_OAUTH_REFRESH_TOKEN"))
    if not (client_id and client_secret and refresh_token):
        return ""

    token_url = "https://oauth2.googleapis.com/token"
    body = urllib.parse.urlencode({
        "client_id": client_id,
        "client_secret": client_secret,
        "refresh_token": refresh_token,
        "grant_type": "refresh_token",
    }).encode("utf-8")
    req = urllib.request.Request(url=token_url, data=body, method="POST")
    req.add_header("Content-Type", "application/x-www-form-urlencoded")
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            payload = json.loads(resp.read().decode("utf-8", errors="replace"))
            return clean(payload.get("access_token"))
    except Exception:
        return ""

def api_request(token: str, method: str, url: str, payload: dict | None = None) -> dict:
    data = None
    req = urllib.request.Request(url=url, method=method)
    req.add_header("Authorization", f"Bearer {token}")
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
        req.add_header("Content-Type", "application/json; charset=utf-8")
    try:
        with urllib.request.urlopen(req, data=data, timeout=60) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} {method} {url}: {detail[:600]}") from exc

def get_sheet_id(token: str, spreadsheet_id: str, sheet_title: str) -> int | None:
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}"
    resp = api_request(token, "GET", url)
    for sheet in resp.get("sheets", []):
        props = sheet.get("properties", {})
        if props.get("title") == sheet_title:
            return props.get("sheetId")
    return None

def main():
    print("Loading Excel workbook...")
    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook[SHEET_TAB]

    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean(value) for value in header_row]

    name_idx = headers.index("Stall / Restaurant Name") if "Stall / Restaurant Name" in headers else -1
    norm_idx = headers.index("Business Name Normalized") if "Business Name Normalized" in headers else -1

    if name_idx == -1 and norm_idx == -1:
        print("Could not find name columns in Excel.")
        return

    # Using O(N) copy strategy for Excel
    print("Filtering rows...")
    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = SHEET_TAB
    new_ws.append(list(header_row))

    rows_to_delete = []

    for idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        name = clean(row[name_idx]) if name_idx != -1 and name_idx < len(row) else ""
        norm = clean(row[norm_idx]) if norm_idx != -1 and norm_idx < len(row) else ""

        if PATTERN.search(name) or PATTERN.search(norm):
            rows_to_delete.append((idx, name, norm))
        else:
            new_ws.append(list(row))

    print(f"Found {len(rows_to_delete)} non-consumer rows to delete.")

    if rows_to_delete:
        print("Saving Excel file...")
        new_wb.save(EXCEL_PATH)
    else:
        print("No non-consumer rows found.")

    # Google Sheets deletion
    if not rows_to_delete:
        return

    token = clean(os.getenv("GOOGLE_SHEETS_ACCESS_TOKEN"))
    token = fetch_access_token_from_refresh() or token
    if not token:
        print("Missing GOOGLE_SHEETS_ACCESS_TOKEN, skipping Google Sheets update.")
        return

    sheet_id = get_sheet_id(token, SHEET_ID, SHEET_TAB)
    if sheet_id is None:
        print("Could not find sheet ID for tab", SHEET_TAB)
        return

    requests = []
    # Sheet indices are 0-based. Row 1 is header (index 0).
    # Excel row_idx corresponds to Sheet index (row_idx - 1).
    for row_idx, _, _ in sorted(rows_to_delete, key=lambda x: x[0], reverse=True):
        sheet_idx = row_idx - 1
        requests.append({
            "deleteDimension": {
                "range": {
                    "sheetId": sheet_id,
                    "dimension": "ROWS",
                    "startIndex": sheet_idx,
                    "endIndex": sheet_idx + 1
                }
            }
        })

    payload = {"requests": requests}
    update_url = f"https://sheets.googleapis.com/v4/spreadsheets/{SHEET_ID}:batchUpdate"
    print("Updating Google Sheets...")
    api_request(token, "POST", update_url, payload)
    print("Successfully deleted rows from Google Sheets.")

if __name__ == "__main__":
    main()
