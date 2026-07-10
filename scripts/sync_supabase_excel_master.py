#!/usr/bin/env python3
"""Sync the Hungry Excel workbook with canonical Supabase data.

The first sheet keeps the existing Excel import contract. Extra sheets make the
workbook a reviewable master without changing the importer surface.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib import error, parse, request

import xlsxwriter
from openpyxl import load_workbook

from validate_excel_workbook import DEFAULT_WORKBOOK, EXPECTED_COLUMNS, clean, dedupe_key


DEFAULT_OUTPUT = DEFAULT_WORKBOOK
DEFAULT_ARCHIVE_DIR = Path("data/archive/20260707-master-sync")
DEFAULT_STATE_FILE = Path("data/supabase_excel_sync_state.json")
MENU_ENRICHMENT = Path("data/menu_enrichment_results.json")
SUPABASE_URL_ENV = "EXPO_PUBLIC_SUPABASE_URL"
SUPABASE_SERVICE_KEY_ENV = "SUPABASE_SERVICE_ROLE_KEY"
REST_PAGE_SIZE = 1000

BAD_URL_PATTERNS = (
    "images.unsplash.com",
    "source.unsplash.com",
    "picsum.photos",
    "hungry-logo",
    "hungry-app-logo",
)


VENUE_COLUMNS = [
    "id",
    "building_id",
    "business_id",
    "stall_id",
    "licence_no",
    "name",
    "unit_number",
    "floor",
    "hygiene_grade",
    "operating_hours",
    "address",
    "postal_code",
    "planning_area",
    "latitude",
    "longitude",
    "google_place_id",
    "google_rating",
    "google_review_count",
    "source_type",
    "source_confidence",
    "last_verified_at",
    "updated_at",
    "media_approved_for_display",
    "menu_approved_for_display",
    "media_rights_status",
    "menu_rights_status",
    "image_url",
    "logo_url",
    "menu_source_url",
    "cuisine_tags",
    "dietary_tags",
]


DISH_COLUMNS = [
    "id",
    "venue_id",
    "building_id",
    "name",
    "description",
    "price",
    "price_range",
    "image_url",
    "media_approved_for_display",
    "menu_approved_for_display",
    "media_rights_status",
    "menu_rights_status",
    "menu_source_url",
    "source_type",
    "food_type_tag",
    "updated_at",
]


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


def require_supabase_env() -> tuple[str, str]:
    load_env_file(Path(".env.local"))
    url = clean(os.environ.get(SUPABASE_URL_ENV))
    key = clean(os.environ.get(SUPABASE_SERVICE_KEY_ENV))
    if not url or not key:
        raise SystemExit(f"Missing {SUPABASE_URL_ENV} or {SUPABASE_SERVICE_KEY_ENV}")
    return url.rstrip("/"), key


def rest_get(
    base_url: str,
    service_key: str,
    table: str,
    columns: list[str],
    limit: int,
    last_id: str | None = None,
) -> list[dict[str, Any]]:
    params = {
        "select": ",".join(columns),
        "order": "id.asc",
        "limit": str(limit),
    }
    if last_id:
        params["id"] = f"gt.{last_id}"
    query = parse.urlencode(params)
    endpoint = f"{base_url}/rest/v1/{table}?{query}"
    req = request.Request(endpoint, method="GET")
    req.add_header("apikey", service_key)
    req.add_header("Authorization", f"Bearer {service_key}")
    try:
        with request.urlopen(req, timeout=120) as response:
            return json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")
        raise SystemExit(f"Supabase REST GET {table} failed: {exc.code} {detail}") from exc
    except error.URLError as exc:
        raise SystemExit(f"Supabase REST GET {table} failed: {exc.reason}") from exc


def fetch_table(base_url: str, service_key: str, table: str, columns: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    last_id: str | None = None
    while True:
        batch = rest_get(base_url, service_key, table, columns, REST_PAGE_SIZE, last_id)
        rows.extend(batch)
        if len(batch) < REST_PAGE_SIZE:
            return rows
        last_id = clean(batch[-1].get("id"))
        if not last_id:
            raise SystemExit(f"Supabase REST GET {table} failed: missing id in paginated batch")


def normalize_sync_timestamp(value: Any, *, label: str) -> str:
    raw = clean(value)
    if not raw:
        raise SystemExit(f"Missing {label} timestamp")
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError as exc:
        raise SystemExit(f"Invalid {label} timestamp: {raw}") from exc
    if parsed.tzinfo is None:
        raise SystemExit(f"Invalid {label} timestamp (timezone required): {raw}")
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def rest_get_updated(
    base_url: str,
    service_key: str,
    table: str,
    columns: list[str],
    limit: int,
    since: str,
    until: str,
    last_updated_at: str | None = None,
    last_id: str | None = None,
) -> list[dict[str, Any]]:
    filters = [("select", ",".join(columns)), ("updated_at", f"lte.{until}")]
    if last_updated_at and last_id:
        filters.append(
            (
                "or",
                f"(updated_at.gt.{last_updated_at},and(updated_at.eq.{last_updated_at},id.gt.{last_id}))",
            )
        )
    else:
        filters.append(("updated_at", f"gt.{since}"))
    filters.extend([("order", "updated_at.asc,id.asc"), ("limit", str(limit))])
    query = parse.urlencode(filters)
    endpoint = f"{base_url}/rest/v1/{table}?{query}"
    req = request.Request(endpoint, method="GET")
    req.add_header("apikey", service_key)
    req.add_header("Authorization", f"Bearer {service_key}")
    try:
        with request.urlopen(req, timeout=120) as response:
            return json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")
        raise SystemExit(f"Supabase incremental REST GET {table} failed: {exc.code} {detail}") from exc
    except error.URLError as exc:
        raise SystemExit(f"Supabase incremental REST GET {table} failed: {exc.reason}") from exc


def fetch_updated_table(
    base_url: str,
    service_key: str,
    table: str,
    columns: list[str],
    since: str,
    until: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    last_updated_at: str | None = None
    last_id: str | None = None
    while True:
        batch = rest_get_updated(
            base_url,
            service_key,
            table,
            columns,
            REST_PAGE_SIZE,
            since,
            until,
            last_updated_at,
            last_id,
        )
        rows.extend(batch)
        if len(batch) < REST_PAGE_SIZE:
            return rows
        last_updated_at = clean(batch[-1].get("updated_at"))
        last_id = clean(batch[-1].get("id"))
        if not last_updated_at or not last_id:
            raise SystemExit(f"Supabase incremental REST GET {table} failed: missing updated_at/id cursor")


def norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (clean(value) or "").casefold()).strip()


def key_parts(*parts: Any) -> str:
    return "|".join(norm(part) for part in parts if clean(part))


def is_good_url(value: Any, *, allow_grab: bool = False) -> bool:
    url = clean(value)
    if not url or not re.match(r"^https?://", url):
        return False
    lower = url.casefold()
    if "," in url:
        return False
    if any(pattern in lower for pattern in BAD_URL_PATTERNS):
        return False
    if "huawei-food-cms.grab.com" in lower:
        return allow_grab
    return True


def choose_venue(current: dict[str, Any] | None, candidate: dict[str, Any]) -> dict[str, Any]:
    if current is None:
        return candidate

    def score(row: dict[str, Any]) -> tuple[int, int, int, str]:
        return (
            1 if is_good_url(row.get("logo_url")) else 0,
            1 if row.get("media_approved_for_display") else 0,
            1 if is_good_url(row.get("image_url")) else 0,
            clean(row.get("updated_at")),
        )

    return candidate if score(candidate) > score(current) else current


def build_venue_indexes(venues: list[dict[str, Any]]) -> dict[str, dict[str, dict[str, Any]]]:
    indexes: dict[str, dict[str, dict[str, Any]]] = {
        "business_id": {},
        "stall_id": {},
        "licence_no": {},
        "name_postal_unit": {},
        "name_address": {},
    }
    for venue in venues:
        for field in ("business_id", "stall_id", "licence_no"):
            value = norm(venue.get(field))
            if value:
                indexes[field][value] = choose_venue(indexes[field].get(value), venue)
        npu = key_parts(venue.get("name"), venue.get("postal_code"), venue.get("unit_number"))
        if npu:
            indexes["name_postal_unit"][npu] = choose_venue(indexes["name_postal_unit"].get(npu), venue)
        na = key_parts(venue.get("name"), venue.get("address"))
        if na:
            indexes["name_address"][na] = choose_venue(indexes["name_address"].get(na), venue)
    return indexes


def match_venue(row: dict[str, Any], indexes: dict[str, dict[str, dict[str, Any]]]) -> tuple[dict[str, Any] | None, str]:
    candidates = [
        ("business_id", norm(row.get("Business ID"))),
        ("stall_id", norm(row.get("Stall ID"))),
        ("licence_no", norm(row.get("Licence No"))),
        ("name_postal_unit", key_parts(row.get("Stall / Restaurant Name"), row.get("Postal Code"), row.get("Unit No."))),
        ("name_address", key_parts(row.get("Stall / Restaurant Name"), row.get("Address"))),
    ]
    for index_name, value in candidates:
        if value and value in indexes[index_name]:
            return indexes[index_name][value], index_name
    return None, "unmatched"


def format_menu_items(dishes: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    seen: set[str] = set()
    for dish in sorted(dishes, key=lambda item: (norm(item.get("name")), clean(item.get("price_range")) or "")):
        name = clean(dish.get("name"))
        if not name or norm(name) in seen:
            continue
        seen.add(norm(name))
        price = clean(dish.get("price_range")) or (f"${dish.get('price'):.2f}" if isinstance(dish.get("price"), (int, float)) else "")
        item = f"{name} ({price})" if price else name
        if len("; ".join(parts + [item])) > 32000:
            break
        parts.append(item)
    return "; ".join(parts)


def normalize_coverage_status(value: Any) -> str | None:
    status = clean(value)
    if not status:
        return None
    lowered = status.casefold().strip()
    if lowered in {"complete", "partial", "unknown", "out_of_scope"}:
        return status
    if lowered == "active":
        return "Partial"
    if any(marker in lowered for marker in ("hidden", "excluded", "out of scope", "non-consumer", "non consumer")):
        return "out_of_scope"
    return "Unknown"


def load_sheet_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            rows.append({**{headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}, "_row_number": row_number})
        return headers, rows
    finally:
        workbook.close()


def load_cached_table_rows(workbook_path: Path, sheet_name: str, expected_columns: list[str]) -> list[dict[str, Any]] | None:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return None
        sheet = workbook[sheet_name]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_values is None:
            return None
        headers = [clean(value) for value in header_values]
        if any(column not in headers for column in expected_columns):
            return None
        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in values):
                continue
            rows.append({header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header})
        return rows
    finally:
        workbook.close()


def merge_rows_by_id(existing: list[dict[str, Any]], changed: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged = {clean(row.get("id")): row for row in existing if clean(row.get("id"))}
    for row in changed:
        row_id = clean(row.get("id"))
        if not row_id:
            raise SystemExit("Incremental Supabase row is missing id")
        merged[row_id] = row
    return sorted(merged.values(), key=lambda row: clean(row.get("id")) or "")


def read_sync_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Cannot read sync state {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise SystemExit(f"Cannot read sync state {path}: expected a JSON object")
    return data


def load_sync_state(path: Path, *, workbook: Path | None = None) -> str | None:
    data = read_sync_state(path)
    if not data:
        return None
    state_workbook = clean(data.get("workbook"))
    if workbook is not None and state_workbook and Path(state_workbook).resolve() != workbook.resolve():
        return None
    return normalize_sync_timestamp(data.get("last_successful_sync_at"), label="state last_successful_sync_at")


def write_sync_state(
    path: Path,
    *,
    last_successful_sync_at: str,
    output: Path,
    last_full_reconciliation_at: str | None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    payload = {
        "version": 2,
        "last_successful_sync_at": last_successful_sync_at,
        "last_full_reconciliation_at": last_full_reconciliation_at,
        "deletion_policy": "Hard deletes require a periodic --reconcile-deletions full sync.",
        "workbook": str(output.resolve()),
    }
    try:
        tmp_path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        tmp_path.replace(path)
    except BaseException:
        if tmp_path.exists():
            tmp_path.unlink()
        raise


def update_sheet1_rows(
    rows: list[dict[str, Any]],
    venues: list[dict[str, Any]],
    dishes_by_venue: dict[str, list[dict[str, Any]]],
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    indexes = build_venue_indexes(venues)
    stats = defaultdict(int)
    updated: list[dict[str, Any]] = []
    for row in rows:
        venue, match_type = match_venue(row, indexes)
        updated.append(update_sheet1_row(row, venue, dishes_by_venue, stats, match_type))
    return updated, dict(stats)


def update_sheet1_row(
    row: dict[str, Any],
    venue: dict[str, Any] | None,
    dishes_by_venue: dict[str, list[dict[str, Any]]],
    stats: defaultdict[str, int],
    match_type: str,
) -> dict[str, Any]:
    out = {column: row.get(column) for column in EXPECTED_COLUMNS}
    out["Coverage Status"] = normalize_coverage_status(out.get("Coverage Status"))
    stats[f"match_{match_type}"] += 1

    existing_image = out.get("Image URL")
    if not is_good_url(existing_image):
        out["Image URL"] = None
        if existing_image:
            stats["removed_bad_image_url"] += 1

    if not is_good_url(out.get("Dishes Image URL")):
        out["Dishes Image URL"] = None

    if not venue:
        return out

    if is_good_url(venue.get("logo_url")):
        out["Logo URL"] = venue.get("logo_url")
        stats["filled_logo_url"] += 1
    else:
        out["Logo URL"] = None

    if venue.get("media_approved_for_display") and is_good_url(venue.get("image_url")):
        out["Image URL"] = venue.get("image_url")
        stats["filled_approved_venue_image_url"] += 1

    if clean(venue.get("planning_area")):
        out["Area"] = venue.get("planning_area")
    if clean(venue.get("google_rating")):
        out["Google Rating"] = venue.get("google_rating")
    if clean(venue.get("google_review_count")):
        out["Google Review Count"] = venue.get("google_review_count")
    if clean(venue.get("operating_hours")) and clean(out.get("Opening Hours")) in {"", "Hours not provided"}:
        out["Opening Hours"] = venue.get("operating_hours")
    if clean(venue.get("menu_source_url")) and not clean(out.get("Menu Source")):
        out["Menu Source"] = venue.get("menu_source_url")

    venue_dishes = dishes_by_venue.get(clean(venue.get("id")), [])
    approved_images = [
        dish.get("image_url")
        for dish in venue_dishes
        if dish.get("media_approved_for_display")
        and dish.get("menu_approved_for_display")
        and is_good_url(dish.get("image_url"))
    ]
    if approved_images:
        out["Dishes Image URL"] = approved_images[0]
        stats["filled_approved_dish_image_url"] += 1
    menu_items = format_menu_items([dish for dish in venue_dishes if dish.get("menu_approved_for_display")])
    if menu_items:
        out["Menu Items"] = menu_items
        stats["updated_menu_items_from_supabase"] += 1
    return out


def update_sheet1_rows_incremental(
    rows: list[dict[str, Any]],
    previous_venues: list[dict[str, Any]],
    venues: list[dict[str, Any]],
    dishes_by_venue: dict[str, list[dict[str, Any]]],
    affected_venue_ids: set[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    previous_indexes = build_venue_indexes(previous_venues)
    current_indexes = build_venue_indexes(venues)
    venues_by_id = {clean(venue.get("id")): venue for venue in venues if clean(venue.get("id"))}
    stats: defaultdict[str, int] = defaultdict(int)
    updated: list[dict[str, Any]] = []
    changed_rows: list[dict[str, Any]] = []
    for row in rows:
        current_venue, current_match = match_venue(row, current_indexes)
        previous_venue, previous_match = match_venue(row, previous_indexes)
        current_id = clean(current_venue.get("id")) if current_venue else None
        previous_id = clean(previous_venue.get("id")) if previous_venue else None
        if current_id not in affected_venue_ids and previous_id not in affected_venue_ids:
            updated.append({column: row.get(column) for column in EXPECTED_COLUMNS})
            stats["unchanged_sheet1_rows"] += 1
            continue

        venue = current_venue
        match_type = current_match
        if venue is None and previous_id:
            venue = venues_by_id.get(previous_id)
            match_type = f"previous_{previous_match}"
        changed = update_sheet1_row(row, venue, dishes_by_venue, stats, match_type)
        updated.append(changed)
        stats["affected_sheet1_rows"] += 1
        if any(changed.get(column) != row.get(column) for column in EXPECTED_COLUMNS):
            changed_rows.append({**changed, "_row_number": row.get("_row_number")})
            stats["updated_sheet1_rows"] += 1
    return updated, changed_rows, dict(stats)


def append_rows(sheet: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header)
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            values.append(value)
        sheet.append(values)


def xlsx_value(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if value is None:
        return None
    if isinstance(value, str) and len(value) > 32767:
        return value[:32760] + " [truncated]"
    return value


def write_xlsx_rows(workbook: xlsxwriter.Workbook, sheet_name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.add_worksheet(sheet_name)
    header_format = workbook.add_format({"bold": True, "bg_color": "#F2F4F7"})
    for column_index, header in enumerate(headers):
        worksheet.write(0, column_index, header, header_format)
    for row_index, row in enumerate(rows, start=1):
        for column_index, header in enumerate(headers):
            value = xlsx_value(row.get(header))
            if value is None:
                continue
            worksheet.write(row_index, column_index, value)
    worksheet.freeze_panes(1, 0)
    worksheet.autofilter(0, 0, min(len(rows), 1_048_575), max(len(headers) - 1, 0))


def load_grab_image_evidence() -> list[dict[str, Any]]:
    if not MENU_ENRICHMENT.exists():
        return []
    data = json.loads(MENU_ENRICHMENT.read_text(encoding="utf-8"))
    evidence: list[dict[str, Any]] = []
    for record in data:
        for dish in record.get("dishes") or []:
            image_url = clean(dish.get("imageUrl"))
            image_url = "" if image_url is None else str(image_url)
            host = parse.urlparse(image_url).netloc.casefold()
            if host not in {"huawei-food-cms.grab.com", "food.grab.com"} or not is_good_url(image_url, allow_grab=True):
                continue
            evidence.append(
                {
                    "stall_name": record.get("stall_name"),
                    "building": record.get("building"),
                    "address": record.get("address"),
                    "source": record.get("source"),
                    "dish_name": dish.get("name"),
                    "price": dish.get("price"),
                    "url_kind": "direct_image" if host == "huawei-food-cms.grab.com" else "grab_page_link",
                    "media_or_source_url": image_url,
                }
            )
    return evidence


def write_master_workbook(
    output_path: Path,
    sheet1_rows: list[dict[str, Any]],
    venues: list[dict[str, Any]],
    dishes: list[dict[str, Any]],
    grab_evidence: list[dict[str, Any]],
    report_rows: list[dict[str, Any]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_path.with_name(f".{output_path.name}.{os.getpid()}.tmp{output_path.suffix}")
    if tmp_path.exists():
        tmp_path.unlink()
    workbook = xlsxwriter.Workbook(
        str(tmp_path),
        {
            "strings_to_urls": False,
            "strings_to_formulas": False,
            "nan_inf_to_errors": True,
        },
    )
    workbook.set_properties(
        {
            "title": "Hungry F&B Database Master",
            "comments": "Generated by scripts/sync_supabase_excel_master.py",
        }
    )

    write_xlsx_rows(workbook, "Sheet1", EXPECTED_COLUMNS, sheet1_rows)
    write_xlsx_rows(workbook, "Supabase Venues", VENUE_COLUMNS, venues)
    write_xlsx_rows(workbook, "Supabase Dishes", DISH_COLUMNS, dishes)
    write_xlsx_rows(
        workbook,
        "Source Image Evidence",
        ["stall_name", "building", "address", "source", "dish_name", "price", "url_kind", "media_or_source_url"],
        grab_evidence,
    )
    write_xlsx_rows(workbook, "Cleanup Report", ["metric", "value", "notes"], report_rows)
    try:
        workbook.close()
        # Confirm the workbook is readable before replacing the existing master.
        load_workbook(tmp_path, read_only=True, data_only=True).close()
        tmp_path.replace(output_path)
    except BaseException:
        if tmp_path.exists():
            tmp_path.unlink()
        raise


def set_openpyxl_cell(cell: Any, value: Any) -> None:
    normalized = xlsx_value(value)
    cell.value = normalized
    # Match the full writer's strings_to_formulas=False safety behavior.
    if isinstance(normalized, str) and normalized.startswith("="):
        cell.data_type = "s"


def patch_cached_table_sheet(
    workbook: Any,
    sheet_name: str,
    columns: list[str],
    changed_rows: list[dict[str, Any]],
) -> int:
    if not changed_rows:
        return 0
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"Incremental workbook patch failed: missing {sheet_name} sheet")
    sheet = workbook[sheet_name]
    headers = {clean(cell.value): cell.column for cell in sheet[1] if clean(cell.value)}
    missing = [column for column in columns if column not in headers]
    if missing:
        raise SystemExit(f"Incremental workbook patch failed: {sheet_name} missing columns: {', '.join(missing)}")
    id_column = headers["id"]
    rows_by_id: dict[str, int] = {}
    for row_number in range(2, sheet.max_row + 1):
        row_id = clean(sheet.cell(row=row_number, column=id_column).value)
        if row_id:
            rows_by_id[row_id] = row_number

    patched_cells = 0
    for changed in changed_rows:
        row_id = clean(changed.get("id"))
        if not row_id:
            raise SystemExit(f"Incremental workbook patch failed: {sheet_name} row missing id")
        is_new = row_id not in rows_by_id
        row_number = rows_by_id.get(row_id, sheet.max_row + 1)
        rows_by_id[row_id] = row_number
        for column in columns:
            if not is_new and column not in changed:
                continue
            cell = sheet.cell(row=row_number, column=headers[column])
            value = changed.get(column)
            if cell.value != xlsx_value(value):
                set_openpyxl_cell(cell, value)
                patched_cells += 1
    return patched_cells


def patch_sheet1_rows(workbook: Any, changed_rows: list[dict[str, Any]]) -> int:
    if not changed_rows:
        return 0
    sheet = workbook["Sheet1"]
    headers = {clean(cell.value): cell.column for cell in sheet[1] if clean(cell.value)}
    missing = [column for column in EXPECTED_COLUMNS if column not in headers]
    if missing:
        raise SystemExit(f"Incremental workbook patch failed: Sheet1 missing columns: {', '.join(missing)}")
    patched_cells = 0
    for changed in changed_rows:
        row_number = changed.get("_row_number")
        if not isinstance(row_number, int) or row_number < 2 or row_number > sheet.max_row:
            raise SystemExit(f"Incremental workbook patch failed: invalid Sheet1 row number {row_number}")
        for column in EXPECTED_COLUMNS:
            cell = sheet.cell(row=row_number, column=headers[column])
            value = changed.get(column)
            if cell.value != xlsx_value(value):
                set_openpyxl_cell(cell, value)
                patched_cells += 1
    return patched_cells


def append_incremental_report(workbook: Any, report_rows: list[dict[str, Any]]) -> None:
    if "Cleanup Report" not in workbook.sheetnames:
        return
    sheet = workbook["Cleanup Report"]
    headers = {clean(cell.value): cell.column for cell in sheet[1] if clean(cell.value)}
    if any(column not in headers for column in ("metric", "value", "notes")):
        return
    for report in report_rows:
        row_number = sheet.max_row + 1
        for column in ("metric", "value", "notes"):
            set_openpyxl_cell(sheet.cell(row=row_number, column=headers[column]), report.get(column))


def write_incremental_workbook(
    source_path: Path,
    output_path: Path,
    sheet1_changed_rows: list[dict[str, Any]],
    changed_venues: list[dict[str, Any]],
    changed_dishes: list[dict[str, Any]],
    report_rows: list[dict[str, Any]],
) -> dict[str, int]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_path.with_name(f".{output_path.name}.{os.getpid()}.tmp{output_path.suffix}")
    if tmp_path.exists():
        tmp_path.unlink()
    try:
        shutil.copy2(source_path, tmp_path)
        workbook = load_workbook(tmp_path, read_only=False, data_only=False)
        try:
            stats = {
                "patched_sheet1_cells": patch_sheet1_rows(workbook, sheet1_changed_rows),
                "patched_venue_cells": patch_cached_table_sheet(workbook, "Supabase Venues", VENUE_COLUMNS, changed_venues),
                "patched_dish_cells": patch_cached_table_sheet(workbook, "Supabase Dishes", DISH_COLUMNS, changed_dishes),
            }
            append_incremental_report(workbook, report_rows)
            workbook.save(tmp_path)
        finally:
            workbook.close()
        load_workbook(tmp_path, read_only=True, data_only=True).close()
        tmp_path.replace(output_path)
        return stats
    except BaseException:
        if tmp_path.exists():
            tmp_path.unlink()
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--archive-dir", type=Path, default=DEFAULT_ARCHIVE_DIR)
    parser.add_argument("--no-backup", action="store_true")
    parser.add_argument(
        "--incremental",
        action="store_true",
        help="Fetch rows newer than the last successful timestamp in --state-file; first run safely falls back to a full sync.",
    )
    parser.add_argument(
        "--since",
        help="Fetch only rows with updated_at later than this timezone-aware ISO timestamp (implies --incremental).",
    )
    parser.add_argument(
        "--reconcile-deletions",
        action="store_true",
        help="Force a full sync so hard-deleted Supabase rows are removed from workbook cache sheets.",
    )
    parser.add_argument("--state-file", type=Path, default=DEFAULT_STATE_FILE)
    args = parser.parse_args()
    if args.reconcile_deletions and (args.incremental or args.since):
        parser.error("--reconcile-deletions cannot be combined with --incremental or --since")

    base_url, service_key = require_supabase_env()
    started_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    headers, workbook_rows = load_sheet_rows(args.workbook)
    missing = [column for column in EXPECTED_COLUMNS if column not in headers]
    if missing:
        raise SystemExit(f"Workbook is missing required columns: {', '.join(missing)}")

    incremental_requested = args.incremental or bool(args.since)
    state_requested = incremental_requested or args.reconcile_deletions
    previous_state = read_sync_state(args.state_file) if incremental_requested else {}
    since = normalize_sync_timestamp(args.since, label="--since") if args.since else None
    if args.incremental and since is None:
        since = load_sync_state(args.state_file, workbook=args.workbook)

    previous_venues: list[dict[str, Any]] = []
    previous_dishes: list[dict[str, Any]] = []
    incremental_active = False
    fallback_reason: str | None = None
    if args.reconcile_deletions:
        fallback_reason = "explicit deletion reconciliation requested"
    elif incremental_requested and since:
        previous_venues_cache = load_cached_table_rows(args.workbook, "Supabase Venues", VENUE_COLUMNS)
        previous_dishes_cache = load_cached_table_rows(args.workbook, "Supabase Dishes", DISH_COLUMNS)
        if previous_venues_cache is None or previous_dishes_cache is None:
            fallback_reason = "workbook is missing complete Supabase Venues/Supabase Dishes cache sheets"
        else:
            since_dt = datetime.fromisoformat(since.replace("Z", "+00:00"))
            until_dt = datetime.fromisoformat(started_at.replace("Z", "+00:00"))
            if since_dt > until_dt:
                raise SystemExit(f"Incremental --since timestamp is after this sync's upper watermark: {since}")
            previous_venues = previous_venues_cache
            previous_dishes = previous_dishes_cache
            incremental_active = True
    elif incremental_requested:
        fallback_reason = f"sync state is missing or does not match {args.workbook}"

    if incremental_active:
        changed_venues = fetch_updated_table(base_url, service_key, "venues", VENUE_COLUMNS, since, started_at)
        changed_dishes = fetch_updated_table(base_url, service_key, "dishes", DISH_COLUMNS, since, started_at)
        old_dishes_by_id = {
            clean(dish.get("id")): dish for dish in previous_dishes if clean(dish.get("id"))
        }
        venues = merge_rows_by_id(previous_venues, changed_venues)
        dishes = merge_rows_by_id(previous_dishes, changed_dishes)
        affected_venue_ids = {
            clean(venue.get("id")) for venue in changed_venues if clean(venue.get("id"))
        }
        for dish in changed_dishes:
            venue_id = clean(dish.get("venue_id"))
            if venue_id:
                affected_venue_ids.add(venue_id)
            old_dish = old_dishes_by_id.get(clean(dish.get("id")))
            old_venue_id = clean(old_dish.get("venue_id")) if old_dish else None
            if old_venue_id:
                affected_venue_ids.add(old_venue_id)
    else:
        changed_venues = []
        changed_dishes = []
        affected_venue_ids = set()
        venues = fetch_table(base_url, service_key, "venues", VENUE_COLUMNS)
        dishes = fetch_table(base_url, service_key, "dishes", DISH_COLUMNS)

    dishes_by_venue: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for dish in dishes:
        venue_id = clean(dish.get("venue_id"))
        if venue_id:
            dishes_by_venue[venue_id].append(dish)

    if incremental_active:
        synced_rows, sheet1_changed_rows, sync_stats = update_sheet1_rows_incremental(
            workbook_rows,
            previous_venues,
            venues,
            dishes_by_venue,
            affected_venue_ids,
        )
    else:
        sheet1_changed_rows = []
        synced_rows, sync_stats = update_sheet1_rows(workbook_rows, venues, dishes_by_venue)
    grab_evidence = load_grab_image_evidence()

    duplicate_keys = defaultdict(int)
    for row in synced_rows:
        key = dedupe_key(row)
        if key:
            duplicate_keys[key] += 1

    report_rows = [
        {"metric": "generated_at", "value": started_at, "notes": "UTC"},
        {"metric": "sync_mode", "value": "incremental" if incremental_active else "full_reconciliation" if args.reconcile_deletions else "full", "notes": fallback_reason or ""},
        {"metric": "deletion_policy", "value": "periodic_full_reconciliation", "notes": "Incremental updated_at filters do not observe hard deletes; run --reconcile-deletions."},
        {"metric": "incremental_since", "value": since or "", "notes": "Exclusive lower updated_at watermark" if incremental_active else ""},
        {"metric": "incremental_until", "value": started_at if incremental_requested else "", "notes": "Inclusive upper updated_at watermark" if incremental_active else ""},
        {"metric": "changed_venues_rows", "value": len(changed_venues), "notes": "Fetched during this run"},
        {"metric": "changed_dishes_rows", "value": len(changed_dishes), "notes": "Fetched during this run"},
        {"metric": "sheet1_rows", "value": len(synced_rows), "notes": "Existing Excel import-contract rows"},
        {"metric": "supabase_venues_rows", "value": len(venues), "notes": "Canonical venues exported to Supabase Venues sheet"},
        {"metric": "supabase_dishes_rows", "value": len(dishes), "notes": "Canonical dishes exported to Supabase Dishes sheet"},
        {"metric": "source_grab_media_evidence_rows", "value": len(grab_evidence), "notes": "Source evidence only; not display-approved catalog media"},
        {"metric": "source_grab_direct_image_rows", "value": sum(1 for row in grab_evidence if row["url_kind"] == "direct_image"), "notes": "Direct Grab CDN image URLs"},
        {"metric": "source_grab_page_link_rows", "value": sum(1 for row in grab_evidence if row["url_kind"] == "grab_page_link"), "notes": "Grab page links stored in source image fields"},
        {"metric": "source_grab_unique_urls", "value": len({row['media_or_source_url'] for row in grab_evidence}), "notes": "Current repo evidence count"},
        {"metric": "duplicate_source_keys", "value": sum(1 for count in duplicate_keys.values() if count > 1), "notes": "Warnings only; importer contract still accepts them"},
    ]
    for key, value in sorted(sync_stats.items()):
        report_rows.append({"metric": key, "value": value, "notes": ""})

    if args.output.resolve() == args.workbook.resolve() and not args.no_backup:
        args.archive_dir.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%dT%H%M%SZ", time.gmtime())
        backup_path = args.archive_dir / f"{args.workbook.stem}.pre-master-sync-{stamp}{args.workbook.suffix}"
        shutil.copy2(args.workbook, backup_path)
        report_rows.append({"metric": "workbook_backup", "value": str(backup_path), "notes": "Pre-sync copy"})

    if incremental_active:
        write_stats = write_incremental_workbook(
            args.workbook,
            args.output,
            sheet1_changed_rows,
            changed_venues,
            changed_dishes,
            report_rows,
        )
        sync_stats.update(write_stats)
    else:
        write_master_workbook(args.output, synced_rows, venues, dishes, grab_evidence, report_rows)
    if state_requested:
        # Advance only after the selected writer has validated and atomically
        # replaced the destination workbook. Incremental runs retain the most
        # recent explicit/full-fallback deletion-reconciliation watermark.
        previous_full_reconciliation = clean(previous_state.get("last_full_reconciliation_at"))
        if previous_full_reconciliation:
            previous_full_reconciliation = normalize_sync_timestamp(
                previous_full_reconciliation,
                label="state last_full_reconciliation_at",
            )
        last_full_reconciliation_at = started_at if not incremental_active else previous_full_reconciliation
        write_sync_state(
            args.state_file,
            last_successful_sync_at=started_at,
            output=args.output,
            last_full_reconciliation_at=last_full_reconciliation_at,
        )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "mode": "incremental" if incremental_active else "full_reconciliation" if args.reconcile_deletions else "full",
                "fallback_reason": fallback_reason,
                "since": since,
                "until": started_at if incremental_requested else None,
                "rows": len(synced_rows),
                "venues": len(venues),
                "dishes": len(dishes),
                "changed_venues": len(changed_venues),
                "changed_dishes": len(changed_dishes),
                "deletions_reconciled": not incremental_active,
                "grab_image_rows": len(grab_evidence),
                "stats": sync_stats,
            },
            indent=2,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
