from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock
from urllib import parse

from openpyxl import Workbook, load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import sync_supabase_excel_master as sync


def sheet1_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {column: None for column in sync.EXPECTED_COLUMNS}
    row.update(overrides)
    return row


class IncrementalSyncTests(unittest.TestCase):
    def test_cached_table_patch_reads_max_row_once_per_phase(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Supabase Venues"
        sheet.append(sync.VENUE_COLUMNS)
        for venue_id in ("venue-1", "venue-2"):
            sheet.append(
                [
                    venue_id if column == "id" else "Old" if column == "name" else None
                    for column in sync.VENUE_COLUMNS
                ]
            )

        worksheet_type = type(sheet)
        original_max_row = worksheet_type.max_row.fget
        max_row_reads = 0

        def counted_max_row(worksheet: object) -> int:
            nonlocal max_row_reads
            max_row_reads += 1
            return original_max_row(worksheet)

        with mock.patch.object(worksheet_type, "max_row", new=property(counted_max_row)):
            patched = sync.patch_cached_table_sheet(
                workbook,
                "Supabase Venues",
                sync.VENUE_COLUMNS,
                [{"id": "venue-1", "name": "New 1"}, {"id": "venue-2", "name": "New 2"}],
            )

        self.assertEqual(patched, 2)
        self.assertEqual(max_row_reads, 2)

    def test_first_incremental_page_uses_exclusive_since_watermark(self) -> None:
        response = mock.MagicMock()
        response.read.return_value = b"[]"
        response.__enter__.return_value = response

        with mock.patch.object(sync.request, "urlopen", return_value=response) as urlopen:
            sync.rest_get_updated(
                "https://example.supabase.co",
                "secret",
                "venues",
                ["id", "updated_at"],
                1000,
                "2026-07-01T00:00:00Z",
                "2026-07-02T00:00:00Z",
            )

        request_object = urlopen.call_args.args[0]
        query = parse.parse_qs(parse.urlparse(request_object.full_url).query)
        self.assertEqual(
            query["updated_at"],
            ["lte.2026-07-02T00:00:00Z", "gt.2026-07-01T00:00:00Z"],
        )

    def test_incremental_request_has_bounded_keyset_filters(self) -> None:
        response = mock.MagicMock()
        response.read.return_value = b"[]"
        response.__enter__.return_value = response

        with mock.patch.object(sync.request, "urlopen", return_value=response) as urlopen:
            rows = sync.rest_get_updated(
                "https://example.supabase.co",
                "secret",
                "venues",
                ["id", "updated_at"],
                1000,
                "2026-07-01T00:00:00Z",
                "2026-07-02T00:00:00Z",
                "2026-07-01T12:00:00Z",
                "venue-2000",
            )

        self.assertEqual(rows, [])
        request_object = urlopen.call_args.args[0]
        query = parse.parse_qs(parse.urlparse(request_object.full_url).query)
        self.assertEqual(query["updated_at"], ["lte.2026-07-02T00:00:00Z"])
        self.assertEqual(
            query["or"],
            [
                "(updated_at.gt.2026-07-01T12:00:00Z,"
                "and(updated_at.eq.2026-07-01T12:00:00Z,id.gt.venue-2000))"
            ],
        )
        self.assertEqual(query["order"], ["updated_at.asc,id.asc"])

    def test_merge_replaces_changed_ids_without_losing_unchanged_rows(self) -> None:
        merged = sync.merge_rows_by_id(
            [{"id": "venue-1", "name": "Old"}, {"id": "venue-2", "name": "Keep"}],
            [{"id": "venue-1", "name": "New"}, {"id": "venue-3", "name": "Added"}],
        )

        self.assertEqual(
            merged,
            [
                {"id": "venue-1", "name": "New"},
                {"id": "venue-2", "name": "Keep"},
                {"id": "venue-3", "name": "Added"},
            ],
        )

    def test_incremental_sheet1_update_preserves_unaffected_rows_exactly(self) -> None:
        rows = [
            {
                **sheet1_row(**{"Business ID": "business-1", "Stall / Restaurant Name": "Changed"}),
                "_row_number": 2,
            },
            sheet1_row(
                **{
                    "Business ID": "business-2",
                    "Stall / Restaurant Name": "Unchanged",
                    "Coverage Status": "Active",
                    "Image URL": "https://images.unsplash.com/keep-existing-value",
                }
            ),
        ]
        previous_venues = [
            {"id": "venue-1", "business_id": "business-1", "name": "Changed"},
            {"id": "venue-2", "business_id": "business-2", "name": "Unchanged"},
        ]
        venues = [
            {
                "id": "venue-1",
                "business_id": "business-1",
                "name": "Changed",
                "logo_url": "https://cdn.example/new-logo.png",
                "media_approved_for_display": True,
                "media_rights_status": "operator_owned",
                "planning_area": "Orchard",
            },
            previous_venues[1],
        ]

        updated, changed_rows, stats = sync.update_sheet1_rows_incremental(
            rows,
            previous_venues,
            venues,
            {"venue-1": []},
            {"venue-1"},
        )

        self.assertEqual(updated[0]["Logo URL"], "https://cdn.example/new-logo.png")
        self.assertEqual(updated[0]["Area"], "Orchard")
        self.assertEqual(updated[1]["Coverage Status"], "Active")
        self.assertEqual(updated[1]["Image URL"], "https://images.unsplash.com/keep-existing-value")
        self.assertEqual(len(changed_rows), 1)
        self.assertEqual(changed_rows[0]["_row_number"], 2)
        self.assertEqual(stats["updated_sheet1_rows"], 1)
        self.assertEqual(stats["unchanged_sheet1_rows"], 1)

    def test_main_merges_cached_sheets_and_advances_state_after_atomic_write(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            output = root / "output.xlsx"
            state = root / "state.json"
            previous_venues = [
                {"id": "venue-1", "business_id": "business-1", "name": "Changed", "updated_at": "2026-07-01T00:00:00Z"},
                {"id": "venue-2", "business_id": "business-2", "name": "Keep", "updated_at": "2026-07-01T00:00:00Z"},
            ]
            previous_dishes = [
                {"id": "dish-1", "venue_id": "venue-1", "name": "Old Dish", "updated_at": "2026-07-01T00:00:00Z"},
                {"id": "dish-2", "venue_id": "venue-2", "name": "Keep Dish", "updated_at": "2026-07-01T00:00:00Z"},
            ]
            sync.write_master_workbook(
                source,
                [
                    sheet1_row(**{"Business ID": "business-1", "Stall / Restaurant Name": "Changed"}),
                    sheet1_row(**{"Business ID": "business-2", "Stall / Restaurant Name": "Keep"}),
                ],
                previous_venues,
                previous_dishes,
                [],
                [],
            )
            workbook = load_workbook(source)
            unrelated = workbook.create_sheet("Unrelated Sheet")
            unrelated["A1"] = "preserve me"
            unrelated["B2"] = "=1+1"
            workbook.save(source)
            workbook.close()
            state.write_text(
                json.dumps({"last_successful_sync_at": "2026-07-01T00:00:00Z"}),
                encoding="utf-8",
            )
            changed_venue = {
                "id": "venue-1",
                "business_id": "business-1",
                "name": "Changed",
                "logo_url": "https://cdn.example/new-logo.png",
                "media_approved_for_display": True,
                "media_rights_status": "operator_owned",
                "updated_at": "2026-07-02T00:00:00Z",
            }
            changed_dish = {
                "id": "dish-1",
                "venue_id": "venue-1",
                "name": "New Dish",
                "menu_approved_for_display": True,
                "menu_rights_status": "operator_owned",
                "updated_at": "2026-07-02T00:00:00Z",
            }

            def fetch_changed(
                _base_url: str,
                _service_key: str,
                table: str,
                _columns: list[str],
                _since: str,
                _until: str,
            ) -> list[dict[str, object]]:
                return [changed_venue] if table == "venues" else [changed_dish]

            argv = [
                "sync_supabase_excel_master.py",
                "--workbook",
                str(source),
                "--output",
                str(output),
                "--incremental",
                "--state-file",
                str(state),
                "--no-backup",
            ]
            with (
                mock.patch.object(sys, "argv", argv),
                mock.patch.object(sync, "require_supabase_env", return_value=("https://example.supabase.co", "secret")),
                mock.patch.object(sync, "fetch_updated_table", side_effect=fetch_changed),
                mock.patch.object(sync, "load_grab_image_evidence", return_value=[]),
                mock.patch("builtins.print"),
            ):
                self.assertEqual(sync.main(), 0)

            venues = sync.load_cached_table_rows(output, "Supabase Venues", sync.VENUE_COLUMNS)
            dishes = sync.load_cached_table_rows(output, "Supabase Dishes", sync.DISH_COLUMNS)
            self.assertIsNotNone(venues)
            self.assertIsNotNone(dishes)
            self.assertEqual({row["id"] for row in venues or []}, {"venue-1", "venue-2"})
            self.assertEqual({row["id"] for row in dishes or []}, {"dish-1", "dish-2"})
            self.assertEqual(next(row for row in venues or [] if row["id"] == "venue-1")["logo_url"], "https://cdn.example/new-logo.png")
            self.assertEqual(next(row for row in venues or [] if row["id"] == "venue-2")["name"], "Keep")

            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                sheet = workbook["Sheet1"]
                headers = [cell.value for cell in sheet[1]]
                first_data_row = [cell.value for cell in sheet[2]]
                values = dict(zip(headers, first_data_row))
                self.assertEqual(values["Logo URL"], "https://cdn.example/new-logo.png")
                self.assertEqual(values["Menu Items"], "New Dish")
            finally:
                workbook.close()

            saved_state = json.loads(state.read_text(encoding="utf-8"))
            self.assertGreater(saved_state["last_successful_sync_at"], "2026-07-01T00:00:00Z")
            self.assertEqual(saved_state["workbook"], str(output.resolve()))
            self.assertIn("--reconcile-deletions", saved_state["deletion_policy"])

            preserved = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(preserved["Unrelated Sheet"]["A1"].value, "preserve me")
                self.assertEqual(preserved["Unrelated Sheet"]["B2"].value, "=1+1")
            finally:
                preserved.close()

    def test_sync_requires_approval_and_rights_and_preserves_verification_date(self) -> None:
        row = sheet1_row(
            **{
                "Business ID": "business-1",
                "Stall / Restaurant Name": "Venue",
                "Last Verified Date": "2026-06-01",
                "Image URL": "https://merchant.example/old.jpg",
                "Menu Items": "Existing Dish",
            }
        )
        venue = {
            "id": "venue-1",
            "business_id": "business-1",
            "name": "Venue",
            "last_verified_at": "2026-07-10T00:00:00Z",
            "updated_at": "2026-07-10T00:00:00Z",
            "media_approved_for_display": True,
            "media_rights_status": "unverified",
            "image_url": "https://merchant.example/new.jpg",
            "logo_url": "https://merchant.example/logo.jpg",
        }
        dishes = {
            "venue-1": [
                {
                    "name": "Unverified Dish",
                    "menu_approved_for_display": True,
                    "menu_rights_status": "unverified",
                    "media_approved_for_display": True,
                    "media_rights_status": "unverified",
                    "image_url": "https://merchant.example/dish.jpg",
                }
            ]
        }

        updated, _stats = sync.update_sheet1_rows([row], [venue], dishes)

        self.assertEqual(updated[0]["Last Verified Date"], "2026-06-01")
        self.assertEqual(updated[0]["Image URL"], "https://merchant.example/old.jpg")
        self.assertEqual(updated[0]["Menu Items"], "Existing Dish")
        self.assertIsNone(updated[0]["Logo URL"])

    def test_state_does_not_advance_when_workbook_write_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            state = root / "state.json"
            sync.write_master_workbook(
                source,
                [sheet1_row(**{"Business ID": "business-1", "Stall / Restaurant Name": "Venue"})],
                [{"id": "venue-1", "business_id": "business-1", "name": "Venue"}],
                [],
                [],
                [],
            )
            original_state = {"last_successful_sync_at": "2026-07-01T00:00:00Z"}
            state.write_text(json.dumps(original_state), encoding="utf-8")
            argv = [
                "sync_supabase_excel_master.py",
                "--workbook",
                str(source),
                "--output",
                str(source),
                "--incremental",
                "--state-file",
                str(state),
                "--no-backup",
            ]
            with (
                mock.patch.object(sys, "argv", argv),
                mock.patch.object(sync, "require_supabase_env", return_value=("https://example.supabase.co", "secret")),
                mock.patch.object(sync, "fetch_updated_table", return_value=[]),
                mock.patch.object(sync, "load_grab_image_evidence", return_value=[]),
                mock.patch.object(sync, "write_incremental_workbook", side_effect=RuntimeError("write failed")),
            ):
                with self.assertRaisesRegex(RuntimeError, "write failed"):
                    sync.main()

            self.assertEqual(json.loads(state.read_text(encoding="utf-8")), original_state)

    def test_reconcile_deletions_forces_full_cache_replacement_and_records_watermark(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.xlsx"
            output = root / "output.xlsx"
            state = root / "state.json"
            sync.write_master_workbook(
                source,
                [sheet1_row(**{"Business ID": "business-1", "Stall / Restaurant Name": "Keep"})],
                [
                    {"id": "venue-1", "business_id": "business-1", "name": "Keep"},
                    {"id": "venue-deleted", "business_id": "deleted", "name": "Deleted"},
                ],
                [
                    {"id": "dish-1", "venue_id": "venue-1", "name": "Keep Dish"},
                    {"id": "dish-deleted", "venue_id": "venue-deleted", "name": "Deleted Dish"},
                ],
                [],
                [],
            )
            state.write_text(
                json.dumps(
                    {
                        "last_successful_sync_at": "2026-07-01T00:00:00Z",
                        "last_full_reconciliation_at": "2026-06-01T00:00:00Z",
                        "workbook": str(source.resolve()),
                    }
                ),
                encoding="utf-8",
            )

            def fetch_full(
                _base_url: str,
                _service_key: str,
                table: str,
                _columns: list[str],
            ) -> list[dict[str, object]]:
                if table == "venues":
                    return [{"id": "venue-1", "business_id": "business-1", "name": "Keep"}]
                return [{"id": "dish-1", "venue_id": "venue-1", "name": "Keep Dish"}]

            argv = [
                "sync_supabase_excel_master.py",
                "--workbook",
                str(source),
                "--output",
                str(output),
                "--reconcile-deletions",
                "--state-file",
                str(state),
                "--no-backup",
            ]
            with (
                mock.patch.object(sys, "argv", argv),
                mock.patch.object(sync, "require_supabase_env", return_value=("https://example.supabase.co", "secret")),
                mock.patch.object(sync, "fetch_table", side_effect=fetch_full),
                mock.patch.object(sync, "load_grab_image_evidence", return_value=[]),
                mock.patch("builtins.print"),
            ):
                self.assertEqual(sync.main(), 0)

            venues = sync.load_cached_table_rows(output, "Supabase Venues", sync.VENUE_COLUMNS)
            dishes = sync.load_cached_table_rows(output, "Supabase Dishes", sync.DISH_COLUMNS)
            self.assertEqual([row["id"] for row in venues or []], ["venue-1"])
            self.assertEqual([row["id"] for row in dishes or []], ["dish-1"])
            saved_state = json.loads(state.read_text(encoding="utf-8"))
            self.assertGreater(saved_state["last_full_reconciliation_at"], "2026-06-01T00:00:00Z")
            self.assertEqual(
                saved_state["last_full_reconciliation_at"],
                saved_state["last_successful_sync_at"],
            )

    def test_incremental_keyset_queries_have_supporting_indexes(self) -> None:
        migration = (
            Path(__file__).resolve().parents[1]
            / "supabase/migrations/20260710130000_enrichment_incremental_sync_indexes.sql"
        ).read_text(encoding="utf-8")
        normalized = " ".join(migration.lower().split())
        self.assertIn("alter table public.venues add column if not exists updated_at", normalized)
        self.assertIn("alter table public.dishes add column if not exists updated_at", normalized)
        self.assertIn("create trigger trg_venues_updated_at", normalized)
        self.assertIn("create trigger trg_dishes_updated_at", normalized)
        self.assertIn("on public.venues (updated_at, id)", normalized)
        self.assertIn("on public.dishes (updated_at, id)", normalized)

    def test_live_followup_replays_cursor_triggers_and_preserves_merchant_guards(self) -> None:
        migration = (
            Path(__file__).resolve().parents[1]
            / "supabase/migrations/20260710133000_enrichment_cursor_live_followup.sql"
        ).read_text(encoding="utf-8")
        normalized = " ".join(migration.lower().split())
        self.assertIn("add column if not exists updated_at", normalized)
        self.assertIn("create trigger trg_venues_updated_at", normalized)
        self.assertIn("create trigger trg_dishes_updated_at", normalized)
        self.assertIn("current_user in ('postgres', 'supabase_admin', 'service_role')", normalized)
        self.assertIn("if auth.uid() is null then raise exception 'authentication required'", normalized)

    def test_known_source_scan_has_a_guarded_partial_index(self) -> None:
        migration = (
            Path(__file__).resolve().parents[1]
            / "supabase/migrations/20260710163000_enrichment_known_source_index.sql"
        ).read_text(encoding="utf-8")
        normalized = " ".join(migration.lower().split())
        self.assertIn("column_name = 'menu_source_url'", normalized)
        self.assertIn("create index if not exists idx_venues_known_menu_source_id", normalized)
        self.assertIn("on public.venues (id) where menu_source_url is not null", normalized)


if __name__ == "__main__":
    unittest.main()
