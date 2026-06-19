#!/usr/bin/env python3
"""Validate the Hungry F&B Excel workbook before staging import.

The validator is intentionally read-only. It checks workbook shape, row-level
data quality, and duplicate source keys, then prints a JSON report that can be
used before running the guarded staging importer.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "Hungry F&B Database.xlsx"
DEFAULT_SHEET = "Sheet1"

EXPECTED_COLUMNS = [
    "Business ID",
    "Stall ID",
    "Licence No",
    "Stall / Restaurant Name",
    "Business Name Normalized",
    "Stall Type",
    "Cuisine",
    "Signature Dish",
    "Menu",
    "Price",
    "Price Min (SGD)",
    "Price Max (SGD)",
    "Opening Hours",
    "Closed Days",
    "NEA Hygiene Grade",
    "Michelin Recognition",
    "Dietary Options",
    "Delivery Platforms",
    "Contact",
    "Instagram",
    "Image URL",
    "Hawker Centre / Building",
    "Building Name",
    "Unit No.",
    "Address",
    "Postal Code",
    "Region",
    "Location",
    "Location Key",
    "Latitude",
    "Longitude",
    "Source Type",
    "Source URL",
    "Last Verified Date",
    "Coverage Status",
    "Notes",
    "Google Place ID",
    "Google Display Name",
    "Google Formatted Address",
    "Google Latitude",
    "Google Longitude",
    "Google Primary Type",
    "Google Maps URI",
    "Google Rating",
    "Google Review Count",
    "Google Opening Hours JSON",
    "Google Price Level",
    "Google Website",
    "Google Photo Metadata",
    "Google Photo Last Refreshed",
    "Menu Source",
    "Menu Items",
    "Dishes Image URL",
    "Area",
]

VALID_STALL_TYPES = {
    "Restaurant",
    "Chain Outlet",
    "Cafe",
    "Bar / Bistro",
    "Hawker Stall",
    "Food Court Stall",
    "Other",
    "Zi Char",
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    normalized = str(value).strip()
    return normalized or None


def normalize_key(value: Any) -> str | None:
    value = clean(value)
    if not value:
        return None
    return re.sub(r"\s+", " ", value).casefold()


def parse_decimal(value: Any) -> Decimal | None:
    value = clean(value)
    if not value:
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def parse_float(value: Any) -> float | None:
    value = clean(value)
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def parse_int(value: Any) -> int | None:
    parsed = parse_decimal(value)
    if parsed is None:
        return None
    try:
        return int(parsed)
    except (OverflowError, ValueError):
        return None


def parse_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    value = clean(value)
    if not value:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return value
    return None


def dedupe_key(row: dict[str, Any]) -> str | None:
    for column in ("Business ID", "Stall ID", "Licence No"):
        value = normalize_key(row.get(column))
        if value:
            return f"{column}:{value}"

    name = normalize_key(row.get("Business Name Normalized") or row.get("Stall / Restaurant Name"))
    postal = normalize_key(row.get("Postal Code"))
    unit = normalize_key(row.get("Unit No."))
    if name and (postal or unit):
        return f"name-location:{name}:{postal or ''}:{unit or ''}"
    return None


def issue(row_number: int, code: str, message: str, severity: str = "error") -> dict[str, Any]:
    return {
        "row_number": row_number,
        "error_code": code,
        "message": message,
        "severity": severity,
    }


def load_workbook_rows(
    workbook_path: Path,
    sheet_name: str,
    limit: int | None,
) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"Workbook is missing sheet {sheet_name!r}. Available: {', '.join(workbook.sheetnames)}")

    worksheet = workbook[sheet_name]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean(value) or "" for value in header_row]
    rows: list[dict[str, Any]] = []
    for row_index, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: row_values[index] if index < len(row_values) else None for index, header in enumerate(headers)}
        row["_row_number"] = row_index
        rows.append(row)
        if limit is not None and len(rows) >= limit:
            break
    return headers, rows


def validate_row(row: dict[str, Any], duplicate_counts: Counter[str]) -> list[dict[str, Any]]:
    row_number = int(row["_row_number"])
    issues: list[dict[str, Any]] = []
    name = clean(row.get("Stall / Restaurant Name"))
    stall_type = clean(row.get("Stall Type"))
    latitude = parse_float(row.get("Latitude"))
    longitude = parse_float(row.get("Longitude"))
    price_min = parse_decimal(row.get("Price Min (SGD)"))
    price_max = parse_decimal(row.get("Price Max (SGD)"))
    postal_code = clean(row.get("Postal Code"))
    last_verified = parse_date(row.get("Last Verified Date"))
    google_rating = parse_decimal(row.get("Google Rating"))
    google_reviews = parse_int(row.get("Google Review Count"))
    google_price_level = parse_int(row.get("Google Price Level"))

    if not name or len(name) < 2:
        issues.append(issue(row_number, "missing_stall_name", "Stall / Restaurant Name is required."))
    if latitude is None:
        issues.append(issue(row_number, "missing_latitude", "Latitude is required and must be numeric."))
    elif not 1.1 <= latitude <= 1.6:
        issues.append(issue(row_number, "invalid_latitude", "Latitude must be within Singapore bounds."))
    if longitude is None:
        issues.append(issue(row_number, "missing_longitude", "Longitude is required and must be numeric."))
    elif not 103.5 <= longitude <= 104.2:
        issues.append(issue(row_number, "invalid_longitude", "Longitude must be within Singapore bounds."))
    if postal_code and not re.fullmatch(r"\d{5,6}", postal_code):
        issues.append(issue(row_number, "invalid_postal_code", "Postal Code must be 5 or 6 digits.", "warning"))
    if stall_type and stall_type not in VALID_STALL_TYPES:
        issues.append(issue(row_number, "invalid_stall_type", f"Unknown Stall Type: {stall_type}."))
    if clean(row.get("Price Min (SGD)")) and price_min is None:
        issues.append(issue(row_number, "invalid_price_min", "Price Min (SGD) must be numeric."))
    if clean(row.get("Price Max (SGD)")) and price_max is None:
        issues.append(issue(row_number, "invalid_price_max", "Price Max (SGD) must be numeric."))
    if price_min is not None and price_max is not None and price_min > price_max:
        issues.append(issue(row_number, "invalid_price_range", "Price Min (SGD) cannot exceed Price Max (SGD)."))
    if clean(row.get("Last Verified Date")) and last_verified is None:
        issues.append(issue(row_number, "invalid_last_verified_date", "Last Verified Date must be YYYY-MM-DD."))
    if google_rating is not None and not Decimal("0") <= google_rating <= Decimal("5"):
        issues.append(issue(row_number, "invalid_google_rating", "Google Rating must be between 0 and 5.", "warning"))
    if google_reviews is not None and google_reviews < 0:
        issues.append(issue(row_number, "invalid_google_review_count", "Google Review Count must be non-negative.", "warning"))
    if google_price_level is not None and not 0 <= google_price_level <= 4:
        issues.append(issue(row_number, "invalid_google_price_level", "Google Price Level must be 0 to 4.", "warning"))

    # Data Quality & Coverage checks
    coverage_status = clean(row.get("Coverage Status"))
    menu_source = clean(row.get("Menu Source"))
    image_url = clean(row.get("Image URL"))
    dishes_image_url = clean(row.get("Dishes Image URL"))

    if coverage_status == "Non-Consumer":
        issues.append(issue(row_number, "non_consumer_entity", f"Stall is flagged as non-consumer: {name}.", "warning"))

    if menu_source and ("Cuisine Category Seed" in menu_source or "category_seed" in menu_source):
        issues.append(issue(row_number, "generic_menu_fallback", "Stall uses generic menu fallback seed.", "warning"))

    generic_storefronts = {
        "photo-1517248135467-4c7edcad34c4", "photo-1552566626-52f8b828add9",
        "photo-1554118811-1e0d58224f24", "photo-1514933651103-005eec06c04b",
        "photo-1569058242253-92a9c755a0ec", "photo-1555396273-367ea4eb4db5"
    }
    if image_url:
        for img_id in generic_storefronts:
            if img_id in image_url:
                issues.append(issue(row_number, "generic_storefront_image", f"Stall uses generic storefront image: {image_url}.", "warning"))
                break

    generic_dishes = {
        "photo-1546069901-ba9599a7e63c", "photo-1507133750040-4a8f57021571",
        "photo-1568901346375-23c9450c58cd", "photo-1534080391025-09795d197a5b",
        "photo-1509440159596-0249088772ff", "photo-1585938338392-50a592202c7b"
    }
    if dishes_image_url:
        for img_id in generic_dishes:
            if img_id in dishes_image_url:
                issues.append(issue(row_number, "generic_dish_image", f"Stall uses generic dish image: {dishes_image_url}.", "warning"))
                break

    key = dedupe_key(row)
    if key and duplicate_counts[key] > 1:
        issues.append(issue(row_number, "duplicate_source_key", f"Duplicate source key: {key}.", "warning"))

    return issues


def validate_workbook(
    workbook_path: Path,
    sheet_name: str = DEFAULT_SHEET,
    limit: int | None = None,
) -> dict[str, Any]:
    headers, rows = load_workbook_rows(workbook_path, sheet_name, limit)
    missing = [column for column in EXPECTED_COLUMNS if column not in headers]
    extra = [column for column in headers if column and column not in EXPECTED_COLUMNS]
    duplicate_counts = Counter(key for row in rows if (key := dedupe_key(row)))

    row_reports = []
    issue_counts: Counter[str] = Counter()
    severity_counts: Counter[str] = Counter()
    for row in rows:
        issues = validate_row(row, duplicate_counts)
        for item in issues:
            issue_counts[item["error_code"]] += 1
            severity_counts[item["severity"]] += 1
        row_reports.append(
            {
                "row_number": row["_row_number"],
                "business_id": clean(row.get("Business ID")),
                "name": clean(row.get("Stall / Restaurant Name")),
                "status": "invalid" if any(item["severity"] == "error" for item in issues) else "valid",
                "issues": issues,
            }
        )

    invalid_rows = sum(1 for row in row_reports if row["status"] == "invalid")
    report = {
        "workbook": str(workbook_path),
        "sheet": sheet_name,
        "headers_ok": not missing,
        "missing_columns": missing,
        "extra_columns": extra,
        "rows_checked": len(rows),
        "valid_rows": len(rows) - invalid_rows,
        "invalid_rows": invalid_rows,
        "issue_counts": dict(sorted(issue_counts.items())),
        "severity_counts": dict(sorted(severity_counts.items())),
        "duplicate_source_key_count": sum(1 for count in duplicate_counts.values() if count > 1),
        "sample_issues": [
            row for row in row_reports if row["issues"]
        ][:25],
    }
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--allow-row-errors",
        action="store_true",
        help="Exit 0 even when row-level validation errors are found.",
    )
    args = parser.parse_args()

    report = validate_workbook(args.workbook, args.sheet, args.limit)
    rendered = json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)

    if not report["headers_ok"]:
        return 2
    if report["invalid_rows"] and not args.allow_row_errors:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
